"""Testy drzwi kart czasu (ADR 0035) — stan idempotencji, blokada instancji, konfiguracja.

Bez sieci: sprawdzamy to, co decyduje o tym, czy ktoś dostanie wiadomość DWA razy albo wcale.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest

from workmate.adapters.inbound.single_instance import (
    AlreadyRunningError,
    acquire_single_instance_lock,
)
from workmate.adapters.inbound.worklogi import app
from workmate.adapters.inbound.worklogi import state as state_store
from workmate.config import TeamsPushSettings, WorklogiSettings
from workmate.core.domain.week import reported_week, week_label

# --- stan idempotencji ------------------------------------------------------------


def test_roundtrip_preserves_entries(tmp_path: Path) -> None:
    path = tmp_path / "state.json"
    state_store.save(path, {"2026-W29:EMP-042": "2026-07-17T16:00:00+02:00"})
    assert state_store.load(path) == {"2026-W29:EMP-042": "2026-07-17T16:00:00+02:00"}


def test_missing_state_starts_empty(tmp_path: Path) -> None:
    assert state_store.load(tmp_path / "nie-ma.json") == {}


def test_corrupt_state_starts_empty_instead_of_crashing(tmp_path: Path) -> None:
    """Martwy proces nie wyśle nic NIKOMU przez tydzień — gorsze niż nadmiarowa wiadomość."""
    path = tmp_path / "state.json"
    path.write_text("{nie json", encoding="utf-8")
    assert state_store.load(path) == {}


def test_state_of_wrong_shape_starts_empty(tmp_path: Path) -> None:
    path = tmp_path / "state.json"
    path.write_text("[1,2,3]", encoding="utf-8")
    assert state_store.load(path) == {}


def test_save_is_atomic_leaving_no_temp_file(tmp_path: Path) -> None:
    path = tmp_path / "state.json"
    state_store.save(path, {"a": "b"})
    assert list(p.name for p in tmp_path.iterdir()) == ["state.json"]


def test_save_creates_missing_parent_directories(tmp_path: Path) -> None:
    path = tmp_path / "glebiej" / "state.json"
    state_store.save(path, {"a": "b"})
    assert path.exists()


def test_key_is_scoped_per_week_and_person() -> None:
    assert state_store.key("2026-W29", "EMP-042") == "2026-W29:EMP-042"
    assert state_store.key("2026-W30", "EMP-042") != state_store.key("2026-W29", "EMP-042")


def test_prune_keeps_only_requested_weeks() -> None:
    state = {
        "2026-W29:EMP-042": "x",
        "2026-W28:EMP-042": "x",
        "2026-W10:EMP-017": "x",
    }
    pruned = state_store.prune(state, keep_weeks=("2026-W29", "2026-W28"))
    assert set(pruned) == {"2026-W29:EMP-042", "2026-W28:EMP-042"}


# --- blokada jednej instancji -----------------------------------------------------


def test_second_instance_is_refused(tmp_path: Path) -> None:
    """Dwa procesy obeszłyby idempotencję i wysłały ludziom po dwie wiadomości."""
    path = tmp_path / "state.json"
    with acquire_single_instance_lock(path), pytest.raises(AlreadyRunningError, match="Inna"):
        acquire_single_instance_lock(path)


def test_lock_is_released_after_the_context_exits(tmp_path: Path) -> None:
    path = tmp_path / "state.json"
    with acquire_single_instance_lock(path):
        pass
    with acquire_single_instance_lock(path):  # ponowne zajęcie musi się udać
        pass


def test_lock_file_records_the_pid(tmp_path: Path) -> None:
    """PID w pliku pozwala operatorowi znaleźć proces trzymający blokadę.

    Czytamy PO zwolnieniu: na Windows ``msvcrt`` blokuje bajt 0 wyłącznie, więc odczyt
    w trakcie trzymania blokady kończy się ``PermissionError``.
    """
    import os

    path = tmp_path / "state.json"
    with acquire_single_instance_lock(path):
        pass
    assert (tmp_path / "state.json.lock").read_text(encoding="utf-8") == str(os.getpid())


# --- konfiguracja -----------------------------------------------------------------


def _settings(**kw) -> WorklogiSettings:
    base: dict = {
        "enabled": True,
        "output_dir": Path("D:/worklogi"),
        "team_id": "team-1",
        "hours_path": Path("hours.json"),
    }
    base.update(kw)
    return WorklogiSettings(**base)


def test_disabled_settings_pass_validation(tmp_path: Path) -> None:
    WorklogiSettings().validate(data_dir=tmp_path / "data")


@pytest.mark.parametrize(
    ("field", "expected"),
    [
        ("output_dir", "OUTPUT_DIR"),
        ("identities_path", "IDENTITIES"),
        ("team_id", "TEAM_ID"),
    ],
)
def test_enabled_requires_targets(field: str, expected: str, tmp_path: Path) -> None:
    empty = Path() if field != "team_id" else ""
    with pytest.raises(ValueError, match=expected):
        _settings(**{field: empty}).validate(data_dir=tmp_path / "data")


def test_missing_identity_file_is_a_hard_error(tmp_path: Path) -> None:
    """Bez mapy tożsamości NIKT nie dostanie arkusza — to musi paść przy starcie."""
    with pytest.raises(ValueError, match="mapa tożsamości nie istnieje"):
        _settings(identities_path=tmp_path / "nie-ma.yaml").validate(data_dir=tmp_path / "data")


def _identity_file(tmp_path: Path) -> Path:
    """Minimalna, POPRAWNA mapa tożsamości — walidacja ma paść na czymś innym niż jej brak."""
    path = tmp_path / "id.yaml"
    path.write_text("EMP-1:\n  aad_user_id: a\n  jira_user: b\n", encoding="utf-8")
    return path


def test_output_dir_inside_data_is_rejected(tmp_path: Path) -> None:
    """Arkusze z godzinami ludzi nie mogą trafić do bazy wiedzy, którą agent czyta."""
    data = tmp_path / "data"
    identities = _identity_file(tmp_path)
    with pytest.raises(ValueError, match="nie może leżeć wewnątrz katalogu danych"):
        _settings(output_dir=data / "arkusze", identities_path=identities).validate(data_dir=data)


def test_output_dir_inside_the_repository_is_rejected(tmp_path: Path) -> None:
    """Imienne godziny w drzewie roboczym czekają na pierwsze ``git add .`` — i zostają w historii.

    Dokumentacja obiecywała „poza data/ I poza repo" od początku; kontrola sprawdzała tylko
    pierwszą połowę, więc ``OUTPUT_DIR=arkusze`` przechodziło bez słowa.
    """
    identities = _identity_file(tmp_path)
    inside_repo = Path(__file__).resolve().parents[3] / "arkusze"
    with pytest.raises(ValueError, match="wewnątrz repozytorium"):
        _settings(output_dir=inside_repo, identities_path=identities).validate(
            data_dir=tmp_path / "data"
        )


# --- bramka potwierdzenia nagłówków (D6) ------------------------------------------


def _confirmable(tmp_path: Path, **kw) -> WorklogiSettings:
    return _settings(identities_path=_identity_file(tmp_path), output_dir=tmp_path / "out", **kw)


def test_live_run_refuses_to_start_with_unconfirmed_headers(tmp_path: Path) -> None:
    """Nagłówki są HIPOTEZĄ — bez potwierdzenia szablonem tryb bojowy rozesłałby makulaturę.

    WorklogPRO dopasowuje kolumny PO NAZWIE, więc jedna literówka unieważnia KAŻDY plik.
    Bez tej bramki dowiedzielibyśmy się o tym dopiero od ludzi, którym import nie przeszedł.
    """
    with pytest.raises(ValueError, match="HEADERS_CONFIRMED"):
        _confirmable(tmp_path, dry_run=False).validate(data_dir=tmp_path / "data")


def test_dry_run_starts_without_confirmation(tmp_path: Path) -> None:
    """Przebieg na sucho MA działać — to on generuje arkusz do porównania z szablonem."""
    _confirmable(tmp_path, dry_run=True).validate(data_dir=tmp_path / "data")


def test_live_run_starts_once_headers_are_confirmed(tmp_path: Path) -> None:
    _confirmable(tmp_path, dry_run=False, headers_confirmed=True).validate(
        data_dir=tmp_path / "data"
    )


def test_headers_confirmation_defaults_to_false() -> None:
    """Domyślnie NIEPOTWIERDZONE — to musi być czynność operatora, nie stan zastany."""
    assert WorklogiSettings().headers_confirmed is False


def test_unknown_timezone_is_rejected_even_when_disabled(tmp_path: Path) -> None:
    """Literówka w strefie nie może spać do dnia, w którym ktoś przestawi bramkę."""
    with pytest.raises(ValueError, match="strefą czasową"):
        WorklogiSettings(tz_name="Nie/Ma").validate(data_dir=tmp_path / "data")


@pytest.mark.parametrize(
    ("field", "value", "expected"),
    [
        ("run_weekday", 9, "RUN_WEEKDAY"),
        ("run_hour", 25, "RUN_HOUR"),
        ("run_minute", 61, "RUN_MINUTE"),
        ("start_hour", -1, "START_HOUR"),
        ("max_hours_per_day", 0, "MAX_HOURS_PER_DAY"),
        ("max_hours_per_day", 99.0, "MAX_HOURS_PER_DAY"),
        ("max_catchup_days", 99, "MAX_CATCHUP_DAYS"),
    ],
)
def test_out_of_range_values_are_rejected(field, value, expected, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match=expected):
        WorklogiSettings(**{field: value}).validate(data_dir=tmp_path / "data")


def test_defaults_are_conservative() -> None:
    """Bramka OFF, tryb próbny ON, piątek 16:00 — nic nie wychodzi bez świadomej decyzji."""
    settings = WorklogiSettings()
    assert settings.enabled is False
    assert settings.dry_run is True
    assert (settings.run_weekday, settings.run_hour) == (4, 16)


def test_from_env_reads_the_surface(monkeypatch) -> None:
    monkeypatch.setenv("WORKMATE_WORKLOGI_ENABLED", "true")
    monkeypatch.setenv("WORKMATE_WORKLOGI_DRY_RUN", "false")
    monkeypatch.setenv("WORKMATE_WORKLOGI_TEAM_ID", " team-9 ")
    monkeypatch.setenv("WORKMATE_WORKLOGI_ONLY_SOURCE_IDS", "EMP-042, EMP-017")
    monkeypatch.setenv("WORKMATE_WORKLOGI_RUN_HOUR", "15")
    settings = WorklogiSettings.from_env()
    assert settings.enabled is True and settings.dry_run is False
    assert settings.team_id == "team-9"
    assert settings.only_source_ids == ("EMP-042", "EMP-017")
    assert settings.run_hour == 15


# --- źródło "shifts" (ADR 0036) ---------------------------------------------------


def _shifts(tmp_path: Path, **kw) -> WorklogiSettings:
    base: dict = {
        "hours_source": "shifts",
        "identities_path": _identity_file(tmp_path),
        "output_dir": tmp_path / "out",
        "fallback_issue": "BIAP-1",
        "summary_dir": tmp_path / "summaries",
    }
    base.update(kw)
    return _settings(**base)


def test_shifts_source_passes_validation_with_its_targets(tmp_path: Path) -> None:
    _shifts(tmp_path).validate(data_dir=tmp_path / "data")


def test_shifts_requires_fallback_issue(tmp_path: Path) -> None:
    """Bez koszyka dni bez klucza z commitów dałyby wiersze bez issue_key — import by padł."""
    with pytest.raises(ValueError, match="FALLBACK_ISSUE"):
        _shifts(tmp_path, fallback_issue="").validate(data_dir=tmp_path / "data")


def test_shifts_fallback_issue_must_look_like_a_jira_key(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="PROJ-123"):
        _shifts(tmp_path, fallback_issue="BADKEY").validate(data_dir=tmp_path / "data")


def test_shifts_requires_summary_dir(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="SUMMARY_DIR"):
        _shifts(tmp_path, summary_dir=Path()).validate(data_dir=tmp_path / "data")


def test_json_source_does_not_require_shifts_fields(tmp_path: Path) -> None:
    """Zero zmian zachowania dla JSON: koszyk i katalog claude_summary są nieistotne."""
    _settings(identities_path=_identity_file(tmp_path), output_dir=tmp_path / "out").validate(
        data_dir=tmp_path / "data"
    )


def test_from_env_reads_shifts_fields(monkeypatch) -> None:
    monkeypatch.setenv("WORKMATE_WORKLOGI_HOURS_SOURCE", "shifts")
    monkeypatch.setenv("WORKMATE_WORKLOGI_FALLBACK_ISSUE", "biap-2")  # normalizowane do wielkich
    monkeypatch.setenv("WORKMATE_WORKLOGI_SUMMARY_DIR", "D:/summaries")
    settings = WorklogiSettings.from_env()
    assert settings.hours_source == "shifts"
    assert settings.fallback_issue == "BIAP-2"
    assert str(settings.summary_dir) not in ("", ".")


def test_unknown_hours_source_is_rejected(tmp_path: Path) -> None:
    """Rozszerzenie listy o „shifts" nie może rozszczelnić bramki — literówka w źródle godzin
    ma PAŚĆ przy starcie, a nie po cichu udać JSON (ADR 0036 dołożył tylko jedną wartość)."""
    with pytest.raises(ValueError, match="HOURS_SOURCE musi być jednym z"):
        _settings(
            identities_path=_identity_file(tmp_path),
            output_dir=tmp_path / "out",
            hours_source="rcp",
        ).validate(data_dir=tmp_path / "data")


def test_teams_push_scopes_include_team_members() -> None:
    """Lista osób z Graph wymaga TeamMember.Read.All — ta sama aplikacja co Powiadomienia_teams."""
    from workmate.config import TeamsPushSettings

    assert "TeamMember.Read.All" in TeamsPushSettings().scopes


# --- nadrabianie i przeżywalność pętli --------------------------------------------


def _friday(day: int) -> datetime:
    """Piątek 16:00 w Europe/Warsaw — domyślny termin przebiegu."""
    return datetime(2026, 7, day, 16, 0, tzinfo=ZoneInfo("Europe/Warsaw"))


def test_missed_deadline_returns_the_deadline_not_just_a_flag(tmp_path: Path, monkeypatch) -> None:
    """Nadrabianie w poniedziałek musi wskazać PIĄTKOWY termin, nie »teraz«.

    Regresja: funkcja zwracała ``bool``, a tydzień raportu liczono z zegara. Piątkowy termin
    i poniedziałkowe podniesienie leżą po DWÓCH stronach granicy tygodnia, więc raport
    przeskakiwał o siedem dni: tydzień, który przepadł, nie trafiał do nikogo NIGDY, a osoby
    zapisane pod etykietą następnego były pomijane w jego prawdziwym przebiegu — traciły oba.
    """
    monday = datetime(2026, 7, 20, 9, 0, tzinfo=ZoneInfo("Europe/Warsaw"))
    monkeypatch.setattr(app, "_now", lambda: monday)

    settings = _settings(state_path=tmp_path / "s.json")
    missed = app._missed_deadline(settings, ZoneInfo("Europe/Warsaw"))

    assert missed == _friday(17)  # piątkowy termin, nie poniedziałkowe „teraz"
    # Termin z piątku W29 raportuje tydzień ZAMKNIĘTY, czyli W28. Liczone z „teraz"
    # (poniedziałek W30) wyszłoby W29 — o tydzień za daleko, i to jest właśnie ten błąd.
    warsaw = ZoneInfo("Europe/Warsaw")
    assert week_label(reported_week(missed, warsaw)[0]) == "2026-W28"
    assert week_label(reported_week(_friday(17), warsaw)[0]) != week_label(
        reported_week(datetime(2026, 7, 20, 9, 0, tzinfo=warsaw), warsaw)[0]
    )


def test_missed_deadline_is_none_when_the_week_was_already_reported(
    tmp_path: Path, monkeypatch
) -> None:
    monday = datetime(2026, 7, 20, 9, 0, tzinfo=ZoneInfo("Europe/Warsaw"))
    monkeypatch.setattr(app, "_now", lambda: monday)
    state = tmp_path / "s.json"
    state_store.save(state, {"2026-W28:EMP-042": monday.isoformat()})

    assert app._missed_deadline(_settings(state_path=state), ZoneInfo("Europe/Warsaw")) is None


def test_failed_run_does_not_kill_the_loop(monkeypatch, caplog) -> None:
    """Awaria przebiegu (np. plik godzin w trakcie zapisu) nie może wywrócić procesu.

    Bez tego systemd restartował usługę, nadrabianie znów było należne, znów padało —
    pętla restartów, w której nikt nie dostaje nic.
    """

    def boom(*_args, **_kwargs):
        raise RuntimeError("plik godzin w trakcie zapisu")

    monkeypatch.setattr(app, "_run_once", boom)

    app._safe_run_once(_settings(), TeamsPushSettings(), lambda: "t", as_of=_friday(17))

    assert "nie powiódł się" in caplog.text


# --- przycinanie stanu ------------------------------------------------------------


def test_prune_keeps_the_week_that_was_just_reported(tmp_path: Path) -> None:
    """Nadrabianie nie może wyciąć tygodnia, który WŁAŚNIE zapisało — stąd ``moment``, nie zegar."""
    state = tmp_path / "s.json"
    settings = _settings(state_path=state, dry_run=False)
    tz = ZoneInfo("Europe/Warsaw")
    moment = _friday(24)
    label = week_label(reported_week(moment, tz)[0])
    saved = {f"{label}:EMP-042": "x", "2020-W01:EMP-042": "stare"}

    app._prune_state(settings, saved, tz, moment)

    kept = state_store.load(state)
    assert f"{label}:EMP-042" in kept
    assert "2020-W01:EMP-042" not in kept


def test_prune_writes_nothing_in_dry_run(tmp_path: Path) -> None:
    """Tryb PRÓBNY nie dotyka stanu — także przy przycinaniu; „na sucho" nie zostawia śladu."""
    state = tmp_path / "s.json"
    settings = _settings(state_path=state, dry_run=True)
    app._prune_state(
        settings, {"2020-W01:EMP-042": "stare"}, ZoneInfo("Europe/Warsaw"), _friday(24)
    )

    assert not state.exists()


# --- wymagania startowe -----------------------------------------------------------


def test_missing_teams_identity_is_a_hard_start_error() -> None:
    """Bez tożsamości Graph nie ma jak wysłać — lepiej nie wystartować niż milczeć w piątek."""
    with pytest.raises(SystemExit, match="CLIENT_ID"):
        app._require_teams(TeamsPushSettings(client_id="", tenant_id=""))


def test_configured_teams_identity_passes() -> None:
    app._require_teams(TeamsPushSettings(client_id="c", tenant_id="t"))


# --- pełne wpięcie drzwi (przebieg na sucho, bez Graph) ---------------------------


def _wired(tmp_path: Path, monkeypatch, *, hours: str) -> WorklogiSettings:
    """Postaw prawdziwe adaptery (JSON, xlsx, YAML) i podmień JEDYNIE listę członków z Graph."""
    identities = tmp_path / "id.yaml"
    identities.write_text(
        "EMP-042:\n  aad_user_id: aad-mikolaj\n  jira_user: mikolaj@example.com\n", encoding="utf-8"
    )
    hours_path = tmp_path / "hours.json"
    hours_path.write_text(hours, encoding="utf-8")
    monkeypatch.setattr(
        "workmate.adapters.outbound.graph_identity_directory.fetch_team_members",
        lambda *_a, **_k: {"aad-mikolaj": "Mikołaj"},
    )
    return _settings(
        output_dir=tmp_path / "out",
        identities_path=identities,
        hours_path=hours_path,
        state_path=tmp_path / "s.json",
        dry_run=True,
    )


def test_run_once_reports_the_week_of_the_deadline_not_of_the_clock(
    tmp_path: Path, monkeypatch
) -> None:
    """Sedno nadrabiania: raport liczy się z ``as_of``, więc tydzień, który przepadł, wraca.

    Bez tego podniesienie w poniedziałek raportowało tydzień bieżący, a ten sprzed awarii
    nie trafiał do nikogo NIGDY.
    """
    settings = _wired(
        tmp_path,
        monkeypatch,
        hours='[{"source_id": "EMP-042", "day": "2026-07-15", "issue_key": "WT-12", "hours": 3}]',
    )
    monkeypatch.setattr(app, "_now", lambda: datetime(2026, 8, 30, 9, 0, tzinfo=ZoneInfo("UTC")))

    report = app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert report.week_label == "2026-W29"  # tydzień terminu, nie „teraz" (koniec sierpnia)
    assert [outcome.source_id for outcome in report.sent] == ["EMP-042"]


def test_run_once_writes_the_sheet_but_sends_nothing_in_dry_run(
    tmp_path: Path, monkeypatch
) -> None:
    """Tryb PRÓBNY ma dać artefakt do obejrzenia i NIC poza tym — ani wiadomości, ani stanu."""
    settings = _wired(
        tmp_path,
        monkeypatch,
        hours='[{"source_id": "EMP-042", "day": "2026-07-15", "issue_key": "WT-12", "hours": 3}]',
    )

    app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert list((tmp_path / "out").glob("*.xlsx"))
    assert not (tmp_path / "s.json").exists()


def test_run_once_fails_closed_for_a_person_outside_the_identity_map(
    tmp_path: Path, monkeypatch
) -> None:
    """Nieznane ``source_id`` NIE dostaje pliku ani wiadomości — zgadywanie tożsamości jest gorsze.

    Zły ``jira_user`` zaimportowałby czyjeś godziny na CUDZE konto Jiry, create-only.
    """
    settings = _wired(
        tmp_path,
        monkeypatch,
        hours='[{"source_id": "OBCY-1", "day": "2026-07-15", "issue_key": "WT-12", "hours": 3}]',
    )

    report = app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert [outcome.reason for outcome in report.failed] == ["unknown_person"]
    assert not list((tmp_path / "out").glob("*.xlsx"))


def test_build_hours_source_dispatches_on_configured_source(tmp_path: Path) -> None:
    """Fabryka wybiera źródło po configu — bez sieci (konstrukcja adaptera nie dotyka Graph)."""
    from workmate.adapters.outbound.json_hours_source import JsonHoursSource
    from workmate.core.application.shift_hours_source import ShiftsHoursSource

    tz = ZoneInfo("Europe/Warsaw")
    json_src = app._build_hours_source(
        _settings(hours_path=tmp_path / "h.json"), object(), lambda: "t", tz
    )
    assert isinstance(json_src, JsonHoursSource)
    shifts_src = app._build_hours_source(
        _settings(hours_source="shifts", fallback_issue="BIAP-1"), object(), lambda: "t", tz
    )
    assert isinstance(shifts_src, ShiftsHoursSource)


def test_run_once_with_shifts_writes_a_sheet_from_real_hours(tmp_path: Path, monkeypatch) -> None:
    """Deliverable S2: dry-run daje arkusz z REALNYCH godzin Shifts na koszykowym issue (ADR 0036).

    Podmieniamy JEDYNIE odczyt bloków z Graph (``read_blocks``); reszta ścieżki — mapowanie
    aad→osoba, budowa zestawienia, zapis xlsx — jest prawdziwa.
    """
    from dataclasses import replace
    from datetime import timezone

    from workmate.adapters.outbound import graph_shift_source
    from workmate.core.domain.shift_hours import ShiftBlock

    block = ShiftBlock(
        user_id="aad-mikolaj",  # ta sama osoba co w mapie tożsamości z ``_wired``
        start=datetime(2026, 7, 15, 6, 0, tzinfo=timezone.utc),  # 08:00–16:00 lokalnie = 8h
        end=datetime(2026, 7, 15, 14, 0, tzinfo=timezone.utc),
    )
    monkeypatch.setattr(graph_shift_source.GraphShiftSource, "read_blocks", lambda self: [block])
    monkeypatch.setattr(app, "_build_commit_source", lambda tz: None)  # bez GitHuba → koszyk
    settings = replace(
        _wired(tmp_path, monkeypatch, hours="[]"),
        hours_source="shifts",
        fallback_issue="BIAP-1",
        summary_dir=tmp_path / "sum",
    )

    report = app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert [outcome.source_id for outcome in report.sent] == ["EMP-042"]
    assert report.sent[0].minutes == 480  # 8h realnie ze zmiany, nie estymacja
    assert list((tmp_path / "out").glob("*.xlsx"))


def test_run_once_with_shifts_attributes_hours_to_commit_issue_keys(
    tmp_path: Path, monkeypatch
) -> None:
    """Deliverable S3: godziny Shifts trafiają na klucz Jira z commitu dnia (nie na koszyk).

    Podmieniamy odczyt zmian (``read_blocks``) i źródło commitów (``_build_commit_source``);
    reszta — mapowanie aad→osoba, przypisanie issue, zapis xlsx — jest prawdziwa. Czytamy
    wygenerowany arkusz i sprawdzamy kolumnę ``Issue Key/ID``.
    """
    from datetime import timezone

    from openpyxl import load_workbook

    from workmate.adapters.outbound import graph_shift_source
    from workmate.core.domain.shift_hours import ShiftBlock
    from workmate.core.domain.worklog import Commit

    identities = tmp_path / "id.yaml"
    identities.write_text(
        "EMP-042:\n  aad_user_id: aad-mikolaj\n  jira_user: mikolaj@example.com\n"
        "  git_email: mikolaj@example.com\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "workmate.adapters.outbound.graph_identity_directory.fetch_team_members",
        lambda *_a, **_k: {"aad-mikolaj": "Mikołaj"},
    )
    block = ShiftBlock(
        user_id="aad-mikolaj",
        start=datetime(2026, 7, 15, 6, 0, tzinfo=timezone.utc),  # 8h dnia 15
        end=datetime(2026, 7, 15, 14, 0, tzinfo=timezone.utc),
    )
    monkeypatch.setattr(graph_shift_source.GraphShiftSource, "read_blocks", lambda self: [block])

    class FakeCommits:
        def commits_for(self, git_email, since, until):
            return [
                Commit(
                    sha="s",
                    message="WT-7 robota",
                    authored_at=datetime(2026, 7, 15, 9, 0, tzinfo=timezone.utc),
                )
            ]

    monkeypatch.setattr(app, "_build_commit_source", lambda tz: FakeCommits())

    settings = _settings(
        hours_source="shifts",
        identities_path=identities,
        output_dir=tmp_path / "out",
        fallback_issue="BIAP-1",
        summary_dir=tmp_path / "sum",
        state_path=tmp_path / "s.json",
        dry_run=True,
    )

    report = app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert report.sent[0].minutes == 480
    xlsx = next((tmp_path / "out").glob("*.xlsx"))
    rows = [tuple(r) for r in load_workbook(xlsx).worksheets[0].iter_rows(values_only=True)]
    assert rows[1][0] == "WT-7"  # kolumna Issue Key/ID — realny klucz z commitu, nie koszyk


def test_run_once_with_shifts_fills_comment_from_claude_summary(
    tmp_path: Path, monkeypatch
) -> None:
    """Deliverable S4: kolumna ``Comment`` arkusza niesie opis dnia z claude_summary.

    Prawdziwy store (katalog JSON), podmieniony jest tylko odczyt zmian i commitów.
    """
    import json as _json
    from datetime import timezone

    from openpyxl import load_workbook

    from workmate.adapters.outbound import graph_shift_source
    from workmate.core.domain.shift_hours import ShiftBlock
    from workmate.core.domain.worklog import Commit

    identities = tmp_path / "id.yaml"
    identities.write_text(
        "EMP-042:\n  aad_user_id: aad-mikolaj\n  jira_user: mikolaj@example.com\n"
        "  git_email: mikolaj@example.com\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "workmate.adapters.outbound.graph_identity_directory.fetch_team_members",
        lambda *_a, **_k: {"aad-mikolaj": "Mikołaj"},
    )
    summary_dir = tmp_path / "sum"
    summary_dir.mkdir()
    (summary_dir / "mikolaj.json").write_text(
        _json.dumps(
            {
                "person": "mikolaj@example.com",
                "days": [
                    {"date": "2026-07-15", "llm_prose": "Robił WT-7.", "commits": [], "prompts": []}
                ],
            }
        ),
        encoding="utf-8",
    )
    block = ShiftBlock(
        user_id="aad-mikolaj",
        start=datetime(2026, 7, 15, 6, 0, tzinfo=timezone.utc),
        end=datetime(2026, 7, 15, 14, 0, tzinfo=timezone.utc),
    )
    monkeypatch.setattr(graph_shift_source.GraphShiftSource, "read_blocks", lambda self: [block])

    class FakeCommits:
        def commits_for(self, git_email, since, until):
            return [
                Commit(
                    sha="s",
                    message="WT-7 robota",
                    authored_at=datetime(2026, 7, 15, 9, 0, tzinfo=timezone.utc),
                )
            ]

    monkeypatch.setattr(app, "_build_commit_source", lambda tz: FakeCommits())

    settings = _settings(
        hours_source="shifts",
        identities_path=identities,
        output_dir=tmp_path / "out",
        fallback_issue="BIAP-1",
        summary_dir=summary_dir,
        state_path=tmp_path / "s.json",
        dry_run=True,
    )

    app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    xlsx = next((tmp_path / "out").glob("*.xlsx"))
    rows = [tuple(r) for r in load_workbook(xlsx).worksheets[0].iter_rows(values_only=True)]
    assert rows[1][0] == "WT-7"  # Issue Key/ID — klucz z commitu
    assert rows[1][4] == "Robił WT-7."  # Comment — opis dnia z claude_summary


def test_run_once_with_shifts_full_assembly_hours_issues_and_comment(
    tmp_path: Path, monkeypatch
) -> None:
    """Złożenie S1-S5: realne godziny Shifts DZIELONE na dwa klucze z commitów, oba z opisem dnia.

    Jeden pełny bieg na sucho przez prawdziwe adaptery (store, xlsx, mapa tożsamości, projekcja);
    podmienione tylko wejścia sieciowe: odczyt zmian (Graph) i commitów (GitHub).
    """
    from datetime import timezone

    from openpyxl import load_workbook

    from workmate.adapters.outbound import graph_shift_source
    from workmate.core.domain.shift_hours import ShiftBlock
    from workmate.core.domain.worklog import Commit

    identities = tmp_path / "id.yaml"
    identities.write_text(
        "EMP-042:\n  aad_user_id: aad-mikolaj\n  jira_user: mikolaj@example.com\n"
        "  git_email: mikolaj@example.com\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "workmate.adapters.outbound.graph_identity_directory.fetch_team_members",
        lambda *_a, **_k: {"aad-mikolaj": "Mikołaj"},
    )
    summary_dir = tmp_path / "sum"
    summary_dir.mkdir()
    (summary_dir / "m.json").write_text(
        __import__("json").dumps(
            {
                "person": "mikolaj@example.com",
                "days": [
                    {
                        "date": "2026-07-15",
                        "llm_prose": "Robił WT-1 i WT-2.",
                        "commits": [],
                        "prompts": [],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    block = ShiftBlock(
        user_id="aad-mikolaj",
        start=datetime(2026, 7, 15, 6, 0, tzinfo=timezone.utc),  # 8h → 480 min
        end=datetime(2026, 7, 15, 14, 0, tzinfo=timezone.utc),
    )
    monkeypatch.setattr(graph_shift_source.GraphShiftSource, "read_blocks", lambda self: [block])

    class FakeCommits:
        def commits_for(self, git_email, since, until):
            at = datetime(2026, 7, 15, 9, 0, tzinfo=timezone.utc)
            return [Commit(sha="a", message="WT-1 rano", authored_at=at),
                    Commit(sha="b", message="WT-2 po", authored_at=at)]

    monkeypatch.setattr(app, "_build_commit_source", lambda tz: FakeCommits())

    settings = _settings(
        hours_source="shifts",
        identities_path=identities,
        output_dir=tmp_path / "out",
        fallback_issue="BIAP-1",
        summary_dir=summary_dir,
        state_path=tmp_path / "s.json",
        dry_run=True,
    )

    report = app._run_once(settings, TeamsPushSettings(), lambda: "tok", as_of=_friday(24))

    assert report.sent[0].minutes == 480  # suma dnia = realne 8h, mimo podziału na dwa klucze
    xlsx = next((tmp_path / "out").glob("*.xlsx"))
    rows = [tuple(r) for r in load_workbook(xlsx).worksheets[0].iter_rows(values_only=True)]
    data = sorted(rows[1:], key=lambda r: r[0])  # po Issue Key/ID
    assert [(r[0], r[2]) for r in data] == [("WT-1", "4h"), ("WT-2", "4h")]  # 240 min = 4h każdy
    assert {r[4] for r in data} == {"Robił WT-1 i WT-2."}  # opis dnia na obu wierszach
