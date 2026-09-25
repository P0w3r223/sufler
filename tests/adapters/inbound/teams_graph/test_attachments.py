"""Testy materializacji załączników Teams (ADR 0016) — referencja → bajty → ``Attachment``.

I/O portu ``GraphChannelClient`` jest wstrzykiwane, więc pełną materializację testujemy
ATRAPĄ portu — bez ``httpx``, bez MSAL, bez sieci. Sedno: sniff obrazu po ZAWARTOŚCI
(magic/Pillow, nie po rozszerzeniu) + downscaling, ekstrakcja tekstu z dokumentów
(docx/xlsx/pptx/pliki tekstowe) oraz łagodna degradacja (limit rozmiaru/liczby, nieobsługiwany
typ, błąd pobrania/uszkodzony plik → NOTKA, a nie wyjątek na zewnątrz).
"""

from __future__ import annotations

import asyncio
import base64
import io
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from sufler.adapters.inbound.teams_graph.attachments import (
    AttachmentLimits,
    AttachmentMaterializer,
    FileBytesMaterializer,
)
from sufler.adapters.inbound.teams_graph.selection import AttachmentRef, ChannelMessage
from sufler.core.errors import AttachmentOutsideChannel

# Atrapy: prefiks magic + zera. Pillow ich nie otworzy → fallback na sniff magicznych bajtów
# (rozpoznaje typ, bez downscalingu) — dokładnie ścieżka dla obrazów inline z Teams.
_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16
_JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 16
_GIF = b"GIF89a" + b"\x00" * 16
_WEBP = b"RIFF" + b"\x00\x00\x00\x00" + b"WEBP" + b"\x00" * 16

_TEAM = "team-1"
_CHAN = "chan-1"


class _FakeGraphClient:
    """Atrapa portu ``GraphChannelClient`` — oddaje zaskryptowane bajty per referencja.

    ``hosted`` mapuje ``hosted_id → bajty``; ``files`` mapuje ``url → bajty``. Podanie
    ``Exception`` zamiast bajtów sprawia, że pobranie rzuca (test degradacji).
    """

    def __init__(
        self,
        *,
        hosted: dict[str, bytes | Exception] | None = None,
        files: dict[str, bytes | Exception] | None = None,
        public: dict[str, bytes | Exception] | None = None,
    ) -> None:
        self._hosted = hosted or {}
        self._files = files or {}
        self._public = public or {}
        self.hosted_calls: list[tuple[str, str, str, str, str]] = []
        self.file_calls: list[str] = []
        self.file_scopes: list[tuple[str, str]] = []
        self.public_calls: list[str] = []

    async def get_hosted_content(
        self, team_id: str, channel_id: str, root_id: str, message_id: str, hosted_id: str
    ) -> bytes:
        self.hosted_calls.append((team_id, channel_id, root_id, message_id, hosted_id))
        return _resolve(self._hosted[hosted_id])

    async def download_channel_file(self, team_id: str, channel_id: str, url: str) -> bytes:
        self.file_calls.append(url)
        self.file_scopes.append((team_id, channel_id))
        return _resolve(self._files[url])

    async def download_public_url(self, url: str) -> bytes:
        self.public_calls.append(url)
        return _resolve(self._public[url])


def _resolve(value: bytes | Exception) -> bytes:
    if isinstance(value, Exception):
        raise value
    return value


def _msg(
    refs: tuple[AttachmentRef, ...], *, msg_id: str = "m-1", root: str = "root-1"
) -> ChannelMessage:
    return ChannelMessage(
        id=msg_id,
        thread_root_id=root,
        created="2024-01-01T10:00:00Z",
        sender_id="u-anna",
        sender_name="Anna",
        text="opis",
        attachment_refs=refs,
    )


def _materialize(
    client: Any,
    refs: tuple[AttachmentRef, ...],
    *,
    max_bytes: int = 5_000_000,
    max_count: int = 5,
    max_total: int = 5_000_000,
    max_extract: int = 50_000_000,
    max_image_edge: int = 2048,
    max_total_text: int = 200_000,
    msg_id: str = "m-1",
    root: str = "root-1",
) -> tuple:
    materializer = AttachmentMaterializer(
        client,
        limits=AttachmentLimits(
            max_bytes=max_bytes,
            max_count=max_count,
            max_total_bytes=max_total,
            max_extract_bytes=max_extract,
            max_image_edge=max_image_edge,
            max_total_text_chars=max_total_text,
        ),
    )
    return asyncio.run(materializer.materialize(_TEAM, _CHAN, _msg(refs, msg_id=msg_id, root=root)))


def _docx_bytes(*, paragraph: str, table_cells: list[list[str]]) -> bytes:
    """Zbuduj mały .docx w pamięci: jeden akapit + tabela — do testu ekstrakcji tekstu."""
    doc = Document()
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=len(table_cells), cols=len(table_cells[0]))
    for r, row in enumerate(table_cells):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(*, sheet: str, rows: list[list[str]]) -> bytes:
    """Zbuduj mały .xlsx w pamięci: jeden arkusz z wierszami — do testu ekstrakcji tekstu."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(*, texts: list[str]) -> bytes:
    """Zbuduj mały .pptx w pamięci: jeden slajd z polami tekstowymi — do testu ekstrakcji."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # pusty układ
    for index, text in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(4), Inches(1))
        box.text_frame.text = text
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _image_bytes(*, size: tuple[int, int], fmt: str = "PNG") -> bytes:
    """Zbuduj PRAWDZIWY obraz danego rozmiaru/formatu (Pillow) — do testu sniffu i downscalingu."""
    buffer = io.BytesIO()
    Image.new("RGB", size, (123, 200, 50)).save(buffer, format=fmt)
    return buffer.getvalue()


# --- hosted (obrazy inline): sniff po zawartości -----------------------------


def test_hosted_png_sniffed_as_image_png():
    client = _FakeGraphClient(hosted={"h1": _PNG})

    (att,) = _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert att.kind == "image"
    assert att.media_type == "image/png"
    assert att.data_base64  # bajty zakodowane base64
    assert client.hosted_calls == [(_TEAM, _CHAN, "root-1", "m-1", "h1")]


def test_hosted_jpeg_sniffed_as_image_jpeg():
    client = _FakeGraphClient(hosted={"h1": _JPEG})

    (att,) = _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert (att.kind, att.media_type) == ("image", "image/jpeg")


def test_hosted_unrecognized_magic_bytes_yields_note():
    """Bajty bez znanego magic (nie obraz) → notka, nie wyjątek."""
    client = _FakeGraphClient(hosted={"h1": b"not-an-image-blob"})

    (att,) = _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert att.kind == "text"
    assert "nieobsługiwany typ" in att.text


def test_hosted_reply_forwards_root_and_reply_ids():
    """Wklejka w ODPOWIEDZI: materializer forwarduje thread_root_id (root) OBOK id repliki — to
    sedno fixu reply-scope 404 (klient buduje wtedy ścieżkę …/messages/{root}/replies/{reply})."""
    client = _FakeGraphClient(hosted={"h1": _PNG})

    _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert client.hosted_calls == [(_TEAM, _CHAN, "root-1", "m-1", "h1")]


def test_hosted_root_post_uses_own_id_as_root():
    """Wklejka w POŚCIE-ROOT: id == thread_root_id → root-scope zachowany (bez segmentu replies)."""
    client = _FakeGraphClient(hosted={"h1": _PNG})

    _materialize(
        client,
        (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),),
        msg_id="root-x",
        root="root-x",
    )

    assert client.hosted_calls == [(_TEAM, _CHAN, "root-x", "root-x", "h1")]


# --- HEIC/HEIF (zdjęcia iPhone): dekodowanie po zawartości + konwersja do JPEG ----


def _heic_bytes(*, size: tuple[int, int]) -> bytes:
    """Zbuduj PRAWDZIWY HEIC (pillow-heif) — do testu dekodowania i konwersji HEIC→JPEG."""
    pytest.importorskip("pillow_heif").register_heif_opener()
    buffer = io.BytesIO()
    Image.new("RGB", size, (123, 200, 50)).save(buffer, format="HEIF")
    return buffer.getvalue()


def test_hosted_heic_converted_to_jpeg():
    """HEIC wklejony inline → dekodowany po zawartości i konwertowany do JPEG."""
    client = _FakeGraphClient(hosted={"h1": _heic_bytes(size=(640, 480))})

    (att,) = _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert att.kind == "image"
    assert att.media_type == "image/jpeg"
    assert Image.open(io.BytesIO(base64.b64decode(att.data_base64))).format == "JPEG"


def test_file_heic_named_heic_converted_to_jpeg():
    """HEIC jako plik .heic → obraz/jpeg (routing po zawartości, nie rozszerzeniu)."""
    client = _FakeGraphClient(files={"u://foto": _heic_bytes(size=(640, 480))})
    ref = AttachmentRef(kind="file", name="foto.heic", url="u://foto")

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("image", "image/jpeg")


def test_large_heic_downscaled_to_jpeg():
    """Zdjęcie 12 MP HEIC (4032×3024) → downscale do 2048 px, wynik JPEG (aspect zachowany)."""
    client = _FakeGraphClient(hosted={"h1": _heic_bytes(size=(4032, 3024))})

    (att,) = _materialize(
        client,
        (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),),
        max_image_edge=2048,
    )

    out = Image.open(io.BytesIO(base64.b64decode(att.data_base64)))
    assert out.format == "JPEG"
    assert out.size == (2048, 1536)


def test_heic_output_fits_under_tight_max_bytes():
    """HEIC 12 MP jako JPEG mieści się pod ciasnym max_bytes (PNG by go przekroczył)."""
    client = _FakeGraphClient(hosted={"h1": _heic_bytes(size=(4032, 3024))})

    (att,) = _materialize(
        client,
        (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),),
        max_bytes=2_000_000,
        max_total=2_000_000,
        max_image_edge=2048,
    )

    assert att.kind == "image"
    assert len(base64.b64decode(att.data_base64)) < 2_000_000


def test_undecodable_heic_yields_note():
    """Bajty z marką HEIC bez dekodowalnej treści → notka (nigdy passthrough do API)."""
    client = _FakeGraphClient(hosted={"h1": b"\x00\x00\x00\x18ftypheic" + b"\x00" * 64})

    (att,) = _materialize(client, (AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),))

    assert att.kind == "text"
    assert "nieobsługiwany typ" in att.text


def test_ensure_heif_registered_missing_plugin_is_silent(monkeypatch):
    """Brak extra pillow-heif → helper NIE rzuca (poller nie może paść) i zapamiętuje False."""
    import builtins

    from sufler.adapters.inbound.teams_graph import attachments as att_mod

    real_import = builtins.__import__

    def fail_pillow_heif(name, *args, **kwargs):
        if name == "pillow_heif":
            raise ImportError("symulacja braku extra")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fail_pillow_heif)
    monkeypatch.setattr(att_mod, "_heif_state", None)

    att_mod._ensure_heif_registered()  # nie może rzucić

    assert att_mod._heif_state is False


def test_ensure_heif_registered_is_idempotent(monkeypatch):
    """Rejestracja opener'a HEIF wykonuje się co najwyżej RAZ na proces."""
    pillow_heif = pytest.importorskip("pillow_heif")

    from sufler.adapters.inbound.teams_graph import attachments as att_mod

    calls = {"n": 0}

    def counting_register() -> None:
        calls["n"] += 1

    monkeypatch.setattr(pillow_heif, "register_heif_opener", counting_register)
    monkeypatch.setattr(att_mod, "_heif_state", None)

    att_mod._ensure_heif_registered()
    att_mod._ensure_heif_registered()

    assert calls["n"] == 1
    assert att_mod._heif_state is True


# --- pliki: mapowanie po rozszerzeniu ----------------------------------------


def test_file_pdf_becomes_document():
    client = _FakeGraphClient(files={"u://umowa": b"%PDF-1.7 ..."})
    ref = AttachmentRef(kind="file", name="umowa.pdf", url="u://umowa")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "document"
    assert att.media_type == "application/pdf"
    assert att.name == "umowa.pdf"
    assert att.data_base64


def test_file_html_is_extracted_to_text_not_handed_over_as_markup():
    """HTML przechodzi ekstraktorem (ADR 0064), a nie gałęzią tekstową.

    Do 1.10.0 plik .html odbijał się notką „nieobsługiwany typ pliku" — użytkownik dostawał
    odmowę na format, który w tym pionie krąży najczęściej (zapisana strona, wyeksportowany
    raport). Gałąź tekstowa byłaby gorsza niż odmowa: model dostałby znaczniki i skrypty.
    """
    html = b"<html><body><script>var x=1;</script><p>Kwota: 12 300 zl</p></body></html>"
    client = _FakeGraphClient(files={"u://raport": html})
    ref = AttachmentRef(kind="file", name="raport.html", url="u://raport")

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("text", "text/plain")
    assert att.text == "Kwota: 12 300 zl"
    assert "<p>" not in att.text and "var x" not in att.text
    assert not att.data_base64  # tekst nie zjada budżetu bajtów API


def test_file_htm_alias_goes_through_the_same_dispatcher_as_html():
    """Alias ``.htm`` nie może zależeć od osobnej listy w drzwiach — to droga cichego rozjazdu."""
    client = _FakeGraphClient(files={"u://r": b"<p>tresc strony</p>"})
    ref = AttachmentRef(kind="file", name="raport.htm", url="u://r")

    (att,) = _materialize(client, (ref,))

    assert att.text == "tresc strony"


def test_readable_file_without_text_yields_a_note_instead_of_an_empty_label():
    """Pusty wynik ekstrakcji ma być NAZWANY — inaczej model dostaje samą etykietę pliku.

    Uwaga: strona z samą grafiką NIE jest tym przypadkiem — ekstraktor policzy obrazy bez opisu
    i to JEST treść (o tym mówi anty-maskowanie). Chodzi o plik faktycznie bez czego czytać.
    """
    client = _FakeGraphClient(files={"u://r": b"<html><body><div></div></body></html>"})
    ref = AttachmentRef(kind="file", name="pusta.html", url="u://r")

    (att,) = _materialize(client, (ref,))

    assert att.name == "status załącznika"
    assert "nie zawiera tekstu" in att.text


def test_file_image_media_type_from_content_overrides_extension():
    """Typ obrazu bierzemy z ZAWARTOŚCI, nie z rozszerzenia — plik JPEG nazwany .png → jpeg."""
    client = _FakeGraphClient(files={"u://foto": _JPEG})
    ref = AttachmentRef(kind="file", name="foto.png", url="u://foto")

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("image", "image/jpeg")


def test_file_gif_and_webp_become_image():
    """REGRESJA symetrii: .gif/.webp jako PLIK też są obrazem (dawniej degradowały)."""
    client = _FakeGraphClient(files={"u://g": _GIF, "u://w": _WEBP})
    refs = (
        AttachmentRef(kind="file", name="anim.gif", url="u://g"),
        AttachmentRef(kind="file", name="logo.webp", url="u://w"),
    )

    gif, webp = _materialize(client, refs)

    assert (gif.kind, gif.media_type) == ("image", "image/gif")
    assert (webp.kind, webp.media_type) == ("image", "image/webp")


def test_file_garbage_named_image_yields_note():
    """Plik nazwany .png, ale treść to nie obraz → notka (nie wysyłamy fałszywego image/png)."""
    client = _FakeGraphClient(files={"u://z": b"whatever-bytes-not-image"})
    ref = AttachmentRef(kind="file", name="foto.png", url="u://z")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nieobsługiwany typ" in att.text


# --- obrazy: downscaling i przepuszczenie ------------------------------------


def test_large_image_is_downscaled_to_max_edge():
    """Duży obraz → przeskalowany tak, że dłuższa krawędź ≤ próg (koszt tokenów)."""
    data = _image_bytes(size=(4000, 3000), fmt="PNG")
    client = _FakeGraphClient(files={"u://big": data})
    ref = AttachmentRef(kind="file", name="wielki.png", url="u://big")

    (att,) = _materialize(client, (ref,), max_image_edge=1024)

    assert (att.kind, att.media_type) == ("image", "image/png")
    out = Image.open(io.BytesIO(base64.b64decode(att.data_base64)))
    assert max(out.size) <= 1024
    assert out.size == (1024, 768)  # proporcje zachowane


def test_small_image_passes_through_unchanged():
    """Obraz poniżej progu → bajty przepuszczone bez re-enkodowania (oryginał zachowany)."""
    data = _image_bytes(size=(320, 240), fmt="PNG")
    client = _FakeGraphClient(files={"u://small": data})
    ref = AttachmentRef(kind="file", name="male.png", url="u://small")

    (att,) = _materialize(client, (ref,), max_image_edge=2048)

    assert att.media_type == "image/png"
    assert base64.b64decode(att.data_base64) == data


def test_url_ref_fetched_via_public_download_and_sniffed():
    """Referencja ``url`` (GIF/emoji) pobierana przez ``download_public_url`` → obraz po sniffie."""
    data = _image_bytes(size=(64, 64), fmt="PNG")
    url = "https://media.giphy.com/media/abc/giphy.gif"
    client = _FakeGraphClient(public={url: data})
    ref = AttachmentRef(kind="url", name="obraz", url=url)

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("image", "image/png")  # typ z ZAWARTOŚCI (sniff)
    assert client.public_calls == [url]  # poszło publiczną ścieżką (bez tokenu Graph)


def test_oversized_pixel_image_not_decoded_locally(monkeypatch):
    """Obraz ponad sufit PIKSELI nie jest dekodowany lokalnie (ochrona pamięci)."""
    from sufler.adapters.inbound.teams_graph import attachments as attachments_mod

    monkeypatch.setattr(attachments_mod, "_MAX_IMAGE_PIXELS", 1000)  # 320×240 = 76800 > 1000
    data = _image_bytes(size=(320, 240), fmt="PNG")
    client = _FakeGraphClient(files={"u://big": data})
    ref = AttachmentRef(kind="file", name="wielki.png", url="u://big")

    # max_image_edge=100 zwykle by przeskalowało; sufit pikseli krótkozwiera PRZED downscalingiem.
    (att,) = _materialize(client, (ref,), max_image_edge=100)

    assert att.media_type == "image/png"
    assert base64.b64decode(att.data_base64) == data  # oryginał, bez lokalnego dekodowania


def test_file_docx_extracts_paragraphs_and_table_text():
    """.docx → kind=text z tekstem: akapit ORAZ komórki tabeli (Claude nie przyjmuje docx)."""
    data = _docx_bytes(
        paragraph="Ustalenia ze spotkania",
        table_cells=[["Zadanie", "Termin"], ["Wdrożenie SCADA", "2024-03"]],
    )
    client = _FakeGraphClient(files={"u://doc": data})
    ref = AttachmentRef(kind="file", name="notatka.docx", url="u://doc")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert att.media_type == "text/plain"
    assert att.data_base64 == ""  # docx idzie jako tekst, nie base64
    assert "Ustalenia ze spotkania" in att.text
    # Sedno: tekst TABELI też jest wyekstrahowany (nie tylko akapity).
    assert "Wdrożenie SCADA" in att.text
    assert "Termin" in att.text


def test_file_xlsx_extracts_cell_text_with_sheet_name():
    """.xlsx → kind=text z tekstem komórek i nazwą arkusza (Claude nie przyjmuje Excela)."""
    data = _xlsx_bytes(
        sheet="Budżet",
        rows=[["Pozycja", "Kwota"], ["Licencje SCADA", "12000"]],
    )
    client = _FakeGraphClient(files={"u://x": data})
    ref = AttachmentRef(kind="file", name="budzet.xlsx", url="u://x")

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("text", "text/plain")
    assert att.data_base64 == ""  # tekst, nie base64
    assert "Budżet" in att.text  # nazwa arkusza
    assert "Licencje SCADA" in att.text
    assert "12000" in att.text


def test_file_pptx_extracts_slide_text():
    """.pptx → kind=text z tekstem slajdów (Claude nie przyjmuje PowerPointa)."""
    data = _pptx_bytes(texts=["Integracja MPWiK", "Kamień milowy: marzec"])
    client = _FakeGraphClient(files={"u://p": data})
    ref = AttachmentRef(kind="file", name="prezentacja.pptx", url="u://p")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "Integracja MPWiK" in att.text
    assert "Kamień milowy: marzec" in att.text


def test_file_text_formats_decoded_as_text():
    """.txt/.csv/.md → kind=text z zdekodowaną treścią UTF-8 (z polskimi znakami)."""
    data = "firma;projekt\nmpwik;scada-integration\nżółć\n".encode()
    client = _FakeGraphClient(files={"u://c": data})
    ref = AttachmentRef(kind="file", name="dane.csv", url="u://c")

    (att,) = _materialize(client, (ref,))

    assert (att.kind, att.media_type) == ("text", "text/plain")
    assert "scada-integration" in att.text
    assert "żółć" in att.text  # UTF-8 zachowane


def test_tekst_z_ekstrakcji_ma_wlasny_laczny_sufit():
    """Regresja: ekstrakcja zaliczała do budżetu ZERO, więc obie bramki bajtów ją przepuszczały.

    Jedynym ogranicznikiem tej ścieżki zostawał ``max_count`` (domyślnie 20) razy 200 000 znaków
    na plik — do czterech milionów znaków w JEDNEJ turze użytkownika, czyli grubo ponad okno
    kontekstu modelu. Żądanie odbijało się na API już po opłaceniu wszystkich ekstrakcji.
    """
    duzy = ("x" * 900).encode()
    client = _FakeGraphClient(
        files={"u://1": duzy, "u://2": duzy, "u://3": duzy},
    )
    refs = tuple(AttachmentRef(kind="file", name=f"d{i}.txt", url=f"u://{i}") for i in (1, 2, 3))

    wyniki = _materialize(client, refs, max_total_text=2000)

    # Notka o statusie też jest załącznikiem tekstowym, więc rozróżniamy po NAZWIE pliku.
    tresci = [a for a in wyniki if a.name.endswith(".txt")]
    assert len(tresci) == 2  # dwa mieszczą się w 2000 znaków, trzeci już nie
    assert any("limit tekstu" in a.text for a in wyniki if a.name == "status załącznika")


def test_sufit_tekstu_nie_rusza_plikow_ktore_sie_miescza():
    """Kontrast: sufit ma przycinać nadmiar, nie blokować zwykłej wiadomości z załącznikiem."""
    client = _FakeGraphClient(files={"u://1": b"krotka tresc"})
    ref = AttachmentRef(kind="file", name="d1.txt", url="u://1")

    (att,) = _materialize(client, (ref,), max_total_text=2000)

    assert att.kind == "text"
    assert "krotka tresc" in att.text


def test_corrupt_xlsx_yields_note_not_raised():
    """Uszkodzony .xlsx (ekstrakcja rzuca) degraduje do notki — nie kładzie pollera."""
    client = _FakeGraphClient(files={"u://bad": b"not a valid xlsx package"})
    ref = AttachmentRef(kind="file", name="uszkodzony.xlsx", url="u://bad")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nie udało się przetworzyć" in att.text
    assert "uszkodzony.xlsx" in att.text


def test_corrupt_pptx_yields_note_not_raised():
    """Uszkodzony .pptx (ekstrakcja rzuca) degraduje do notki — nie kładzie pollera."""
    client = _FakeGraphClient(files={"u://bad": b"not a valid pptx package"})
    ref = AttachmentRef(kind="file", name="uszkodzony.pptx", url="u://bad")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nie udało się przetworzyć" in att.text
    assert "uszkodzony.pptx" in att.text


def test_file_unsupported_extension_yields_note():
    client = _FakeGraphClient(files={"u://z": b"zip-bytes"})
    ref = AttachmentRef(kind="file", name="archiwum.zip", url="u://z")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nieobsługiwany typ" in att.text
    assert "archiwum.zip" in att.text


# --- łagodna degradacja: limity i błędy --------------------------------------


def test_oversized_attachment_yields_note_not_bytes():
    """Plik ponad limit rozmiaru → notka (nie wysyłamy megabajtów do API)."""
    client = _FakeGraphClient(files={"u://big": b"x" * 100})
    ref = AttachmentRef(kind="file", name="wielki.pdf", url="u://big")

    (att,) = _materialize(client, (ref,), max_bytes=10)

    assert att.kind == "text"
    assert "przekracza limit rozmiaru" in att.text
    assert "wielki.pdf" in att.text


def test_over_count_drops_extras_with_note_and_still_fetches_kept():
    """Ponad ``max_count`` → notka o pominiętych + POBRANE tylko dozwolone (reszta odcięta)."""
    client = _FakeGraphClient(files={"u://1": b"%PDF a", "u://2": b"%PDF b"})
    refs = (
        AttachmentRef(kind="file", name="a.pdf", url="u://1"),
        AttachmentRef(kind="file", name="b.pdf", url="u://2"),
        AttachmentRef(kind="file", name="c.pdf", url="u://3"),  # ta już nie pobrana
    )

    result = _materialize(client, refs, max_count=2)

    # Pierwszy element to notka o pominięciu, potem 2 pobrane dokumenty.
    assert result[0].kind == "text"
    assert "Pominięto 1" in result[0].text
    assert [a.kind for a in result[1:]] == ["document", "document"]
    # Trzeci URL nie był w ogóle pobierany (limit odciął przed I/O).
    assert client.file_calls == ["u://1", "u://2"]


def test_download_exception_yields_note_not_raised():
    """Wyjątek pobrania jednego załącznika NIE kładzie pollera — degraduje do notki."""
    client = _FakeGraphClient(files={"u://boom": RuntimeError("Graph 500")})
    ref = AttachmentRef(kind="file", name="feler.pdf", url="u://boom")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nie udało się pobrać" in att.text
    assert "feler.pdf" in att.text


def test_hosted_download_404_hint_to_send_as_file():
    """Wklejony obraz (hosted/AMS) z błędem pobrania → notka z radą wysłania jako plik."""
    client = _FakeGraphClient(hosted={"h1": RuntimeError("404 Not Found")})
    ref = AttachmentRef(kind="hosted", name="obraz", hosted_id="h1")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "jako osobny plik" in att.text  # rada dla użytkownika


def test_large_extracted_file_ignores_base64_limit():
    """Plik ekstrahowany do tekstu (docx) NIE podlega limitowi base64 — liczy się max_extract."""
    data = _docx_bytes(paragraph="Duża notatka", table_cells=[["a", "b"]])
    client = _FakeGraphClient(files={"u://doc": data})
    ref = AttachmentRef(kind="file", name="duza.docx", url="u://doc")

    # max_bytes (base64) drastycznie mały — gdyby docx mu podlegał, odpadłby na rozmiarze.
    (att,) = _materialize(client, (ref,), max_bytes=10, max_extract=1_000_000)

    assert att.kind == "text"
    assert "Duża notatka" in att.text  # wyekstrahowany mimo maleńkiego max_bytes


def test_extracted_file_does_not_consume_api_budget():
    """Ekstrakcja do tekstu zalicza 0 do budżetu base64 — PDF po niej wciąż się mieści."""
    docx = _docx_bytes(paragraph="Notatka", table_cells=[["x", "y"]])
    client = _FakeGraphClient(files={"u://d": docx, "u://p": b"%PDF" + b"a" * 96})
    refs = (
        AttachmentRef(kind="file", name="a.docx", url="u://d"),  # tekst → 0 do budżetu
        AttachmentRef(kind="file", name="b.pdf", url="u://p"),  # 100 B base64
    )

    doc_text, pdf = _materialize(client, refs, max_bytes=1000, max_total=150)

    assert doc_text.kind == "text"  # docx wyekstrahowany
    assert pdf.kind == "document"  # PDF się zmieścił (docx nie zjadł budżetu)


def test_corrupt_docx_yields_note_not_raised():
    """REGRESJA: uszkodzony/podszyty .docx (ekstrakcja rzuca) degraduje do notki —

    inaczej wyjątek propagowałby przez pollera i zapętlał kanał (build był poza try).
    """
    client = _FakeGraphClient(files={"u://bad": b"not a valid zip/docx package"})
    ref = AttachmentRef(kind="file", name="uszkodzony.docx", url="u://bad")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "nie udało się przetworzyć" in att.text
    assert "uszkodzony.docx" in att.text


def test_aggregate_budget_drops_overflow_with_note():
    """Pliki mieszczące się pojedynczo, ale przekraczające ŁĄCZNY budżet → nadmiar to notka."""
    client = _FakeGraphClient(files={"u://1": b"%PDF" + b"a" * 96, "u://2": b"%PDF" + b"b" * 96})
    refs = (
        AttachmentRef(kind="file", name="a.pdf", url="u://1"),  # 100 B — mieści się
        AttachmentRef(kind="file", name="b.pdf", url="u://2"),  # 100 B — już nie (budżet 150)
    )

    first, second = _materialize(client, refs, max_bytes=1000, max_total=150)

    assert first.kind == "document"
    assert second.kind == "text"
    assert "przekroczony łączny limit" in second.text
    assert "b.pdf" in second.text


def test_mixed_refs_materialize_independently():
    """Jedna zła referencja (błąd) nie psuje dobrych — każda materializowana osobno."""
    client = _FakeGraphClient(
        hosted={"h1": _PNG},
        files={"u://good": b"%PDF ok", "u://bad": RuntimeError("padło")},
    )
    refs = (
        AttachmentRef(kind="hosted", name="obraz", hosted_id="h1"),
        AttachmentRef(kind="file", name="dobry.pdf", url="u://good"),
        AttachmentRef(kind="file", name="zly.pdf", url="u://bad"),
    )

    img, good, bad = _materialize(client, refs)

    assert img.kind == "image"
    assert good.kind == "document"
    assert bad.kind == "text" and "nie udało się pobrać" in bad.text


def test_file_materializer_degrades_a_broken_document_instead_of_raising():
    """Uszkodzony plik NIE MOŻE wyjść wyjątkiem — rdzeń woła narzędzie poza ``try``.

    ``DocumentExtractionError`` nie dziedziczy z ``SuflerError``, więc bez osłony TUTAJ
    przelatywał kopertę narzędzia i zabijał całą turę: użytkownik dostawał „chwilowy błąd",
    a tura nie trafiała do pamięci rozmowy. ADR 0064 obiecuje degradację do notki, nigdy crash.
    """
    materializer = FileBytesMaterializer(max_image_edge=2048)

    assert materializer.materialize("umowa.docx", b"to nie jest zip") is None


def test_file_materializer_builds_the_same_attachment_as_the_door():
    """Model i drzwi mają widzieć ten sam plik tak samo — stąd wspólny ``_build``, nie kopia."""
    materializer = FileBytesMaterializer(max_image_edge=2048)

    built = materializer.materialize("raport.html", b"<p>tresc strony</p>")

    assert built is not None
    attachment, sent = built
    assert (attachment.kind, attachment.text, sent) == ("text", "tresc strony", 0)


def test_laczny_sufit_tekstu_jest_ZWIAZANY_z_sufitem_pojedynczego_pliku():
    """Komentarz przy ``max_total_text_chars`` obiecuje „tyle, ile wolno pojedynczemu plikowi",
    ale liczba jest przepisana literałem w drugim module. Podniesienie sufitu per-plik bez tego
    sprawiłoby, że JEDEN duży dokument przekracza łączny budżet wiadomości i degraduje do notki —
    cicha utrata zdolności, nie błąd. Wiążemy obie liczby sondą, bo import stałej z modułu
    ekstrakcji zrobiłby z sufitu wiadomości pochodną cudzej decyzji.
    """
    from sufler.adapters.inbound.document_text import _MAX_TEXT_CHARS

    limity = AttachmentLimits(max_bytes=1, max_count=1, max_total_bytes=1)

    assert limity.max_total_text_chars == _MAX_TEXT_CHARS


def test_zalacznik_spoza_kanalu_dostaje_INNA_notke_niz_nieudane_pobranie():
    """Odmowa granicy i awaria pobrania to dwa różne zdarzenia — i mają być odróżnialne (ADR 0072).

    Gdyby odmowa wpadła do generycznego handlera, człowiek dostałby „nie udało się pobrać" —
    komunikat, który każe spróbować ponownie, choć ponowienie nigdy nie zadziała. Notka ma
    powiedzieć, CO zrobić: wgrać plik do kanału.
    """
    client = _FakeGraphClient(files={"u://obcy": AttachmentOutsideChannel("poza kanałem")})
    ref = AttachmentRef(kind="file", name="wynagrodzenia.xlsx", url="u://obcy")

    (att,) = _materialize(client, (ref,))

    assert att.kind == "text"
    assert "wynagrodzenia.xlsx" in att.text
    assert "nie leży w plikach tego kanału" in att.text
    assert "nie udało się pobrać" not in att.text


def test_materializer_podaje_kanal_przy_pobraniu_zalacznika():
    """Kanał jedzie do klienta jako argument — bez tego granica nie miałaby wobec czego mierzyć."""
    client = _FakeGraphClient(files={"u://a": b"%PDF-1.4 tresc"})
    ref = AttachmentRef(kind="file", name="raport.pdf", url="u://a")

    _materialize(client, (ref,))

    assert client.file_scopes == [(_TEAM, _CHAN)]
