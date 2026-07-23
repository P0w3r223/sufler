"""Testy adapterów wyjściowych kart czasu (ADR 0035): xlsx, źródło JSON, katalog tożsamości.

``tmp_path`` dla realnych plików, ``httpx.MockTransport`` dla Graph — jak reszta repo, bez
bibliotek mockujących. Sedno: adaptery mają być głupie i fail-closed.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import httpx
import pytest

from workmate.adapters.outbound.graph_identity_directory import (
    GraphIdentityDirectory,
    YamlIdentityDirectory,
    fetch_team_members,
)
from workmate.adapters.outbound.graph_shift_source import GraphShiftSource
from workmate.adapters.outbound.json_hours_source import JsonHoursSource
from workmate.adapters.outbound.openpyxl_sheet_writer import OpenpyxlSheetWriter
from workmate.core.domain.timesheet import TimesheetError

_HEADERS = ("Issue Key/ID", "User", "Time Spent", "Start Date & Time", "Comment")
_ROW = ("WT-12", "mikolaj@example.org", "1h 30m", "2026-07-15T08:00:00.000+0200", "")


# --- zapis xlsx -------------------------------------------------------------------


def _read_back(path: Path) -> list[tuple]:
    from openpyxl import load_workbook

    sheet = load_workbook(path).worksheets[0]
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_writes_headers_and_rows(tmp_path: Path) -> None:
    target = tmp_path / "a.xlsx"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, (_ROW,))
    header_row, data_row = _read_back(target)
    assert header_row == _HEADERS
    # Pusty ``Comment`` wraca jako ``None`` — Excel nie odróżnia pustego tekstu od pustej
    # komórki. Dla WorklogPRO to bez różnicy (kolumna jest opcjonalna), więc porównujemy
    # wartości znaczące, a nie dosłowną krotkę.
    assert data_row[:4] == _ROW[:4]
    assert data_row[4] in ("", None)


def test_creates_missing_directories(tmp_path: Path) -> None:
    target = tmp_path / "2026-W29" / "mikolaj.xlsx"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, (_ROW,))
    assert target.exists()


def test_overwrites_existing_file(tmp_path: Path) -> None:
    """Ścieżka jest deterministyczna, więc ponowny przebieg ma NADPISAĆ, nie dołożyć drugi plik."""
    target = tmp_path / "a.xlsx"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, (_ROW, _ROW))
    OpenpyxlSheetWriter().write(str(target), _HEADERS, (_ROW,))
    assert len(_read_back(target)) == 2  # nagłówek + jeden wiersz


def test_timestamp_stays_text_not_a_date_cell(tmp_path: Path) -> None:
    """Excel przekształciłby znacznik we własny typ daty i zepsuł to, co parsuje WorklogPRO."""
    target = tmp_path / "a.xlsx"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, (_ROW,))
    assert _read_back(target)[1][3] == "2026-07-15T08:00:00.000+0200"


def test_formula_injection_lands_as_text_not_a_formula(tmp_path: Path) -> None:
    """Komentarz ze źródła nie może stać się FORMUŁĄ — to wykonanie kodu u człowieka.

    openpyxl wnioskuje typ z treści: napis od ``=`` zapisuje jako formułę. Plik jawnie każemy
    otworzyć przed importem, a źródła godzin jeszcze nie znamy, więc ``=cmd|'/c calc'!A0``
    w komórce jest realnym wektorem (klasyczne wstrzyknięcie formuły do arkusza). Kontrakt
    ``SheetWriter``: KAŻDA komórka zapisana jako tekst — sprawdzamy przez odczyt typu komórki,
    nie samej wartości, bo wartość wygląda tak samo w obu przypadkach.
    """
    from openpyxl import load_workbook

    target = tmp_path / "a.xlsx"
    attack = "=cmd|'/c calc'!A0"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, ((*_ROW[:4], attack),))

    cell = load_workbook(target).worksheets[0].cell(row=2, column=5)
    assert cell.data_type == "s"  # 'f' = formuła
    assert cell.value == attack  # treść NIETKNIĘTA — komentarz ma dojechać do Jiry taki, jaki był


def test_writes_an_empty_sheet_without_rows(tmp_path: Path) -> None:
    target = tmp_path / "a.xlsx"
    OpenpyxlSheetWriter().write(str(target), _HEADERS, ())
    assert _read_back(target) == [_HEADERS]


# --- źródło JSON ------------------------------------------------------------------


def _source(tmp_path: Path, payload: str) -> JsonHoursSource:
    path = tmp_path / "hours.json"
    path.write_text(payload, encoding="utf-8")
    return JsonHoursSource(path)


def test_reads_entries_and_converts_hours(tmp_path: Path) -> None:
    source = _source(
        tmp_path,
        '[{"source_id":"EMP-042","day":"2026-07-15","issue_key":"wt-12","hours":3.5}]',
    )
    (entry,) = source.read(date(2026, 7, 13), date(2026, 7, 20))
    assert entry.minutes == 210
    assert entry.issue_key == "WT-12"  # znormalizowane do wielkich liter


def test_minutes_take_precedence_over_hours(tmp_path: Path) -> None:
    source = _source(
        tmp_path,
        '[{"source_id":"E","day":"2026-07-15","issue_key":"WT-1","hours":9,"minutes":90}]',
    )
    assert source.read(date(2026, 7, 13), date(2026, 7, 20))[0].minutes == 90


def test_filters_to_the_half_open_window(tmp_path: Path) -> None:
    source = _source(
        tmp_path,
        '[{"source_id":"E","day":"2026-07-12","issue_key":"WT-1","minutes":60},'
        '{"source_id":"E","day":"2026-07-15","issue_key":"WT-1","minutes":60},'
        '{"source_id":"E","day":"2026-07-20","issue_key":"WT-1","minutes":60}]',
    )
    entries = source.read(date(2026, 7, 13), date(2026, 7, 20))
    assert [e.day.day for e in entries] == [15]


def test_missing_file_is_a_hard_error_not_silent_emptiness(tmp_path: Path) -> None:
    """Cicha pustka wyglądałaby jak „nikt nie pracował" — tydzień bez wiadomości i bez śladu."""
    with pytest.raises(TimesheetError, match="brak pliku"):
        JsonHoursSource(tmp_path / "nie-ma.json").read(date(2026, 7, 13), date(2026, 7, 20))


def test_corrupt_json_is_a_hard_error(tmp_path: Path) -> None:
    with pytest.raises(TimesheetError, match="uszkodzony"):
        _source(tmp_path, "{nie json").read(date(2026, 7, 13), date(2026, 7, 20))


def test_missing_required_field_is_a_hard_error(tmp_path: Path) -> None:
    with pytest.raises(TimesheetError, match="wymaganego pola"):
        _source(tmp_path, '[{"source_id":"E","day":"2026-07-15"}]').read(
            date(2026, 7, 13), date(2026, 7, 20)
        )


def test_non_list_payload_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(TimesheetError, match="LISTĘ"):
        _source(tmp_path, '{"a":1}').read(date(2026, 7, 13), date(2026, 7, 20))


# --- katalog tożsamości -----------------------------------------------------------

_YAML = """
EMP-042:
  aad_user_id: aad-mikolaj
  jira_user: mikolaj@example.org
  git_email: mikolaj@example.org
  display_name: Mikołaj Anonimowicz
EMP-017:
  aad_user_id: aad-piotr
  jira_user: piotr@example.com
"""


def _identities(tmp_path: Path, payload: str = _YAML) -> Path:
    path = tmp_path / "identities.yaml"
    path.write_text(payload, encoding="utf-8")
    return path


def test_yaml_directory_resolves_a_known_person(tmp_path: Path) -> None:
    person = YamlIdentityDirectory(_identities(tmp_path)).resolve("EMP-042")
    assert person is not None
    assert person.aad_user_id == "aad-mikolaj"
    assert person.jira_user == "mikolaj@example.org"


def test_yaml_directory_is_fail_closed_for_unknown_person(tmp_path: Path) -> None:
    """Brak wpisu = brak pliku i brak wiadomości. Nigdy dopasowania po nazwisku."""
    assert YamlIdentityDirectory(_identities(tmp_path)).resolve("EMP-999") is None


def test_missing_identity_file_fails_at_startup(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="brak pliku mapy"):
        YamlIdentityDirectory(tmp_path / "nie-ma.yaml")


def test_entry_without_jira_user_fails_at_startup(tmp_path: Path) -> None:
    """Niekompletna mapa = ktoś po cichu nie dostanie nic. Padamy przy starcie, nie w piątek."""
    with pytest.raises(ValueError, match="jira_user"):
        YamlIdentityDirectory(_identities(tmp_path, "EMP-1:\n  aad_user_id: a\n"))


def test_two_people_sharing_a_jira_account_fail_at_startup(tmp_path: Path) -> None:
    """Skopiowany blok bez podmiany ``jira_user`` = czyjś tydzień na cudzym koncie Jiry.

    Fail-closed pilnował tylko osi „brak wpisu". Tę awarię strażniki ``assert_single_person``
    przepuszczają, bo zestawienie JEST jednorodne — wskazuje po prostu złą osobę. A worklogi
    są create-only i nieusuwalne z poziomu narzędzi.
    """
    duplikat = (
        "EMP-1:\n  aad_user_id: aad-1\n  jira_user: mikolaj@example.com\n"
        "EMP-2:\n  aad_user_id: aad-2\n  jira_user: mikolaj@example.com\n"
    )
    with pytest.raises(ValueError, match="jira_user"):
        YamlIdentityDirectory(_identities(tmp_path, duplikat))


def test_two_people_sharing_a_teams_account_fail_at_startup(tmp_path: Path) -> None:
    """Ten sam ``aad_user_id`` u dwóch osób = ktoś dostaje cudzą tabelę godzin."""
    duplikat = (
        "EMP-1:\n  aad_user_id: aad-mikolaj\n  jira_user: a@example.com\n"
        "EMP-2:\n  aad_user_id: aad-mikolaj\n  jira_user: b@example.com\n"
    )
    with pytest.raises(ValueError, match="aad_user_id"):
        YamlIdentityDirectory(_identities(tmp_path, duplikat))


# --- most tożsamości: git_email + odwrotne lookupy (ADR 0036) ---------------------


def test_yaml_parses_optional_git_email(tmp_path: Path) -> None:
    directory = YamlIdentityDirectory(_identities(tmp_path))
    with_email = directory.resolve("EMP-042")
    without_email = directory.resolve("EMP-017")
    assert with_email is not None and with_email.git_email == "mikolaj@example.org"
    assert without_email is not None and without_email.git_email == ""  # opcjonalny


def test_resolve_by_git_email_maps_back_to_person(tmp_path: Path) -> None:
    directory = YamlIdentityDirectory(_identities(tmp_path))
    person = directory.resolve_by_git_email("mikolaj@example.org")
    assert person is not None and person.source_id == "EMP-042"


def test_resolve_by_git_email_is_fail_closed(tmp_path: Path) -> None:
    """Nieznany e-mail (commit obcej osoby) i pusty e-mail → None, nigdy zgadywanie."""
    directory = YamlIdentityDirectory(_identities(tmp_path))
    assert directory.resolve_by_git_email("ktos-obcy@example.com") is None
    assert directory.resolve_by_git_email("") is None


def test_resolve_by_aad_user_id_maps_back_to_person(tmp_path: Path) -> None:
    """Wejście od Shifts: aad_user_id → osoba (albo None dla obcego)."""
    directory = YamlIdentityDirectory(_identities(tmp_path))
    person = directory.resolve_by_aad_user_id("aad-mikolaj")
    assert person is not None and person.source_id == "EMP-042"
    assert directory.resolve_by_aad_user_id("aad-obcy") is None


def test_graph_reverse_lookups_honor_team_membership(tmp_path: Path) -> None:
    """Odwrotne lookupy przez Graph też są bramkowane członkostwem — osoba spoza zespołu → None."""
    directory = GraphIdentityDirectory(_identities(tmp_path), {"aad-mikolaj": "Mikołaj A."})
    assert directory.resolve_by_git_email("mikolaj@example.org") is not None
    assert directory.resolve_by_aad_user_id("aad-mikolaj") is not None
    # EMP-017 (aad-piotr) nie jest w zespole → oba lookupy fail-closed.
    assert directory.resolve_by_aad_user_id("aad-piotr") is None


def test_two_people_sharing_a_git_email_fail_at_startup(tmp_path: Path) -> None:
    """Współdzielony git_email przypisałby czyjeś commity drugiej osobie (ADR 0036)."""
    duplikat = (
        "EMP-1:\n  aad_user_id: aad-1\n  jira_user: a@x.pl\n  git_email: dev@x.pl\n"
        "EMP-2:\n  aad_user_id: aad-2\n  jira_user: b@x.pl\n  git_email: dev@x.pl\n"
    )
    with pytest.raises(ValueError, match="git_email"):
        YamlIdentityDirectory(_identities(tmp_path, duplikat))


def test_empty_git_email_is_not_a_collision(tmp_path: Path) -> None:
    """Kilka osób bez git_email to norma (opcjonalny), nie kolizja — mapa ma się załadować."""
    bez = (
        "EMP-1:\n  aad_user_id: aad-1\n  jira_user: a@x.pl\n"
        "EMP-2:\n  aad_user_id: aad-2\n  jira_user: b@x.pl\n"
    )
    directory = YamlIdentityDirectory(_identities(tmp_path, bez))
    assert directory.resolve("EMP-1") is not None
    assert directory.resolve_by_git_email("") is None


def test_graph_directory_requires_current_team_membership(tmp_path: Path) -> None:
    directory = GraphIdentityDirectory(_identities(tmp_path), {"aad-mikolaj": "Mikołaj A."})
    assert directory.resolve("EMP-042") is not None
    assert directory.resolve("EMP-017") is None  # nie jest już w zespole


def test_graph_directory_fills_display_name_from_graph(tmp_path: Path) -> None:
    directory = GraphIdentityDirectory(_identities(tmp_path), {"aad-piotr": "Piotr Cząstkiewicz"})
    person = directory.resolve("EMP-017")
    assert person is not None and person.display_name == "Piotr Cząstkiewicz"


def test_explicit_display_name_wins_over_graph(tmp_path: Path) -> None:
    directory = GraphIdentityDirectory(_identities(tmp_path), {"aad-mikolaj": "z Graph"})
    person = directory.resolve("EMP-042")
    assert person is not None and person.display_name == "Mikołaj Anonimowicz"


# --- pobranie członków z Graph ----------------------------------------------------


def test_fetch_team_members_maps_ids_to_names() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["Authorization"] == "Bearer tok"
        return httpx.Response(
            200,
            json={
                "value": [
                    {"userId": "aad-1", "displayName": "Ala"},
                    {"userId": "", "displayName": "bez id"},
                    "śmieć",
                ]
            },
        )

    client = httpx.Client(transport=httpx.MockTransport(handler))
    assert fetch_team_members(client, "team-1", "tok") == {"aad-1": "Ala"}


def test_fetch_team_members_follows_pagination() -> None:
    pages = {
        "/v1.0/teams/team-1/members": {
            "value": [{"userId": "aad-1", "displayName": "Ala"}],
            "@odata.nextLink": "https://graph.microsoft.com/v1.0/next",
        },
        "/v1.0/next": {"value": [{"userId": "aad-2", "displayName": "Bob"}]},
    }

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=pages[request.url.path])

    client = httpx.Client(transport=httpx.MockTransport(handler))
    assert set(fetch_team_members(client, "team-1", "tok")) == {"aad-1", "aad-2"}


def test_fetch_team_members_raises_on_permission_error() -> None:
    """403 = brak TeamMember.Read.All. Musi być głośne, nie pusta lista udająca pusty zespół."""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(403, json={"error": {"message": "Forbidden"}})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    with pytest.raises(httpx.HTTPStatusError):
        fetch_team_members(client, "team-1", "tok")


def test_fetch_team_members_fails_loudly_when_the_page_cap_cuts_the_list() -> None:
    """Ucięcie listy członków było CICHE — a brak członka znaczy „nie dostaje arkusza".

    Katalog tożsamości jest fail-closed, więc osoba spoza tej mapy wygląda w raporcie
    identycznie jak ktoś, kto nie pracował. Przy zespole kilkunastu osób wyczerpanie
    dziesięciu stron jest anomalią, więc lepiej nie wysłać NIC niż pominąć kogoś po cichu
    (ta sama zasada co przy niekompletnej mapie tożsamości).
    """

    def handler(request: httpx.Request) -> httpx.Response:
        # Każda strona wskazuje kolejną — Graph nigdy nie mówi „to już koniec".
        return httpx.Response(
            200,
            json={
                "value": [{"userId": f"aad-{request.url.path}", "displayName": "X"}],
                "@odata.nextLink": "https://graph.microsoft.com/v1.0/next",
            },
        )

    client = httpx.Client(transport=httpx.MockTransport(handler))
    with pytest.raises(ValueError, match="NIEKOMPLETNA"):
        fetch_team_members(client, "team-1", "tok")


def test_fetch_team_members_succeeds_when_the_list_ends() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"value": [{"userId": "aad-1", "displayName": "Ala"}]})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    assert fetch_team_members(client, "team-1", "tok") == {"aad-1": "Ala"}


# --- źródło zmian z Graph (Shifts, ADR 0036) --------------------------------------

_SHARED = {"startDateTime": "2026-07-15T06:00:00Z", "endDateTime": "2026-07-15T14:00:00Z"}


def _shift_source(handler) -> GraphShiftSource:
    client = httpx.Client(transport=httpx.MockTransport(handler))
    return GraphShiftSource(lambda: "tok", "team-1", client=client)


def test_shift_source_parses_only_published_shifts_with_a_user() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["Authorization"] == "Bearer tok"
        return httpx.Response(
            200,
            json={
                "value": [
                    {"userId": "aad-1", "sharedShift": dict(_SHARED)},
                    {"userId": "aad-2", "draftShift": dict(_SHARED)},  # robocza — pomijamy
                    {"userId": "", "sharedShift": dict(_SHARED)},  # bez userId
                    "śmieć",
                ]
            },
        )

    blocks = _shift_source(handler).read_blocks()
    assert [b.user_id for b in blocks] == ["aad-1"]  # tylko opublikowana z userId
    assert blocks[0].start.tzinfo is not None  # świadomy datetime


def test_shift_source_skips_bad_dates_and_inverted_ranges() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "value": [
                    {"userId": "a", "sharedShift": {**_SHARED, "startDateTime": "niedata"}},
                    {  # koniec przed startem
                        "userId": "b",
                        "sharedShift": {
                            "startDateTime": "2026-07-15T14:00:00Z",
                            "endDateTime": "2026-07-15T06:00:00Z",
                        },
                    },
                ]
            },
        )

    assert _shift_source(handler).read_blocks() == []


def test_shift_source_follows_pagination() -> None:
    pages = {
        "/v1.0/teams/team-1/schedule/shifts": {
            "value": [{"userId": "a", "sharedShift": dict(_SHARED)}],
            "@odata.nextLink": "https://graph.microsoft.com/v1.0/next",
        },
        "/v1.0/next": {"value": [{"userId": "b", "sharedShift": dict(_SHARED)}]},
    }

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=pages[request.url.path])

    assert {b.user_id for b in _shift_source(handler).read_blocks()} == {"a", "b"}


def test_shift_source_fails_loudly_when_the_page_cap_cuts_the_list() -> None:
    """Ucięcie grafiku = zaniżone godziny w arkuszu importowanym jako fakt — twardy błąd."""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "value": [{"userId": "a", "sharedShift": dict(_SHARED)}],
                "@odata.nextLink": "https://graph.microsoft.com/v1.0/next",
            },
        )

    with pytest.raises(ValueError, match="NIEKOMPLETNA"):
        _shift_source(handler).read_blocks()


# --- git_email case-insensitive (ADR 0036, L3) ------------------------------------


def test_git_email_is_canonicalized_to_lowercase(tmp_path: Path) -> None:
    """Git ignoruje wielkość liter e-maila autora — kanonizujemy przy ładowaniu do małych."""
    yaml = "EMP-1:\n  aad_user_id: a\n  jira_user: j@x.pl\n  git_email: Dev@Example.PL\n"
    person = YamlIdentityDirectory(_identities(tmp_path, yaml)).resolve("EMP-1")
    assert person is not None and person.git_email == "dev@example.pl"


def test_git_email_lookup_is_case_insensitive(tmp_path: Path) -> None:
    yaml = "EMP-1:\n  aad_user_id: a\n  jira_user: j@x.pl\n  git_email: dev@example.pl\n"
    directory = YamlIdentityDirectory(_identities(tmp_path, yaml))
    assert directory.resolve_by_git_email("DEV@Example.PL") is not None  # inny case → ta sama osoba


def test_shared_git_email_differing_only_in_case_is_rejected(tmp_path: Path) -> None:
    """Dwie osoby z tym samym e-mailem różniącym się wielkością liter to wciąż kolizja."""
    yaml = (
        "EMP-1:\n  aad_user_id: a1\n  jira_user: a@x.pl\n  git_email: Dev@x.pl\n"
        "EMP-2:\n  aad_user_id: a2\n  jira_user: b@x.pl\n  git_email: dev@X.pl\n"
    )
    with pytest.raises(ValueError, match="git_email"):
        YamlIdentityDirectory(_identities(tmp_path, yaml))
