from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ceidg_tool.errors import ExportError
from ceidg_tool.exporter import (
    check_free_space,
    part_paths,
    plan_export,
    write_csv,
    write_jsonl,
    write_workbook,
)
from ceidg_tool.normalizer import SHEETS, NormalizedRecord, columns, normalize
from ceidg_tool.records import RawRecord, RowContext
from tests.conftest import detail_record, list_record

CTX = RowContext(srodowisko="test", pobrano_utc="2026-09-05T10:00:00Z")
META = [("kryteria", "wojewodztwo: podlaskie"), ("srodowisko", "test"), ("count", 2)]


def make_records(n: int = 2, with_detail: bool = True) -> list[NormalizedRecord]:
    out = []
    for idx in range(1, n + 1):
        base = detail_record(idx)
        detail = detail_record(
            idx,
            wlasciciel={**base["wlasciciel"], "nip": "0123456789"},
            adresDzialalnosci={**base["adresDzialalnosci"], "kod": "01-234"},
        )
        r = RawRecord(
            id=list_record(idx)["id"],
            list_json=list_record(idx),
            detail_json=detail if with_detail else None,
            list_utc="2026-09-05T09:00:00Z",
            detail_utc="2026-09-05T09:30:00Z" if with_detail else None,
            detail_state="pobrany" if with_detail else "brak",
            zrodlo="CEIDG_API",
        )
        out.append(normalize(r, CTX))
    return out


def source_of(records: list[NormalizedRecord]):  # type: ignore[no-untyped-def]
    return lambda: iter(records)


def test_workbook_has_tables_freeze_and_text_ids(tmp_path: Path) -> None:
    paths = write_workbook(tmp_path / "out.xlsx", source_of(make_records()), metadata=META)
    assert paths == [tmp_path / "out.xlsx"]
    wb = load_workbook(paths[0])
    assert wb.sheetnames == ["Firmy", "PKD", "Spolki", "Slownik", "Metadane"]

    firmy = wb["Firmy"]
    assert [c.value for c in firmy[1]] == list(columns("Firmy"))
    # zamrożone są nagłówek i kolumny tożsamości (nip, nazwa), nie sam nagłówek —
    # przy 43 kolumnach bez tego kontakty i PKD wiszą bez wiersza, do którego należą
    assert firmy.freeze_panes == "C2"
    assert "tbl_Firmy" in firmy.tables
    assert firmy.tables["tbl_Firmy"].autoFilter is not None

    header = {c.value: i + 1 for i, c in enumerate(firmy[1])}
    nip_cell = firmy.cell(row=2, column=header["nip"])
    assert nip_cell.value == "0123456789"
    assert nip_cell.number_format == "@"
    kod_cell = firmy.cell(row=2, column=header["kod_pocztowy"])
    assert kod_cell.value == "01-234" and kod_cell.number_format == "@"
    date_cell = firmy.cell(row=2, column=header["data_rozpoczecia"])
    assert isinstance(date_cell.value, datetime)
    assert date_cell.value.date() == date(2014, 7, 29)
    assert date_cell.number_format == "YYYY-MM-DD"
    assert firmy.cell(row=2, column=header["pobrano_utc"]).value == "2026-09-05T09:30:00Z"

    pkd = wb["PKD"]
    assert pkd.max_row == 1 + 8  # 2 firmy × 4 PKD
    assert "tbl_PKD" in pkd.tables


def test_slownik_covers_all_columns_and_metadata_written(tmp_path: Path) -> None:
    (dest,) = write_workbook(tmp_path / "out.xlsx", source_of(make_records()), metadata=META)
    wb = load_workbook(dest)
    described = {(r[0].value, r[1].value) for r in wb["Slownik"].iter_rows(min_row=2)}
    for sheet in SHEETS:
        for col in columns(sheet):
            assert (sheet, col) in described
    meta = {r[0].value: r[1].value for r in wb["Metadane"].iter_rows(min_row=2)}
    assert meta["srodowisko"] == "test"
    assert meta["count"] == 2


def test_empty_related_sheets_are_omitted_but_firmy_always_present(tmp_path: Path) -> None:
    (dest,) = write_workbook(
        tmp_path / "out.xlsx", source_of(make_records(with_detail=False)), metadata=META
    )
    assert load_workbook(dest).sheetnames == ["Firmy", "Slownik", "Metadane"]

    (empty,) = write_workbook(tmp_path / "empty.xlsx", source_of([]), metadata=META)
    wb2 = load_workbook(empty)
    assert wb2.sheetnames == ["Firmy", "Slownik", "Metadane"]
    assert wb2["Firmy"].max_row == 1
    assert not wb2["Firmy"].tables  # tabela nad samym nagłówkiem psuje plik w Excelu


def test_export_is_split_into_parts_above_row_limit(tmp_path: Path) -> None:
    records = make_records(5)  # każda firma: 1 wiersz Firmy, 4 PKD, 1 Spolki
    plan = plan_export(source_of(records), row_limit=9)
    assert plan.parts == (2, 2, 1)  # 2 firmy = 8 wierszy PKD ≤ 9; trzecia przekroczyłaby
    assert plan.rows["PKD"] == 20

    paths = write_workbook(tmp_path / "duzy.xlsx", source_of(records), metadata=META, row_limit=9)
    assert [p.name for p in paths] == [
        "duzy_czesc01.xlsx",
        "duzy_czesc02.xlsx",
        "duzy_czesc03.xlsx",
    ]
    firms = 0
    for i, p in enumerate(paths, start=1):
        wb = load_workbook(p)
        assert wb.sheetnames == ["Firmy", "PKD", "Spolki", "Slownik", "Metadane"]
        firms += wb["Firmy"].max_row - 1
        assert wb["PKD"].max_row - 1 <= 9
        meta = {r[0].value: r[1].value for r in wb["Metadane"].iter_rows(min_row=2)}
        assert meta["czesc"] == f"{i}/3"
    assert firms == 5


def test_part_paths_and_free_space_check(tmp_path: Path) -> None:
    assert part_paths(Path("a/b.xlsx"), 1) == [Path("a/b.xlsx")]
    assert [p.name for p in part_paths(Path("a/b.xlsx"), 2)] == ["b_czesc01.xlsx", "b_czesc02.xlsx"]
    with pytest.raises(ExportError, match="miejsca"):
        check_free_space(tmp_path, expected_rows=10, min_free_bytes=10**18)


def test_csv_and_jsonl_share_columns(tmp_path: Path) -> None:
    paths = write_csv(tmp_path / "csv", source_of(make_records()))
    assert {p.name for p in paths} == {"firmy.csv", "pkd.csv", "spolki.csv"}
    with (tmp_path / "csv" / "firmy.csv").open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    assert rows[0] == list(columns("Firmy"))
    assert rows[1][rows[0].index("nip")] == "0123456789"
    assert rows[1][rows[0].index("data_rozpoczecia")] == "2014-07-29"

    jsonl = write_jsonl(tmp_path / "out.jsonl", source_of(make_records()))
    lines = jsonl.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2
    obj = json.loads(lines[0])
    assert obj["nip"] == "0123456789"
    assert len(obj["pkd"]) == 4
    assert obj["data_rozpoczecia"] == "2014-07-29"


def test_source_is_reiterated_so_multiple_exports_are_complete(tmp_path: Path) -> None:
    records = make_records(3)
    source = source_of(records)
    write_workbook(tmp_path / "a.xlsx", source, metadata=META)
    paths = write_csv(tmp_path / "csv", source)
    with (tmp_path / "csv" / "firmy.csv").open(encoding="utf-8-sig", newline="") as fh:
        assert len(list(csv.reader(fh, delimiter=";"))) == 4
    assert paths
