"""Kreator (`ui/wizard`): menu z §A, kolejność pozycji i zbieranie kryteriów.

Kreator ma być cienki — całą kolejność decyzji trzyma `flow`. Testy pilnują więc trzech
rzeczy, za które odpowiada sam kreator: wznowienie pojawia się **pierwsze i tylko wtedy**,
gdy w bazie czeka niedokończone zadanie; każda pozycja menu prowadzi do właściwej operacji;
puste kryteria są odrzucane przy zbieraniu, **zanim** poleci jakiekolwiek żądanie.

Wszystko bez terminala: `ScriptedPrompter` odpowiada po identyfikatorze pytania.
"""

from __future__ import annotations

from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from ceidg_tool.config import Settings
from ceidg_tool.errors import CeidgError, ConfigError
from ceidg_tool.pipeline import Deps, build_deps
from ceidg_tool.ui import texts, wizard
from ceidg_tool.ui.prompts import CancelledError, Question, ScriptedPrompter
from tests.conftest import FakeClock, list_record
from tests.support import FakeApi, RecordingView, criteria

BASE = "https://test-dane.biznes.gov.pl/api/ceidg/v3"
VERSION = "0.1.0"

# Odpowiedzi na komplet pytań o kryteria; test nadpisuje tylko to, co go interesuje.
BLANK_CRITERIA: dict[str, object] = {
    "wojewodztwo": "",
    "miasto": "",
    "pkd": "",
    "status": "",
    "data_od": "",
    "data_do": "",
    "nazwa": "",
    "max_rekordow": "",
}


def criteria_answers(**overrides: object) -> dict[str, object]:
    return {**BLANK_CRITERIA, **overrides}


def deps_for(tmp_path: Path, clock: FakeClock, api: FakeApi) -> Deps:
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    return build_deps(settings, clock=clock, http=api.client())


def silent_api() -> FakeApi:
    """API wywracające test przy każdym żądaniu — dla ścieżek, które nie mają sięgać sieci."""
    api = FakeApi()

    def fallback(request: httpx.Request) -> httpx.Response:
        raise AssertionError(f"nieoczekiwane żądanie: {request.url}")

    api.fallback = fallback
    return api


def seed_unfinished(deps: Deps, run_id: str = "run-przerwany") -> str:
    source = criteria(wojewodztwo="podlaskie")
    deps.store.start_run(
        run_id=run_id,
        criteria_json=source.model_dump_json(),
        criteria_hash=source.fingerprint(),
        profile_hash=deps.profile.profile_hash(),
        mode="lista",
        tool_version=VERSION,
        cursor_mode="links",
    )
    deps.store.update_run_status(run_id, "przerwany")
    return run_id


# ----------------------------------------------------------------------------- menu


def test_the_menu_lists_the_five_actions_from_section_a(tmp_path: Path, clock: FakeClock) -> None:
    """§A wymienia pięć działań plus wyjście; brak pozycji = funkcja nieosiągalna z kreatora."""
    deps = deps_for(tmp_path, clock, silent_api())

    keys = [item.key for item in wizard._menu_items(deps)]

    assert keys == ["pobierz", "aktualizuj", "raport", "nip", "wyjdz"]
    deps.store.close()


def test_resume_is_absent_when_nothing_was_interrupted(tmp_path: Path, clock: FakeClock) -> None:
    """Pozycja „wznów” bez czego wznawiać byłaby ślepą uliczką."""
    deps = deps_for(tmp_path, clock, silent_api())

    assert "wznow" not in [item.key for item in wizard._menu_items(deps)]
    deps.store.close()


def test_resume_is_the_first_item_when_a_job_is_waiting(tmp_path: Path, clock: FakeClock) -> None:
    """§A punkt 4: niedokończona praca ma się rzucać w oczy, zanim ktoś zacznie nową."""
    deps = deps_for(tmp_path, clock, silent_api())
    run_id = seed_unfinished(deps)

    items = wizard._menu_items(deps)

    assert items[0].key == "wznow"
    assert run_id[:8] in items[0].hint  # widać, o które zadanie chodzi
    deps.store.close()


def test_a_finished_run_does_not_bring_back_the_resume_item(
    tmp_path: Path, clock: FakeClock
) -> None:
    deps = deps_for(tmp_path, clock, silent_api())
    run_id = seed_unfinished(deps)
    deps.store.update_run_status(run_id, "zakonczony")

    assert "wznow" not in [item.key for item in wizard._menu_items(deps)]
    deps.store.close()


def test_choosing_exit_ends_the_session_with_code_zero(tmp_path: Path, clock: FakeClock) -> None:
    """Wyjście z menu to normalne zakończenie sesji, nie błąd."""
    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()
    # Bez wznowienia „wyjdź” jest piątą pozycją.
    prompter = ScriptedPrompter({"menu": "5"})

    code = wizard.run_wizard(deps, prompter, view, version=VERSION)

    assert code == 0
    assert "ceidg-tool" in view.text()  # pierwszy ekran padł przed menu
    deps.store.close()


def test_the_first_screen_is_shown_once_before_the_menu(tmp_path: Path, clock: FakeClock) -> None:
    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()

    wizard.run_wizard(deps, ScriptedPrompter({"menu": "5"}), view, version=VERSION)

    titles = view.titles()
    assert titles[0].startswith("ceidg-tool")
    assert titles.count(titles[0]) == 1
    deps.store.close()


def test_an_error_in_one_action_returns_to_the_menu_instead_of_ending_the_session(
    tmp_path: Path, clock: FakeClock
) -> None:
    """ADR-0008: w sesji interaktywnej błąd jednego działania nie może wywalić całej sesji."""
    api = FakeApi()
    api.fallback = lambda request: httpx.Response(500, json={"code": "E", "message": "awaria"})
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()
    # 1 = pobierz (padnie na 500 przy szukaniu raportu), potem 5 = wyjdź.
    prompter = ScriptedPrompter(
        {
            "menu": ["1", "5"],
            **criteria_answers(wojewodztwo="podlaskie"),
        }
    )

    code = wizard.run_wizard(deps, prompter, view, version=VERSION)

    assert code == 0
    assert view.errors  # błąd pokazany...
    assert prompter.asked.count("menu") == 2  # ...a menu wróciło
    deps.store.close()


# ----------------------------------------------------------------------------- zbieranie kryteriów


def test_empty_criteria_are_refused_at_collection_before_any_request(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Bez filtra zapytanie objęłoby cały rejestr — kreator pyta ponownie, nie woła API."""
    api = silent_api()
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()
    answers: dict[str, object] = {
        key: ["", "podlaskie"] if key == "wojewodztwo" else "" for key in BLANK_CRITERIA
    }
    answers["ponow_kryteria"] = True
    prompter = ScriptedPrompter(answers)

    result = wizard.collect_criteria(prompter, view)

    assert result.wojewodztwo == ("podlaskie",)
    assert any("cały rejestr" in e for e in view.errors)
    assert api.requests == []
    deps.store.close()


def test_invalid_criteria_are_reported_and_asked_again() -> None:
    """Literówka w województwie ma wrócić jako pytanie, a nie wywrócić kreator."""
    view = RecordingView()
    prompter = ScriptedPrompter(
        {
            **BLANK_CRITERIA,
            "wojewodztwo": ["nieistniejace", "podlaskie"],
            "ponow_kryteria": True,
        }
    )

    result = wizard.collect_criteria(prompter, view)

    assert result.wojewodztwo == ("podlaskie",)
    assert any("Niepoprawne kryteria" in e for e in view.errors)


def test_declining_the_retry_leaves_the_criteria_loop(tmp_path: Path, clock: FakeClock) -> None:
    """Bez tej furtki operator, który przeklikał puste odpowiedzi, nie miałby jak wrócić
    do menu inaczej niż zabijając proces."""
    view = RecordingView()
    prompter = ScriptedPrompter({**BLANK_CRITERIA, "ponow_kryteria": False})

    with pytest.raises(CancelledError):
        wizard.collect_criteria(prompter, view)


def test_collected_criteria_accept_a_comma_separated_list() -> None:
    view = RecordingView()
    prompter = ScriptedPrompter(criteria_answers(miasto="Białystok, Łomża"))

    result = wizard.collect_criteria(prompter, view)

    assert result.miasto == ("Białystok", "Łomża")


def test_the_allowed_values_are_shown_before_the_questions() -> None:
    """Operator nie zna zamkniętych list — bez podpowiedzi każde pytanie to zgadywanka."""
    view = RecordingView()

    wizard.collect_criteria(ScriptedPrompter(criteria_answers(wojewodztwo="podlaskie")), view)

    assert "Dozwolone wartości" in view.titles()


# ------------------------------------------------------------------ powtórzenie zapytania


def test_powtorzenie_pokazuje_polecenie_i_nie_pisze_zadnego_pliku(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Kreator kończy poleceniem do skopiowania, nie plikiem (ADR-0022).

    Plik zapytania YAML był tu do 2026-09-10 i kosztował operatora osobne pytanie oraz
    format, którego musiał się nauczyć, żeby zrobić to samo, co robi jedna linia. Asercja
    o braku plików jest w tym teście celowo: gdyby zapis został gdzieś z tyłu, katalog
    wyników znów zbierałby pliki, o które nikt nie prosił.
    """
    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()

    wizard.show_repeat_command(criteria(wojewodztwo="podlaskie", pkd="62.01.Z"), deps, view)

    tekst = view.text()
    assert "ceidg-tool pobierz --wojewodztwo podlaskie --pkd 6201Z --tak" in tekst
    assert not list(deps.settings.output_dir.glob("*.yaml"))
    deps.store.close()


def test_powtorzenie_w_pokazie_niesie_flage_demo(tmp_path: Path, clock: FakeClock) -> None:
    """Polecenie z pokazu musi zostać poleceniem pokazu — inaczej wklejone sięga rejestru.

    To jest znacznik ADR-0014 na czwartym kanale: pierwszy ekran, wiersz `Metadane`, prefiks
    `DEMO_` i osobny katalog mówią „to pokaz", a polecenie bez `--demo` mówiłoby coś innego —
    i to akurat w miejscu, które operator kopiuje do harmonogramu.
    """
    deps = deps_for(tmp_path, clock, silent_api())
    deps.demo = True
    view = RecordingView()

    wizard.show_repeat_command(criteria(wojewodztwo="podlaskie"), deps, view)

    assert "ceidg-tool pobierz --demo --wojewodztwo podlaskie --tak" in view.text()
    deps.store.close()


# ----------------------------------------------------------------------------- routing pozycji


def test_resume_without_any_interrupted_run_says_so_instead_of_failing(
    tmp_path: Path, clock: FakeClock
) -> None:
    api = silent_api()
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()

    wizard.handle_resume(deps, ScriptedPrompter({}), view)

    assert "Brak przerwanych pobrań." in view.messages
    assert api.requests == []
    deps.store.close()


def test_a_resumable_run_with_unreadable_criteria_is_reported_as_configuration_error(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Uszkodzony wiersz w bazie ma dać czytelny komunikat, nie `ValidationError` na ekran."""
    api = silent_api()
    deps = deps_for(tmp_path, clock, api)
    deps.store.start_run(
        run_id="run-zepsuty",
        criteria_json="{nie-json",
        criteria_hash="hash",
        profile_hash=deps.profile.profile_hash(),
        mode="lista",
        tool_version=VERSION,
        cursor_mode="links",
    )
    deps.store.update_run_status("run-zepsuty", "przerwany")

    with pytest.raises(CeidgError, match="nieczytelne kryteria"):
        wizard.handle_resume(deps, ScriptedPrompter({}), RecordingView())

    assert api.requests == []
    deps.store.close()


def test_the_nip_action_shows_a_card_without_exporting_when_nothing_was_found(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Brak wpisu to wynik: karta z wyjaśnieniem i żadnego skoroszytu."""
    api = FakeApi()
    api.fallback = lambda request: httpx.Response(204)
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()

    wizard.handle_nip(deps, ScriptedPrompter({"nip": "3563457932"}), view)

    assert "NIP 3563457932" in view.titles()
    assert "brak wpisu" in view.text()
    assert not list(deps.settings.output_dir.glob("*.xlsx"))
    deps.store.close()


def test_the_nip_action_refuses_a_bad_checksum_without_a_request(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Ta sama bramka co w `sprawdz-nip`: literówka nie kosztuje żądania."""
    api = silent_api()
    deps = deps_for(tmp_path, clock, api)

    with pytest.raises(CeidgError, match="Niepoprawny NIP"):
        wizard.handle_nip(deps, ScriptedPrompter({"nip": "3563457931"}), RecordingView())

    assert api.requests == []
    deps.store.close()


def test_an_update_with_no_changes_does_not_export(tmp_path: Path, clock: FakeClock) -> None:
    """Zero zmienionych wpisów to zero powodów, żeby cokolwiek robić.

    Od 2026-09-06 `aktualizuj` pyta przed startem, więc pusty zakres w ogóle nie dochodzi
    do pytania: tanie zapytanie o `count` mówi, że nie ma czego pobierać, i na tym koniec."""
    api = FakeApi()
    api.fallback = lambda request: httpx.Response(204)
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()

    wizard.handle_update(deps, ScriptedPrompter({}), view)

    assert any(texts.update_declined(0) in m for m in view.messages)
    assert not list(deps.settings.output_dir.glob("*.xlsx"))
    deps.store.close()


def test_an_update_the_operator_declines_makes_no_fetch(tmp_path: Path, clock: FakeClock) -> None:
    """Sedno zmiany: koszt pada przed pracą, a odmowa naprawdę zatrzymuje.

    Wcześniej wybranie „Zaktualizować bazę" startowało pobranie od razu — przy 2891 zmianach
    trzydzieści kilka minut, o których nikt nie uprzedził."""
    ids = [{"id": f"id-{i}"} for i in range(40)]
    api = FakeApi()
    api.fallback = lambda request: httpx.Response(
        200,
        json={"identyfikatoryWpisow": ids, "count": 40, "links": {"self": str(request.url)}},
    )
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()

    wizard.handle_update(deps, ScriptedPrompter({"aktualizowac": False}), view)

    assert any("40" in str(row) for block in view.blocks for row in block.rows)
    assert any(texts.update_declined(40) in m for m in view.messages)
    # jedno tanie zapytanie o `count`, ani jednego o szczegóły
    assert len(api.requests) == 1
    assert not any("/firma" in url for url in api.requests)
    deps.store.close()


def test_the_fetch_action_stops_when_the_operator_leaves_at_the_cost_table(
    tmp_path: Path, clock: FakeClock
) -> None:
    """„Wyjdź” przy tabeli kosztów kończy działanie bez pobrania i bez pliku."""
    api = FakeApi()

    def fallback(request: httpx.Request) -> httpx.Response:
        if "/firmy" in str(request.url) and request.url.params.get("limit") == "1":
            return httpx.Response(200, json={"count": 1_240, "firmy": []})
        raise AssertionError(f"pobranie mimo wyjścia: {request.url}")

    api.fallback = fallback
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()
    prompter = ScriptedPrompter(
        {
            **criteria_answers(miasto="Białystok"),
            "co_dalej": "wyjdz",
        }
    )

    wizard.handle_fetch(deps, prompter, view, source="auto")

    assert len(api.requests) == 1
    assert not list(deps.settings.output_dir.glob("*.xlsx"))
    deps.store.close()


def test_correcting_the_criteria_collects_them_again_and_recounts(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Pętla „popraw kryteria”: drugie kryteria dostają własne, jedno zapytanie o `count`."""
    api = FakeApi()

    def fallback(request: httpx.Request) -> httpx.Response:
        if "/firmy" in str(request.url) and request.url.params.get("limit") == "1":
            return httpx.Response(200, json={"count": 1_240, "firmy": []})
        raise AssertionError(f"nieoczekiwane pobranie: {request.url}")

    api.fallback = fallback
    deps = deps_for(tmp_path, clock, api)
    view = RecordingView()
    prompter = ScriptedPrompter(
        {
            **{key: ["", ""] for key in BLANK_CRITERIA if key != "miasto"},
            "miasto": ["Białystok", "Łomża"],
            "co_dalej": ["popraw", "wyjdz"],
        }
    )

    wizard.handle_fetch(deps, prompter, view, source="auto")

    assert prompter.asked.count("co_dalej") == 2
    assert len([r for r in api.requests if "limit=1" in r]) == 2
    deps.store.close()


def test_ctrl_c_at_the_menu_ends_the_session_like_choosing_exit(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Przerwanie w menu to rezygnacja, nie błąd konfiguracji — inaczej sesja kończyłaby się
    kodem 3 i komunikatem o błędzie za naciśnięcie Ctrl+C."""

    class Interrupting(ScriptedPrompter):
        def ask(self, question: Question) -> str:
            raise CancelledError("Przerwano wybór.")

    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()

    code = wizard.run_wizard(deps, Interrupting({}), view, version=VERSION)

    assert code == 0
    assert not view.errors
    assert any("Przerwano" in message for message in view.messages)
    deps.store.close()


def resumable_api() -> FakeApi:
    """API, które oddaje jedną stronę listy — tyle wystarczy, by wznowienie się zakończyło."""
    api = FakeApi()

    def fallback(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/firmy" not in url:
            raise AssertionError(f"nieoczekiwane żądanie: {url}")
        return httpx.Response(
            200, json={"count": 1, "firmy": [list_record(1, id="rec-1")], "links": {}}
        )

    api.fallback = fallback
    return api


def test_the_resume_action_asks_for_the_purpose_before_it_starts_fetching(
    tmp_path: Path, clock: FakeClock
) -> None:
    """uzupelnienie-01.md §B wymaga pola „cel pobrania” w Metadanych, a wznowienie potrafi
    trwać godziny — pytanie po pobraniu zostawiłoby gotowy run bez skoroszytu, gdyby
    operator je przerwał, bo skończony run wypada z listy do wznowienia."""
    deps = deps_for(tmp_path, clock, resumable_api())
    seed_unfinished(deps)
    view = RecordingView()
    prompter = ScriptedPrompter({"cel": "audyt zgodności"})

    wizard.handle_resume(deps, prompter, view)

    assert prompter.asked == ["cel"]  # pytanie padło i to jako pierwsze, przed pobraniem
    workbook = next((tmp_path / "dane" / "wyniki").glob("*.xlsx"))
    metadata = {
        row[0].value: row[1].value
        for row in load_workbook(workbook)["Metadane"].iter_rows(min_row=2)
    }
    assert metadata["cel_pobrania"] == "audyt zgodności"
    deps.store.close()


def test_a_mistyped_menu_answer_asks_again_instead_of_ending_the_session(
    tmp_path: Path, clock: FakeClock
) -> None:
    """W trybie awaryjnym (`input` zamiast strzałek) da się wpisać cokolwiek. Pomyłka przy
    menu to nie błąd konfiguracji — sesja kończyłaby się wtedy kodem 3 za literówkę."""

    class Fumbling(ScriptedPrompter):
        def __init__(self) -> None:
            super().__init__({})
            self.attempts = 0

        def ask(self, question: Question) -> str:
            self.attempts += 1
            if self.attempts == 1:
                raise ConfigError("Nieznana odpowiedź '9' na pytanie 'menu'.")
            return question.check("5")

    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()
    prompter = Fumbling()

    code = wizard.run_wizard(deps, prompter, view, version=VERSION)

    assert code == 0
    assert prompter.attempts == 2
    assert any("Nieznana odpowiedź" in error for error in view.errors)
    deps.store.close()


def test_cancelling_at_the_criteria_stage_does_not_promise_data_that_was_never_fetched(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Rezygnacja przy kryteriach nie zakłada żadnego runu, więc podpowiedź o eksporcie
    byłaby nieprawdą — a `eksportuj` bez `--run-id` wydałby najnowszy run w bazie,
    czyli cudzy wynik podany jako rezultat tego, co operator właśnie robił."""
    deps = deps_for(tmp_path, clock, silent_api())
    view = RecordingView()
    prompter = ScriptedPrompter({"menu": ["1", "5"], **BLANK_CRITERIA, "ponow_kryteria": False})

    code = wizard.run_wizard(deps, prompter, view, version=VERSION)

    assert code == 0
    assert any("Rezygnacja z kryteriów" in message for message in view.messages)
    assert not any("eksportuj" in message for message in view.messages)
    assert deps.store.list_runs() == []
    deps.store.close()


def test_cancelling_at_the_save_question_names_the_run_to_export(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Tam, gdzie pytanie pada **po** pobraniu, run już istnieje — więc podpowiedź musi paść,
    i z identyfikatorem, bo `eksportuj` bez `--run-id` bierze najnowszy run w bazie."""
    deps = deps_for(tmp_path, clock, silent_api())
    run_id = seed_unfinished(deps, "run-gotowy")
    deps.store.update_run_status(run_id, "zakonczony")
    view = RecordingView()

    class CancelAtSave(ScriptedPrompter):
        def confirm(self, question: Question, *, default: bool) -> bool:
            raise CancelledError("Przerwano pytanie.")

    wizard._export_on_request(deps, CancelAtSave({}), view, run_id)

    hint = " ".join(view.messages)
    assert "Przerwano pytanie." in hint
    assert f"eksportuj --run-id {run_id}" in hint
    assert not list((tmp_path / "dane" / "wyniki").glob("*.xlsx"))
    deps.store.close()


# ------------------------------------------- ścieżka opisowa musi być widoczna z menu


def test_menu_mowi_ze_wystarczy_opisac_zdaniem_gdy_asystent_dziala(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Funkcja istniejąca i niewidoczna jest z punktu widzenia operatora nieistniejąca.

    Opis zdaniem jest pierwszym pytaniem tej ścieżki od fazy 4 (ADR-0011, decyzja 5), ale
    menu mówiło „Pobrać firmy **według kryteriów** — lista albo szczegóły", więc operator
    czytał „formularz" i nie miał skąd wiedzieć, że wystarczy napisać zdanie. To narzędzie
    istnieje dla kogoś, kto nie zna API ani kodów PKD — jeśli nie widzi drogi bez kodów PKD,
    lepiej obsłuży go rządowa wyszukiwarka.
    """
    deps = deps_for(tmp_path, clock, FakeApi())
    deps.assistant = object()  # type: ignore[assignment]

    pozycja = next(i for i in wizard._menu_items(deps) if i.key == "pobierz")

    assert "opisz zdaniem" in pozycja.label.casefold()
    assert "salony fryzjerskie" in pozycja.hint, "przykład pokazuje zakres, sama zachęta nie"
    deps.store.close()


def test_menu_nie_obiecuje_opisu_gdy_asystenta_nie_ma(tmp_path: Path, clock: FakeClock) -> None:
    """Kontrola pozytywna, i nie formalna: bez klucza ta droga **nie działa**, więc
    zapowiedzenie jej byłoby wysłaniem operatora w ścianę — dokładnie odwrotnie niż zamierza
    ta poprawka."""
    deps = deps_for(tmp_path, clock, FakeApi())
    deps.assistant = None

    pozycja = next(i for i in wizard._menu_items(deps) if i.key == "pobierz")

    assert "opisz zdaniem" not in pozycja.label.casefold()
    assert "według kryteriów" in pozycja.label
    deps.store.close()


def test_pytanie_o_opis_niesie_przyklad_z_dwiema_branzami_i_dwoma_miastami() -> None:
    """„Opisz, czego szukasz" nie mówi, **jak dużo** wolno napisać. Przykład mówi — i jest
    zmierzony: to zdanie przeszło przez prawdziwego asystenta 2026-09-09 i dało
    `miasto: Gdańsk, Wrocław`, `PKD: 6622Z, 9621Z`."""
    from ceidg_tool.ui.prompts import OPIS

    assert texts.PRZYKLAD_OPISU in OPIS.hint
    assert "Enter" in OPIS.hint, "wyjście do pytań po kolei musi zostać widoczne"
