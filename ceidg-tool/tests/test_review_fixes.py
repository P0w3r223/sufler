"""Testy dla poprawek z przeglądu kodu (2026-09-05) i scenariuszy z uzupelnienie-01.md §D."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from ceidg_tool.apiprofile import ApiProfile
from ceidg_tool.errors import ProfileMismatchError
from ceidg_tool.exporter import write_csv, write_workbook
from ceidg_tool.normalizer import normalize
from ceidg_tool.ratelimit import REASON_COOLDOWN, REASON_WINDOW, InMemoryHistory, RateLimiter
from ceidg_tool.records import RawRecord, RowContext
from ceidg_tool.safetext import sanitize_text
from ceidg_tool.store import Store
from tests.conftest import FakeClock, detail_record, list_record
from tests.support import criteria

CTX = RowContext(srodowisko="test", pobrano_utc="2026-09-05T10:00:00Z")
HOSTILE_NAMES = ["=CMD()", "+1", "@SUM(A1)", "-2+3", '=HYPERLINK("http://evil","klik")', "\tTAB"]


def raw_with_name(name: str, idx: int = 1) -> RawRecord:
    return RawRecord(
        id=list_record(idx)["id"],
        list_json=list_record(idx, nazwa=name),
        detail_json=detail_record(idx, nazwa=name, email="\x01ctrl\x0bchars@x.test"),
        list_utc="2026-09-05T09:00:00Z",
        detail_utc="2026-09-05T09:30:00Z",
        detail_state="pobrany",
        zrodlo="CEIDG_API",
    )


# --- scenariusz 4: dane z rejestru jako wrogie ---------------------------------------


def test_sanitize_text_neutralizes_formula_prefixes_and_control_chars() -> None:
    for name in HOSTILE_NAMES:
        assert sanitize_text(name).startswith("'")
    assert sanitize_text("ACME\x01\x0bSp") == "ACMESp"
    assert sanitize_text("linia1\nlinia2\tx") == "linia1\nlinia2\tx"
    assert sanitize_text("zwykła nazwa") == "zwykła nazwa"


def test_workbook_cells_with_hostile_names_are_text_not_formulas(tmp_path: Path) -> None:
    records = [normalize(raw_with_name(n, i), CTX) for i, n in enumerate(HOSTILE_NAMES, 1)]
    records.append(normalize(raw_with_name("x" * 5000, 9), CTX))
    (dest,) = write_workbook(
        tmp_path / "hostile.xlsx",
        lambda: iter(records),
        metadata=[("kryteria", "=1+1"), ("cel", "@test")],
    )
    wb = load_workbook(dest)
    ws = wb["Firmy"]
    header = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for row_idx, name in enumerate(HOSTILE_NAMES, start=2):
        cell = ws.cell(row=row_idx, column=header["nazwa"])
        assert cell.data_type == "s"
        assert cell.value == "'" + name
    email = ws.cell(row=2, column=header["email"]).value
    assert email == "ctrlchars@x.test"
    meta = {r[0].value: r[1] for r in wb["Metadane"].iter_rows(min_row=2)}
    assert meta["kryteria"].value == "'=1+1" and meta["kryteria"].data_type == "s"
    assert meta["cel"].value == "'@test"


def test_csv_cells_with_hostile_names_are_prefixed(tmp_path: Path) -> None:
    records = [normalize(raw_with_name(n, i), CTX) for i, n in enumerate(HOSTILE_NAMES, 1)]
    write_csv(tmp_path, lambda: iter(records))
    with (tmp_path / "firmy.csv").open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    col = rows[0].index("nazwa")
    for row, name in zip(rows[1:], HOSTILE_NAMES, strict=True):
        assert row[col] == "'" + name


def test_export_is_atomic_leaves_no_tmp_file(tmp_path: Path) -> None:
    records = [normalize(raw_with_name("ok"), CTX)]
    (dest,) = write_workbook(tmp_path / "a.xlsx", lambda: iter(records), metadata=[])
    assert dest.exists()
    assert not list(tmp_path.glob(".*.tmp"))


# --- kryteria: odcisk niezależny od kolejności, limit długości --------------------------


def test_fingerprint_independent_of_list_element_order() -> None:
    a = criteria(status=["ZAWIESZONY", "AKTYWNY"], miasto=["Łomża", "Białystok"])
    b = criteria(status=["AKTYWNY", "ZAWIESZONY"], miasto=["Białystok", "Łomża"])
    assert a.fingerprint() == b.fingerprint()
    assert a.to_params(ApiProfile(base_url="https://test-dane.biznes.gov.pl/api/ceidg/v3")) == (
        b.to_params(ApiProfile(base_url="https://test-dane.biznes.gov.pl/api/ceidg/v3"))
    )


def test_text_fields_have_length_limit() -> None:
    with pytest.raises(ValidationError, match="limit 200"):
        criteria(nazwa="a" * 201)


# --- profil: tylko dozwolone hosty -----------------------------------------------------


@pytest.mark.parametrize(
    "url",
    [
        "https://evil.example.com/api/ceidg/v3",
        "https://dane.biznes.gov.pl.evil.example/api",
        "http://dane.biznes.gov.pl/api/ceidg/v3",
    ],
)
def test_profile_rejects_hosts_outside_allowlist(url: str) -> None:
    with pytest.raises(ValidationError, match="dozwolon|hostów"):
        ApiProfile(base_url=url)


# --- normalizer: klucze nadrzędne wygrywają z danymi z rejestru ------------------------


def test_registry_data_cannot_override_parent_keys() -> None:
    rec = RawRecord(
        id="AAA",
        list_json=list_record(1),
        detail_json=detail_record(
            1,
            adresyDzialalnosciDodatkowe=[{"id": "PODMIENIONE", "nip": "9999999999", "miasto": "X"}],
            pkd=[{"kod": "62.01.Z", "nazwa": "n", "id": "ZLE"}],
            pkdGlowny={"kod": "6201Z", "nazwa": "n"},
        ),
        list_utc=None,
        detail_utc=None,
        detail_state="pobrany",
        zrodlo="CEIDG_API",
    )
    out = normalize(rec, CTX)
    assert out.adresy[0]["id"] == "AAA"
    assert out.adresy[0]["nip"] == "3563457932"
    assert out.pkd[0]["id"] == "AAA"
    assert out.pkd[0]["czy_glowny"] is True  # 62.01.Z vs 6201Z porównane po normalizacji


# --- limiter: skok zegara ściennego nie kasuje okien ani blokady -----------------------


def make_limiter(
    clock: FakeClock, history: InMemoryHistory | None = None
) -> tuple[RateLimiter, list[tuple[float, str, float]]]:
    waits: list[tuple[float, str, float]] = []

    class Rec:
        def on_request(self, *a: object) -> None: ...
        def on_wait(self, seconds: float, reason: str, resume: float) -> None:
            waits.append((seconds, reason, resume))

        def on_page(self, *a: object) -> None: ...
        def on_details(self, *a: object) -> None: ...
        def on_export(self, *a: object) -> None: ...
        def on_download(self, *a: object) -> None: ...
        def on_model(self, *a: object) -> None: ...
        def on_message(self, *a: object) -> None: ...
        def close(self) -> None: ...

    limiter = RateLimiter(
        windows=((50, 180.0), (1000, 3600.0)),
        min_spacing_s=0.0,
        cooldown_s=185.0,
        clock=clock,
        history=history or InMemoryHistory(),
        events=Rec(),
    )
    return limiter, waits


def test_window_survives_wall_clock_jump(clock: FakeClock) -> None:
    limiter, waits = make_limiter(clock)
    for _ in range(50):
        limiter.acquire("firmy")
        clock.advance(1.0)
    clock.jump_wall(+7200.0)
    limiter.acquire("firmy")
    assert waits[-1][1] == REASON_WINDOW
    assert clock.sleeps[-1] == pytest.approx(130.0)


def test_cooldown_survives_wall_clock_jump(clock: FakeClock) -> None:
    limiter, waits = make_limiter(clock)
    limiter.acquire("firmy")
    limiter.note_response(429)
    clock.jump_wall(-7200.0)
    limiter.acquire("firmy")
    assert waits[-1][1] == REASON_COOLDOWN
    assert clock.sleeps[-1] == pytest.approx(185.0)


def test_identical_timestamps_in_history_do_not_crash(clock: FakeClock) -> None:
    from ceidg_tool.ratelimit import RequestStamp

    now = clock.wall()
    history = InMemoryHistory([RequestStamp(now - 10, "a", 200), RequestStamp(now - 10, "b", None)])
    limiter, _ = make_limiter(clock, history)
    limiter.acquire("firmy")


# --- store: nietypowy JSON, profil przy wznowieniu, wyścig blokady ---------------------


def test_store_accepts_odd_json_shapes(tmp_path: Path, clock: FakeClock) -> None:
    with Store(tmp_path / "s.sqlite", environment="test", clock=clock) as store:
        run_id = store.start_run(
            run_id="r",
            criteria_json="{}",
            criteria_hash="h",
            profile_hash="p",
            mode="lista",
            tool_version="0",
            cursor_mode="links",
        )
        odd: list[dict[str, Any]] = [
            {"id": "A", "wlasciciel": ["lista"], "adresDzialalnosci": "tekst"},
            {"id": "B", "wlasciciel": None, "adresDzialalnosci": {}},
        ]
        store.save_page(run_id, page_index=0, records=odd, next_cursor=None)
        assert store.count_run_records(run_id) == 2


def test_resume_refuses_different_profile(tmp_path: Path, clock: FakeClock) -> None:
    with Store(tmp_path / "s.sqlite", environment="test", clock=clock) as store:
        store.start_run(
            run_id="r",
            criteria_json="{}",
            criteria_hash="h",
            profile_hash="old",
            mode="lista",
            tool_version="0",
            cursor_mode="links",
        )
        with pytest.raises(ProfileMismatchError):
            store.find_resumable_run("h", profile_hash="new")
        found = store.find_resumable_run("h", profile_hash="old")
        assert found is not None and found.run_id == "r"


def test_lock_is_taken_atomically_and_same_pid_does_not_steal(
    tmp_path: Path, clock: FakeClock
) -> None:
    import sqlite3

    with Store(tmp_path / "s.sqlite", environment="test", clock=clock) as store:
        store.acquire_lock()
        # cudzy, żywy wpis o tym samym PID (po restarcie systemu PID-y się powtarzają)
        conn = sqlite3.connect(tmp_path / "s.sqlite")
        conn.execute("UPDATE run_lock SET started_utc = '2000-01-01T00:00:00Z'")
        conn.commit()
        conn.close()
        from ceidg_tool.errors import StoreLockedError

        with pytest.raises(StoreLockedError):
            store.acquire_lock()
        clock.advance(601)
        store.acquire_lock()
        store.release_lock()
