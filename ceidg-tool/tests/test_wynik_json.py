"""`--wynik json`: koperta na każdej drodze wyjścia i klasyfikacja zakończeń (ADR-0024).

Najważniejszy test w tym pliku jest ten o `finally`. Koperta pisana z `except CeidgError`
zachowuje się identycznie na **każdej** drodze, którą taksonomia obejmuje — a różni się
dokładnie tam, gdzie wołający jest najbardziej bezradny: przy wyjątku, którego nikt nie
przewidział. Pusty stdout i kod wyjścia bez treści to zakończenie, które ADR nazywa najgorszym
z możliwych, więc obserwatorem jest tu przejście po **wszystkich** drogach, nie po typowej.
"""

from __future__ import annotations

import io
import json
import sys
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
import typer
from openpyxl import load_workbook
from rich.console import Console
from typer.testing import CliRunner

from ceidg_tool import cli
from ceidg_tool.cli import KOPERTA_OBSLUGIWANA, app
from ceidg_tool.console import LineEvents
from ceidg_tool.criteria import Criteria
from ceidg_tool.errors import AuthError, ConfigError
from ceidg_tool.ui import flow, prompts
from ceidg_tool.ui.wynik import KLUCZE_WSPOLNE, WERSJA_KOPERTY, Wynik
from tests.support import zarejestrowane_polecenia

EXIT_CONFIG = 3
EXIT_PUSTO = 4


@pytest.fixture
def runner(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> CliRunner:
    """CLI w pustym katalogu — `szukaj-pkd` nie dotyka bazy, ale sąsiedzi już tak."""
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("CEIDG_DATA_DIR", str(tmp_path / "dane"))
    monkeypatch.delenv("FORCE_COLOR", raising=False)
    monkeypatch.setenv("TERM", "dumb")
    # Szerokość przypięta z tego samego powodu co w `tests/test_cli.py:46`: asercja o treści
    # zdania nie ma przy okazji mierzyć szerokości maszyny, bo `rich` zawija je gdzie indziej
    # na Windowsie i gdzie indziej w CI.
    monkeypatch.setattr(cli.console, "width", 200)
    return CliRunner()


@pytest.fixture
def env(tmp_path: Path) -> dict[str, str]:
    """Token i katalog danych — pokaz ich nie używa, ale `_settings` je rozwiązuje."""
    return {"CEIDG_TOKEN": "token-testowy-nie-jwt", "CEIDG_DATA_DIR": str(tmp_path / "dane")}


@contextmanager
def przechwycony_stdout(monkeypatch: pytest.MonkeyPatch) -> Iterator[io.StringIO]:
    """Podmieniony stdout do testów samego menedżera, bez przechodzenia przez typera.

    Funkcja, a nie fixture, i to jest nauczka z pierwszego podejścia: pytest przestawia
    `sys.stdout` **między** przygotowaniem a wywołaniem testu, więc podmiana zrobiona
    w fixture jest natychmiast nadpisywana. Koperta szła wtedy do przechwytu pytesta,
    a bufor testu zostawał pusty — czyli test oblewał z powodu, który z jego treścią nie
    miał nic wspólnego.
    """
    bufor = io.StringIO()
    monkeypatch.setattr(sys, "stdout", bufor)
    yield bufor


# --------------------------------------------- pokrycie dróg: każda kończy się kopertą


@pytest.mark.parametrize(
    ("argumenty", "status", "kod"),
    [
        (["9621Z"], "ok", 0),  # kod z PKD 2025
        (["9602Z"], "ok", 0),  # kod z 2007 — czytany z tablicy przejścia
        (["fryzjer"], "ok", 0),  # fraza
        (["9999Z"], "brak_trafien", EXIT_PUSTO),  # kod spoza obu roczników
        (["czegoś takiego nie ma"], "brak_trafien", EXIT_PUSTO),  # fraza bez trafień
        (["produkcja"], "ok", 0),  # fraza ponad sufitem wierszy
        (["produkcja", "--wszystkie"], "ok", 0),  # ta sama, bez sufitu
    ],
)
def test_every_terminal_path_ends_with_a_parseable_envelope(
    runner: CliRunner, argumenty: list[str], status: str, kod: int
) -> None:
    """Każda droga wyjścia pisze **dokładnie jedną** kopertę, a kod w niej jest kodem procesu.

    Puste stdout przy kodzie 0 jest dla wołającego nie do odróżnienia od „nic nie znalazłem",
    więc ścieżka bez koperty byłaby defektem cichym. Drugie porównanie — koperta kontra
    `exit_code` — pilnuje rozjazdu między liczbą wydrukowaną a rzeczywistą; są liczone raz,
    ale dopiero ta asercja czyni to twierdzeniem.
    """
    wynik = runner.invoke(app, ["szukaj-pkd", *argumenty, "--wynik", "json"])

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == status
    assert koperta["kod_wyjscia"] == wynik.exit_code == kod


def test_the_screen_and_the_envelope_travel_on_different_streams(runner: CliRunner) -> None:
    """Nic nie jest wyciszone: agent dostaje całą ludzką narrację, tylko na stderr."""
    wynik = runner.invoke(app, ["szukaj-pkd", "9621Z", "--wynik", "json"])

    assert json.loads(wynik.stdout)["polecenie"] == "szukaj-pkd"
    assert "9621Z" in wynik.stderr
    assert "jeden rocznik" in wynik.stderr


def test_without_the_flag_stdout_stays_empty_and_zero_hits_exit_zero(runner: CliRunner) -> None:
    """Decyzja 3B: kod 4 obowiązuje **tylko** pod flagą, bo flaga jest deklaracją wołającego.

    Harmonogram, który dziś czyta niezerowy kod jako „sprawdź, co się stało", zacząłby bez
    tego alarmować w dniu, w którym po prostu nic nie pasowało.
    """
    wynik = runner.invoke(app, ["szukaj-pkd", "9999Z"])

    assert wynik.exit_code == 0
    assert wynik.stdout == ""
    assert "9999Z" in wynik.stderr


def test_an_unknown_wynik_value_is_refused_without_an_envelope(runner: CliRunner) -> None:
    """Odmowa pada, zanim wiadomo, o jakie wyjście chodziło — więc nie ma czego wypisywać.

    Ciche zostanie przy ekranie przy `--wynik jsonl` jest dokładnie tym, przez co wołający
    kończy na parsowaniu ramki `rich`: prosił o dokument i nie dostał ani jednego znaku,
    który by mu powiedział, że go nie będzie.
    """
    wynik = runner.invoke(app, ["szukaj-pkd", "9621Z", "--wynik", "jsonl"])

    assert wynik.exit_code == EXIT_CONFIG
    assert wynik.stdout == ""
    assert "--wynik" in wynik.stderr


# --------------------------------------------- klasyfikacja zakończeń


def test_an_exception_outside_the_taxonomy_still_writes_an_envelope_and_is_re_raised(
    runner: CliRunner, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Test `finally` — jedyny, który odróżnia go od `except CeidgError`.

    Na każdej drodze objętej taksonomią oba rozwiązania zachowują się identycznie. Różnią się
    tam, gdzie wołający jest najbardziej bezradny: przy wyjątku, którego nikt nie przewidział.

    Trzecia asercja jest tą, która naprawdę coś wnosi. Bez niej menedżer mógłby wypisać
    kopertę i **połknąć** wyjątek, a wtedy ślad stosu nie dochodzi do nikogo — czyli program
    milczy o tym, czego sam nie rozumie.
    """
    monkeypatch.setattr(cli, "szukaj", lambda *a, **k: 1 / 0)

    wynik = runner.invoke(app, ["szukaj-pkd", "9621Z", "--wynik", "json"])

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == "blad"
    assert koperta["blad"]["typ"] == "ZeroDivisionError"
    assert isinstance(wynik.exception, ZeroDivisionError)


def test_a_typer_exit_is_not_an_error(monkeypatch: pytest.MonkeyPatch) -> None:
    """Polecenia podnoszą `typer.Exit` na drogach **udanych** (np. `kreator`).

    `finally` zamieniający `raise typer.Exit(code=0)` w `blad` to defekt, który sam się prosi,
    i nie miałby żadnego obserwatora poza tym: koperta byłaby poprawnym JSON-em mówiącym
    nieprawdę o zakończeniu, które się powiodło.
    """
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        cli.WynikPolecenia("pobierz", json=True),
    ):
        raise typer.Exit(code=0)

    assert json.loads(stdout.getvalue())["status"] == "ok"


def test_a_typer_exit_takes_the_envelopes_code_not_its_own(monkeypatch: pytest.MonkeyPatch) -> None:
    """Kod jest liczony, wpisany do koperty i **dopiero potem** podniesiony.

    Polecenie kończące zerem po drodze bez trafień ma pod flagą wyjść czwórką, bo to koperta
    orzeka o zakończeniu. Dwa niezależne źródła kodu wyjścia to dwa źródła, które rozjadą się
    przy pierwszej poprawce jednego z nich.
    """
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(typer.Exit) as podniesione,
        cli.WynikPolecenia("pobierz", json=True) as koperta,
    ):
        koperta.ustaw(Wynik(polecenie="pobierz", status="brak_trafien"))
        raise typer.Exit(code=0)

    assert podniesione.value.exit_code == EXIT_PUSTO
    assert json.loads(stdout.getvalue())["kod_wyjscia"] == EXIT_PUSTO


def test_a_ceidg_error_keeps_its_own_exit_code(monkeypatch: pytest.MonkeyPatch) -> None:
    """Taksonomia `errors.py` zostaje nietknięta — koperta ją cytuje, nie przepisuje."""
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(typer.Exit) as podniesione,
        cli.WynikPolecenia("pobierz", json=True),
    ):
        raise AuthError("token odrzucony")

    assert podniesione.value.exit_code == EXIT_CONFIG
    koperta = json.loads(stdout.getvalue())
    assert koperta["blad"] == {"typ": "AuthError", "komunikat": "token odrzucony"}


def test_an_error_keeps_what_the_command_already_paid_for(monkeypatch: pytest.MonkeyPatch) -> None:
    """Przebieg przerwany po trzydziestu żądaniach kosztował trzydzieści żądań.

    Wyzerowanie licznika przy błędzie odbierałoby wołającemu dokładnie tę informację, na
    której opiera decyzję o ponowieniu — a koperta błędu jest jedynym miejscem, gdzie ona
    jeszcze jest.
    """
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(typer.Exit),
        cli.WynikPolecenia("pobierz", json=True) as koperta,
    ):
        koperta.ustaw(
            Wynik(polecenie="pobierz", status="ok", zapytania=30, srodowisko="test", demo=True)
        )
        raise ConfigError("coś się zepsuło")

    zapisana = json.loads(stdout.getvalue())
    assert zapisana["zapytania"] == 30
    assert zapisana["srodowisko"] == "test"
    assert zapisana["demo"] is True


def test_ctrl_c_is_left_to_run_and_writes_no_envelope(monkeypatch: pytest.MonkeyPatch) -> None:
    """130 należy do `run()` poza `app()` i **nie** jest kodem z taksonomii koperty.

    Wymyślenie dla niego statusu wstawiłoby do `kod_wyjscia` liczbę, której funkcja
    `kod_wyjscia` nie potrafi wyprodukować — czyli ten sam rozjazd, któremu ma zapobiegać.
    Ctrl+C jest decyzją spoza programu, więc program go nie opisuje, tylko przepuszcza.
    """
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(KeyboardInterrupt),
        cli.WynikPolecenia("pobierz", json=True),
    ):
        raise KeyboardInterrupt

    assert stdout.getvalue() == ""


def test_without_the_flag_nothing_is_written_on_any_path(monkeypatch: pytest.MonkeyPatch) -> None:
    """Menedżer obsługuje błąd zawsze, a kopertę pisze tylko na prośbę.

    Dzięki temu droga ludzka nie ma osobnej gałęzi obsługi wyjątków — a dwie gałęzie to ten
    sam kształt, za który ten projekt zapłacił już przy `richtext.safe` (ADR-0009).
    """
    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(typer.Exit) as podniesione,
        cli.WynikPolecenia("pobierz", json=False),
    ):
        raise AuthError("token odrzucony")

    assert podniesione.value.exit_code == EXIT_CONFIG
    assert stdout.getvalue() == ""


# --------------------------------------------- która komenda umie kopertę (ADR-0024, decyzja 6)


def test_the_envelope_table_covers_every_registered_command() -> None:
    """Połowa kompletności: polecenie, które w tabeli nie stoi, ma być widoczne.

    Sama tabela porównana z rzeczywistością nie widzi polecenia, którego w niej nie ma — a to
    jest właśnie ten przypadek, w którym nowe polecenie po cichu odziedziczyłoby „nie umie"
    i odmawiało flagi, której nikt świadomie nie odmówił. Ta sama sztuczka co przy
    `BUDUJE_ASYSTENTA` (ADR-0025) i `POLA_BEZ_WLASNEJ_FLAGI_LISTOWEJ`.
    """
    assert set(KOPERTA_OBSLUGIWANA) == zarejestrowane_polecenia()


@pytest.mark.parametrize(
    "polecenie",
    # `token` jest grupą, nie poleceniem — flagę noszą jej podpolecenia, więc `token --help`
    # pokazuje spis, a nie opcje. Wyliczone osobno niżej, żeby ta lista nie udawała, że grupa
    # jest wyjątkiem od reguły.
    sorted(set(KOPERTA_OBSLUGIWANA) - {"token"}),
)
def test_every_command_knows_the_flag_even_when_it_refuses_it(
    runner: CliRunner, polecenie: str
) -> None:
    """Flaga zadeklarowana **wszędzie**, żeby odmowa była nasza.

    Gdyby `--wynik` stało tylko przy siedmiu obsługiwanych, `kreator --wynik json` kończyłby
    się angielskim „No such option" i kodem 2 od typera — czyli komunikatem o składni tam,
    gdzie chodzi o kontrakt. Sprawdzane po `--help`, bo to jest ta powierzchnia, którą
    wołający czyta, zanim spróbuje.
    """
    pomoc = runner.invoke(app, [polecenie, "--help"])

    assert "--wynik" in pomoc.output.replace("\n", " "), polecenie


def test_a_command_without_an_envelope_refuses_in_polish_with_code_three(
    runner: CliRunner,
) -> None:
    """Kreator jest rozmową — maszynowy wołający nie ma się czym w nim posłużyć."""
    wynik = runner.invoke(app, ["kreator", "--wynik", "json"])

    assert wynik.exit_code == EXIT_CONFIG
    assert "nie wypisuje koperty JSON" in wynik.stderr
    assert wynik.stdout == ""


def test_the_envelope_needs_explicit_yes(runner: CliRunner, env: dict[str, str]) -> None:
    """Dorozumienie `--tak` byłoby odpowiadaniem za wołającego na pytania, o których nie wie.

    Pytania z `safe_default=False` — zgoda na produkcję, start powyżej progu, poszerzenie przy
    zerze — istnieją po to, żeby harmonogram ich **nie** podjął (ADR-0017). Flaga wyjścia nie
    ma prawa ich rozstrzygać bokiem: to nie jest ta sama decyzja, choć pada w tym samym
    wierszu poleceń.
    """
    wynik = runner.invoke(
        app, ["pobierz", "--demo", "-w", "wielkopolskie", "--wynik", "json"], env=env
    )

    assert wynik.exit_code == EXIT_CONFIG
    assert "--tak" in wynik.stderr
    assert wynik.stdout == ""


def test_a_command_that_asks_nothing_does_not_demand_yes(runner: CliRunner) -> None:
    """`runy` i `szukaj-pkd` o nic nie pytają, więc żądanie `--tak` byłoby ceremonią.

    Odmowa, która nie broni żadnej decyzji, uczy obchodzić odmowy — a `--tak` dopisane
    odruchowo do każdego wywołania przestaje cokolwiek znaczyć tam, gdzie znaczy.
    """
    assert runner.invoke(app, ["szukaj-pkd", "9621Z", "--wynik", "json"]).exit_code == 0


# --------------------------------------------- pokrycie: pobranie, od końca do końca


def test_zero_hits_exit_four_only_under_the_flag(runner: CliRunner, env: dict[str, str]) -> None:
    """Decyzja 3B w jednym teście: ta sama droga, dwa kody, różnica wyłącznie we fladze.

    Odwrotność jest tym, co właściciel odrzucił: harmonogram czytający niezerowy kod jako
    „sprawdź, co się stało" zacząłby alarmować w dniu, w którym rejestr zwyczajnie nie miał
    nic pasującego.
    """
    argumenty = ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--pkd", "9999Z"]

    bez_flagi = runner.invoke(app, argumenty, env=env)
    z_flaga = runner.invoke(app, [*argumenty, "--wynik", "json"], env=env)

    assert bez_flagi.exit_code == 0
    assert bez_flagi.stdout == ""
    assert z_flaga.exit_code == EXIT_PUSTO
    assert json.loads(z_flaga.stdout)["status"] == "brak_trafien"


def test_the_envelope_describes_the_workbook_that_is_really_on_disk(
    runner: CliRunner, env: dict[str, str]
) -> None:
    """Anty-rozjazd mierzony **rzeczywistością**, nie drugim renderowaniem.

    Porównanie koperty z ekranem sprawdziłoby tylko, że dwa napisy pochodzą z jednej zmiennej.
    Plik na dysku i liczba wierszy w arkuszu `Firmy` są od obu niezależne, więc dopiero one
    czynią z „ekran i koperta są dwoma renderowaniami tych samych wartości" twierdzeniem.

    Ścieżka w kopercie jest przy okazji tą, której ekran **nie** umie pokazać w całości:
    `summary_table` zawija ją w komórce tabeli (ADR-0024, pierwszy z czterech zmierzonych
    defektów), więc jest to dokładnie ten fakt, dla którego koperta powstała.
    """
    wynik = runner.invoke(
        app,
        ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--maks", "5", "--wynik", "json"],
        env=env,
    )

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == "ok"
    sciezki = koperta["pliki"]
    assert len(sciezki) == 1

    skoroszyt = Path(sciezki[0])
    assert skoroszyt.is_file()
    # Bez `read_only=True`: w tym trybie `max_row` zwraca `None`, dopóki arkusz nie zostanie
    # przejrzany — czyli dokładnie tę liczbę, której ten test szuka.
    arkusz = load_workbook(skoroszyt)["Firmy"]
    assert arkusz.max_row - 1 == koperta["rekordy"]
    assert len(koperta["run_ids"]) == 1
    assert koperta["kryteria"]["wojewodztwo"] == ["wielkopolskie"]


def test_the_demo_marker_reaches_the_channel_nobody_re_reads(
    runner: CliRunner, env: dict[str, str]
) -> None:
    """Znacznik szósty od końca do końca (ADR-0014 + ADR-0024).

    Pozostałe pięć opisują ekran, skoroszyt i nazwę pliku. Za agentem nikt nie ogląda
    pierwszego ekranu, więc koperta jest jedynym miejscem, w którym może się dowiedzieć,
    że dane są wymyślone.
    """
    wynik = runner.invoke(
        app, ["sprawdz-nip", "3563457932", "--demo", "--tak", "--wynik", "json"], env=env
    )

    assert json.loads(wynik.stdout)["demo"] is True


def test_the_help_states_whose_data_the_envelope_carries(runner: CliRunner) -> None:
    """Granica retencji ma stać tam, gdzie wołający ją przeczyta, a nie tylko w ADR.

    `sprawdz-nip` to jedyne polecenie, którego wynikiem jest karta jednej osoby. Narzędzie
    nie zapisuje koperty na dysk, więc przechowanie jej należy do wołającego — i to jest
    zdanie, które musi paść przed pierwszym wywołaniem, nie po nim.
    """
    pomoc = runner.invoke(app, ["sprawdz-nip", "--help"])

    assert "dane jednej osoby" in " ".join(pomoc.output.split())


# --------------------------------------------- licznik żądań: jeden, ten sam co w oznakach życia


def test_the_envelope_reads_the_counter_that_drove_the_signs_of_life(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`LicznikZadan` istnieje wyłącznie po to, żeby te dwie liczby były jedną liczbą.

    Do 2026-09-24 nie miało to żadnego obserwatora: jedyny test nazywający to pole ustawiał
    `zapytania=30` ręcznie na `Wynik` i nie podstawiał licznika, więc gałąź `replace(wynik,
    zapytania=...)` w ogóle się nie wykonywała — test był zielony, bo kopia dataklasy zachowuje
    pole. Usunięcie `koperta.licz_zadania(events)` ze wszystkich pięciu wywołań zostawiało
    suitę zieloną i każdą kopertę z `"zapytania": 0`.
    """
    licznik = LineEvents(Console(file=io.StringIO(), width=120, force_terminal=False))

    with (
        przechwycony_stdout(monkeypatch) as stdout,
        cli.WynikPolecenia("pobierz", json=True) as koperta,
    ):
        koperta.licz_zadania(licznik)
        for _ in range(3):
            licznik.on_request("firmy", 200, 0.1)
        koperta.ustaw(Wynik(polecenie="pobierz", status="ok"))

    assert json.loads(stdout.getvalue())["zapytania"] == 3


def test_an_interrupted_run_still_reports_what_it_spent(monkeypatch: pytest.MonkeyPatch) -> None:
    """Przebieg przerwany po trzech żądaniach kosztował trzy — i to **bez** `ustaw`.

    Odczyt licznika przy wyjściu, a nie przy `ustaw`, jest tu całą treścią: awaria zwykle
    wypada zanim polecenie zdąży cokolwiek ustawić, a wołający opiera na tej liczbie decyzję
    o ponowieniu.
    """
    licznik = LineEvents(Console(file=io.StringIO(), width=120, force_terminal=False))

    with (
        przechwycony_stdout(monkeypatch) as stdout,
        pytest.raises(typer.Exit),
        cli.WynikPolecenia("pobierz", json=True) as koperta,
    ):
        koperta.licz_zadania(licznik)
        for _ in range(3):
            licznik.on_request("firmy", 200, 0.1)
        raise AuthError("token odrzucony w połowie")

    zapisana = json.loads(stdout.getvalue())
    assert zapisana["status"] == "blad"
    assert zapisana["zapytania"] == 3


def test_a_real_run_reports_more_than_zero_requests(runner: CliRunner, env: dict[str, str]) -> None:
    """Obserwator dla **wywołań** `licz_zadania`, nie dla samego menedżera.

    Testy wyżej sprawdzają mechanizm; ten sprawdza, czy ktokolwiek go podłączył. Bez liczby
    dokładnej, bo zależy od korpusu pokazu i byłaby asercją o atrapie rejestru — a mutacja,
    której to pilnuje, zeruje licznik, nie przesuwa go o jeden.
    """
    wynik = runner.invoke(
        app,
        ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--maks", "5", "--wynik", "json"],
        env=env,
    )

    assert json.loads(wynik.stdout)["zapytania"] > 0


# --------------------------------------------- cztery polecenia, które kopertę tylko deklarowały


@pytest.mark.parametrize(
    ("argumenty", "status", "kod"),
    [
        # Pusta baza: nie ma czego wznawiać i nie ma żadnego pobrania do pokazania. To jest
        # `nic_do_zrobienia`, czyli jedyny status, którego żaden test nie oglądał od strony
        # polecenia — dwa istniejące końce z kodem 4 szły przez `brak_trafien`.
        (["wznow", "--demo", "--tak"], "nic_do_zrobienia", EXIT_PUSTO),
        (["runy", "--demo"], "nic_do_zrobienia", EXIT_PUSTO),
        (["aktualizuj", "--demo", "--tak"], "ok", 0),
    ],
)
def test_the_other_commands_really_write_an_envelope(
    runner: CliRunner, env: dict[str, str], argumenty: list[str], status: str, kod: int
) -> None:
    """Cztery z siedmiu poleceń były pokryte wyłącznie tekstem `--help` i tabelą kontra tabelą.

    Żadne z tych sprawdzeń nie **uruchamia** polecenia, więc zlanie wszystkich czterech
    klasyfikacji `nic_do_zrobienia`/`przerwano` do `"ok"` zostawiało suitę zieloną (zmierzone
    2026-09-24). Deklaracja w tabeli nie jest dowodem, że koperta powstaje.
    """
    wynik = runner.invoke(app, [*argumenty, "--wynik", "json"], env=env)

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == status
    assert koperta["kod_wyjscia"] == wynik.exit_code == kod


@pytest.mark.parametrize(
    "argumenty",
    [
        ["wznow", "--demo", "--tak"],
        ["runy", "--demo"],
        ["aktualizuj", "--demo", "--tak"],
        ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--maks", "1"],
        ["sprawdz-nip", "3563457932", "--demo", "--tak"],
    ],
)
def test_the_demo_marker_is_on_every_envelope_a_demo_run_writes(
    runner: CliRunner, env: dict[str, str], argumenty: list[str]
) -> None:
    """Znacznik szósty (ADR-0014) na **każdej** drodze, która go wystawia.

    Dotąd pilnowały go dwa polecenia i trzy asercje o samej dataklasie — pętla po nazwach
    poleceń, która budowała `Wynik(polecenie=…)` wprost, więc nazwa była ozdobą. Ustawienie
    `demo=False` w `runy`, `wznow` i `aktualizuj` przechodziło (zmierzone 2026-09-24).
    Znaczniki ADR-0014 obowiązują **łącznie**, a ten jest na kanale bez człowieka po drugiej
    stronie.
    """
    wynik = runner.invoke(app, [*argumenty, "--wynik", "json"], env=env)

    assert json.loads(wynik.stdout)["demo"] is True


def test_run_notes_reach_the_envelope(
    runner: CliRunner, env: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    """`uwagi=result.notes` — podmienione na `()` przechodziło całą suitę.

    Uwagi przebiegu (niedokończone szczegóły, niedobór przy podziale na partie) nie dają się
    wywołać w pokazie, więc świadek wstrzykuje je w prawdziwy przebieg zamiast go udawać:
    `flow.execute` liczy, co ma policzyć, a test dokłada uwagę do zwróconej wartości.
    """
    prawdziwe = flow.execute

    def z_uwaga(*args: object, **kwargs: object) -> object:
        return replace(prawdziwe(*args, **kwargs), notes=("Coś się nie domknęło.",))  # type: ignore[arg-type]

    monkeypatch.setattr(flow, "execute", z_uwaga)

    wynik = runner.invoke(
        app,
        ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--maks", "1", "--wynik", "json"],
        env=env,
    )

    assert json.loads(wynik.stdout)["uwagi"] == [{"kod": "", "tekst": "Coś się nie domknęło."}]


# --------------------------------------------- `przerwano` pod flagą: dlaczego go tam nie ma


def test_under_yes_giving_up_after_the_cost_table_is_not_an_option() -> None:
    """Mechanizm, na którym stoi nieosiągalność `przerwano` pod `--wynik json`.

    Flaga wymaga `--tak`, a `--tak` odpowiada na każde pytanie jego domyślną. Pytanie „co
    dalej" domyśla się `lista` albo `szczegoly` — nigdy `wyjdz`. Rezygnacja przy zerze trafień
    jest wtedy `brak_trafien` (zmierzone `count == 0`), a rezygnacja powyżej progu podnosi
    `ConfigError` (`safe_default=False`, ADR-0017). Koperty z `przerwano` nie ma dziś jak
    wystawić, i lepiej, żeby wiedział o tym test niż żeby czytelnik zakładał odwrotnie.

    Gdyby `--wynik json` kiedyś przestało wymagać `--tak`, ten test zapali się jako pierwszy —
    i wtedy trzeba będzie sprawdzić, czy `plan.count == 0` nadal znaczy „zapytanie poszło
    i nie pasowało nic". Trzy inne drogi zwracają `count=0` jako wypełniacz.
    """
    prompter = prompts.make_prompter(yes=True, interactive=False)
    pytanie = flow._co_dalej_default(Criteria(wojewodztwo=("wielkopolskie",)))

    assert prompter.ask(pytanie) not in ("wyjdz", "popraw")


# --------------------------------------------- kształt na drucie


def test_the_wire_shape_is_pinned_where_a_change_has_to_be_noticed() -> None:
    """Lista klucze-na-drucie **poza** modułem, który je produkuje.

    `test_the_shared_keys_are_the_ones_the_envelope_always_writes` porównuje dwie listy
    mieszkające w `ui/wynik.py` — przydatne (łapie klucz dodany do `koperta()` i pominięty
    w `KLUCZE_WSPOLNE`), ale nie jest twierdzeniem o kontrakcie, bo obie strony zmienia się
    jednym ruchem. Ta lista jest tutaj, więc usunięcie albo przemianowanie pola wymaga
    dotknięcia testu — i wtedy pod ręką jest zdanie o `WERSJA_KOPERTY`.

    **Zmieniasz ten zbiór przez usunięcie albo zmianę nazwy pola? Podnieś `WERSJA_KOPERTY`.**
    Dołożenie klucza wersji nie zmienia: nie psuje niczyjego odczytu.
    """
    na_drucie = {
        "wersja",
        "polecenie",
        "status",
        "kod_wyjscia",
        "demo",
        "zapytania",
        "uwagi",
        "srodowisko",
        "kryteria",
        "run_ids",
        "rekordy",
        "pliki",
        "blad",
    }

    assert na_drucie == KLUCZE_WSPOLNE
    assert WERSJA_KOPERTY == 1


def test_the_command_registry_reads_the_name_the_app_accepts() -> None:
    """Generator, któremu ufają **obie** tabele kompletności — i który nikogo nie pilnował.

    `zarejestrowane_polecenia` liczyła nazwę z nazwy funkcji, ignorując `CommandInfo.name`.
    Na prawdziwym `app` obie formy dziś się pokrywają (trzy polecenia mają jawną nazwę i każda
    równa się nazwie funkcji z podkreśleniami zamienionymi na myślniki), więc mutacja cofająca
    poprawkę przechodziła całą suitę — dlatego świadek buduje własną aplikację, w której się
    różnią.

    Stawka jest konkretna: `@app.command("pobierz-wszystko")` nad `def pobierz_all` wstawiłby
    do zbioru `"pobierz-all"`, więc `KOPERTA_OBSLUGIWANA` i `BUDUJE_ASYSTENTA` sprawdzałyby się
    względem nazwy, której `app` nie przyjmuje, a prawdziwe polecenie odziedziczyłoby domyślną
    po cichu. To jest ten sam kształt co ręcznie napisany `rokPkd` i anonimizator zamieniający
    identyfikatory na wielkie litery: dowód sanityzowany z tego, co miał wykazać.
    """
    atrapa = typer.Typer()

    @atrapa.command("pobierz-wszystko")
    def pobierz_all() -> None: ...

    assert zarejestrowane_polecenia(atrapa) == {"pobierz-wszystko"}


# --------------------------------------------- bramka, która przeszła nic nie sprawdzając

# NIP z korpusu pokazu (`demo/korpus.py`, pierwszy wpis). Nośny jest fakt, że **trafia**:
# do 2026-09-24 test tej drogi używał NIP-u spoza korpusu, więc `firma` było `null` i żaden
# rekord nigdy nie przeszedł przez serializację.
NIP_Z_KORPUSU = "9995237548"


def test_a_record_really_serialises(runner: CliRunner, env: dict[str, str]) -> None:
    """`sprawdz-nip --wynik json` wywracało się na **każdym trafieniu**, a bramka tego nie widziała.

    Wiersz `Firmy` niesie `datetime.date` (`normalizer._to_date` dla pięciu pól, w tym
    `data_rozpoczecia`, obecnego praktycznie w każdym rekordzie). `zamaskuj` z rozmysłu zostawia
    nieznane typy w spokoju, więc `json.dumps` podnosiło `TypeError`: pusty stdout i kod 1,
    czyli zakończenie, które ADR-0024 nazywa najgorszym możliwym.

    Bramka „każde obsługiwane polecenie uruchomione pod flagą i przepuszczone przez parser"
    była **spełniona** — bo użyty NIP był spoza korpusu pokazu i koperta niosła `"firma": null`.
    Dowód sanityzowany z dokładnie tej własności, którą miał wykazać; ten sam kształt co
    anonimizator zamieniający identyfikatory na wielkie litery i ręcznie napisany `rokPkd`.
    Dlatego ten test pilnuje **wartości daty**, a nie tego, że koperta się parsuje.
    """
    wynik = runner.invoke(
        app, ["sprawdz-nip", NIP_Z_KORPUSU, "--demo", "--tak", "--wynik", "json"], env=env
    )

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == "ok"
    assert koperta["kod_wyjscia"] == wynik.exit_code == 0
    assert isinstance(koperta["firma"]["data_rozpoczecia"], str)
    date.fromisoformat(koperta["firma"]["data_rozpoczecia"])


def test_an_envelope_written_before_anything_succeeds_still_says_it_is_a_demo(
    runner: CliRunner, env: dict[str, str]
) -> None:
    """Znacznik szósty na drodze **błędu** — tam, gdzie polecenie nie zdążyło nic ustawić.

    `--format xml` pada jako `ConfigError` przed pobraniem, więc koperta powstaje wyłącznie
    z tego, co wie menedżer. Brał wtedy `demo` z domyślnej dataklasy i meldował `false` na
    przebiegu z pokazu — a agent przypisujący nieudany pokaz produkcji to jest dokładnie ta
    pomyłka, przed którą ADR-0014 stawia pięć znaczników **łącznie**.
    """
    wynik = runner.invoke(
        app,
        ["pobierz", "--demo", "--tak", "-w", "wielkopolskie", "--format", "xml", "--wynik", "json"],
        env=env,
    )

    koperta = json.loads(wynik.stdout)
    assert koperta["status"] == "blad"
    assert koperta["demo"] is True
