"""Scenariusz 4 (UZUPELNIENIE_01 §D): fixtures z nazwami `=CMD()`, `+1`, `@SUM`, `-2+3`,
znakami sterującymi i 5 000 znaków — skoroszyt otwiera się, komórki są tekstem,
formuły się nie wykonują. Uruchamiany w CI przy każdym przebiegu."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from ceidg_tool.exporter import write_csv, write_jsonl, write_workbook
from ceidg_tool.normalizer import normalize
from ceidg_tool.records import RawRecord, RowContext
from ceidg_tool.safetext import FORMULA_PREFIXES
from tests.conftest import detail_record, list_record

CTX = RowContext(srodowisko="test", pobrano_utc="2026-09-05T10:00:00Z")
HOSTILE = [
    "=CMD()",
    "+1",
    "@SUM(A1)",
    "-2+3",
    '=HYPERLINK("http://evil","klik")',
    "=cmd|'/c calc'!A1",
    "\tTAB",
    "\rCR",
    "ACME\x01\x0b\x7fSp. z o.o.",
    "x" * 5000,
]


def hostile_records() -> list:  # type: ignore[type-arg]
    out = []
    for i, name in enumerate(HOSTILE, start=1):
        raw = RawRecord(
            id=list_record(i)["id"],
            list_json=list_record(i, nazwa=name),
            detail_json=detail_record(
                i,
                nazwa=name,
                email=name + "@x.test",
                telefon=name,
                pkd=[{"kod": name[:5], "nazwa": name}],
                pkdGlowny={"kod": name[:5], "nazwa": name},
                adresKorespondencyjny={
                    "ulica": name,
                    "budynek": "1",
                    "miasto": name,
                    "kod": "00-000",
                },
                spolki=[{"nip": name, "regon": name}],
                obywatelstwa=[{"symbol": name, "kraj": name}],
            ),
            list_utc="2026-09-05T09:00:00Z",
            detail_utc="2026-09-05T09:30:00Z",
            detail_state="pobrany",
            zrodlo="CEIDG_API",
        )
        out.append(normalize(raw, CTX))
    return out


def _is_safe(value: object) -> bool:
    if not isinstance(value, str):
        return True
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return False
    return all(ch in "\t\n" or ord(ch) >= 32 for ch in value)


def test_the_hostile_fixture_exercises_every_prefix_the_exporter_neutralises() -> None:
    """Lista wrogich napisów wyprowadzona z produkcji, nie spisana z §D raz na zawsze.

    `HOSTILE` i `_is_safe` są niezależnym oraclem — i tak ma zostać, bo wspólna stała
    przepuściłaby błąd w samej stałej. Niezależność ma jednak jedną stronę słabą: gdy do
    `FORMULA_PREFIXES` dojdzie nowy znak, nikt nie doda próbki i scenariusz 4 przejdzie,
    nie sprawdziwszy nowego prefiksu. Ta asercja pyta produkcję, czego pilnuje, i wymusza
    dopisanie próbki — a `_is_safe` niżej dalej sprawdza to po swojemu.
    """
    covered = {name[0] for name in HOSTILE if name}

    assert set(FORMULA_PREFIXES) <= covered, (
        f"brak próbki dla prefiksu {sorted(set(FORMULA_PREFIXES) - covered)!r}"
    )


def test_every_text_cell_in_every_sheet_is_neutralized(tmp_path: Path) -> None:
    records = hostile_records()
    (dest,) = write_workbook(
        tmp_path / "hostile.xlsx",
        lambda: iter(records),
        metadata=[("kryteria", "=1+1"), ("cel_pobrania", "@SUM(A1:A9)")],
    )
    wb = load_workbook(dest)
    assert {"Firmy", "PKD", "Spolki", "Slownik", "Metadane"} <= set(wb.sheetnames)
    checked = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type == "s", (ws.title, cell.coordinate, cell.value[:30])
                    assert _is_safe(cell.value), (ws.title, cell.coordinate, cell.value[:30])
                    checked += 1
    assert checked > 100
    firmy = wb["Firmy"]
    header = {c.value: i + 1 for i, c in enumerate(firmy[1])}
    assert firmy.cell(row=1 + len(HOSTILE), column=header["nazwa"]).value == "x" * 5000


def test_csv_is_neutralized_and_jsonl_keeps_raw_but_valid_json(tmp_path: Path) -> None:
    records = hostile_records()
    for path in write_csv(tmp_path / "csv", lambda: iter(records)):
        with path.open(encoding="utf-8-sig", newline="") as fh:
            for row in list(csv.reader(fh, delimiter=";"))[1:]:
                assert all(_is_safe(v) for v in row), path.name
    # JSONL to format dla automatów, nie arkusz: dane zostają surowe, ale każdy wiersz
    # musi być poprawnym JSON-em (znaki sterujące zaescapowane, nie wstrzyknięte)
    jsonl = write_jsonl(tmp_path / "out.jsonl", lambda: iter(records))
    lines = jsonl.read_text(encoding="utf-8").splitlines()
    assert len(lines) == len(HOSTILE)
    names = [json.loads(line)["nazwa"] for line in lines]
    assert names == HOSTILE
