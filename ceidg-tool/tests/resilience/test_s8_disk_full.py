"""Scenariusz 8 (uzupelnienie-01.md §D): dysk zapełnia się w trakcie eksportu.

Zaliczenie wg §D: brak pliku częściowego, baza nietknięta, czytelny komunikat.

Sprawdzenie wolnego miejsca **przed** zapisem (`check_free_space`) ma już swój test
w `tests/test_exporter.py`. To jest druga połowa scenariusza i trudniejsza: miejsce
starczyło na oszacowanie, a skończyło się w połowie pisania. Wtedy na dysku leży
napoczęty plik tymczasowy i to on nie może zostać ani pod swoją nazwą, ani pod docelową.

Czego ten test **nie** zastępuje: §E chce przebiegu na naprawdę pełnym dysku, z datą
i wynikiem w `docs/resilience-report.md`. Wstrzykujemy `ENOSPC` z kodu, a prawdziwy pełny
wolumin potrafi zwrócić błąd także przy zamykaniu pliku, przy `os.replace` i przy zapisie
dziennika — te ścieżki zostają dla przebiegu ręcznego.
"""

from __future__ import annotations

import errno
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from ceidg_tool.errors import ExportError
from ceidg_tool.pipeline import Deps, build_deps, run_export, run_fetch
from tests.conftest import FakeClock
from tests.support import criteria
from tests.test_pipeline_e2e import UNIQUE_IDS, paged_api, settings_for

PARTIAL_BYTES = b"PK\x03\x04" + b"0" * 4096


def completed_run(tmp_path: Path, clock: FakeClock) -> tuple[Deps, str]:
    """Zakończone pobranie w bazie — punkt wyjścia każdego eksportu."""
    api = paged_api()
    deps = build_deps(settings_for(tmp_path), clock=clock, http=api.client())
    deps.profile = deps.profile.model_copy(update={"max_limit_firmy": 5})
    assert deps.client is not None
    deps.client._profile = deps.profile
    result = run_fetch(criteria(wojewodztwo="podlaskie"), deps, known_count=15)
    assert result.status == "zakonczony"
    return deps, result.run_id


def fill_the_disk_mid_write(monkeypatch: pytest.MonkeyPatch) -> None:
    """Zapis zaczyna się, plik tymczasowy powstaje, po czym na woluminie kończy się miejsce."""
    real_save = Workbook.save

    def out_of_space(self: Workbook, filename: object) -> None:
        # Najpierw prawdziwy zapis do bufora: skoroszyt jest strumieniowy, więc dopiero to
        # domyka generatory arkuszy. Bez tego test zostawiałby ostrzeżenia odśmiecacza,
        # które nie mówią nic o badanym zachowaniu.
        real_save(self, io.BytesIO())
        Path(str(filename)).write_bytes(PARTIAL_BYTES)
        raise OSError(errno.ENOSPC, "No space left on device")

    monkeypatch.setattr(Workbook, "save", out_of_space)


def leftovers(directory: Path) -> list[str]:
    return sorted(p.name for p in directory.iterdir() if p.name.endswith(".tmp"))


def test_a_disk_that_fills_mid_write_leaves_no_partial_file(
    tmp_path: Path, clock: FakeClock, monkeypatch: pytest.MonkeyPatch
) -> None:
    deps, run_id = completed_run(tmp_path, clock)
    dest = tmp_path / "wyniki" / "firmy.xlsx"
    fill_the_disk_mid_write(monkeypatch)

    with pytest.raises(ExportError) as failed:
        run_export(run_id, dest, deps)

    assert str(dest) in str(failed.value)  # komunikat mówi, którego pliku nie udało się zapisać
    assert "No space left" in str(failed.value)
    assert not dest.exists()  # kryterium §D: brak pliku częściowego pod docelową nazwą
    assert leftovers(dest.parent) == []  # ani pod tymczasową
    deps.store.close()


def test_the_database_survives_a_failed_export_and_the_next_one_succeeds(
    tmp_path: Path, clock: FakeClock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """„Baza nietknięta" sprawdzone tak, jak to widzi operator: powtórzonym eksportem.

    Samo policzenie wierszy dowodzi mniej — eksport czyta z bazy, więc dopiero drugi
    przebieg po zwolnieniu miejsca pokazuje, że nie zostało po pierwszym nic, co psuje
    kolejny (otwarta transakcja, przestawiony etap runu, zajęty plik na Windows).
    """
    deps, run_id = completed_run(tmp_path, clock)
    dest = tmp_path / "wyniki" / "firmy.xlsx"
    records_before = deps.store.count_run_records(run_id)
    fill_the_disk_mid_write(monkeypatch)
    with pytest.raises(ExportError):
        run_export(run_id, dest, deps)

    monkeypatch.undo()  # miejsce na dysku wróciło
    summary = run_export(run_id, dest, deps)

    assert deps.store.get_run(run_id).status == "zakonczony"
    assert deps.store.count_run_records(run_id) == records_before == UNIQUE_IDS
    assert summary.paths == (dest,)
    assert load_workbook(dest)["Firmy"].max_row == UNIQUE_IDS + 1
    deps.store.close()


def test_the_export_error_is_a_readable_sentence_not_a_traceback(
    tmp_path: Path, clock: FakeClock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`OSError` z systemu plików ma dojść do operatora jako `ExportError` z kodem wyjścia 1."""
    deps, run_id = completed_run(tmp_path, clock)
    fill_the_disk_mid_write(monkeypatch)

    with pytest.raises(ExportError) as failed:
        run_export(run_id, tmp_path / "wyniki" / "firmy.xlsx", deps)

    assert failed.value.exit_code == 1
    assert str(failed.value).startswith("Nie można zapisać ")
    deps.store.close()
