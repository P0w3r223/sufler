"""Pobranie w partiach (`run_batched_fetch`) i eksport z kilku runów naraz.

ADR-0008, decyzja 3: każda partia to osobny run z własnym checkpointem, rozpoznawany po
odcisku kryteriów. Stąd trzy zachowania, które muszą być pewne: partia już pobrana jest
**pomijana bez żądania**, partia przerwana jest **wznawiana**, a partia wciąż za duża jest
**dzielona drobniej, zanim cokolwiek pobierze**. Na koniec kilka runów daje jeden skoroszyt
bez duplikatów, a różnica między sumą partii a pierwotnym `count` jest raportowana, nie ukryta.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from ceidg_tool.batching import BatchPlan, plan_batches
from ceidg_tool.config import Settings
from ceidg_tool.errors import ConfigError
from ceidg_tool.pipeline import BatchOutcome, Deps, build_deps, run_batched_fetch, run_export
from ceidg_tool.ui import flow
from ceidg_tool.ui.flow import FetchPlan
from tests.conftest import FakeClock, list_record
from tests.support import FakeApi, RecordingView, criteria

BASE = "https://test-dane.biznes.gov.pl/api/ceidg/v3"
TODAY = date(2026, 9, 5)
THRESHOLD = 50_000

# Dwa lata → dwie partie roczne, gdy count przekracza próg dla jednej dekady.
SOURCE = criteria(wojewodztwo="mazowieckie", data_od="2021-01-01", data_do="2022-12-31")


def deps_for(tmp_path: Path, clock: FakeClock, api: FakeApi) -> Deps:
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    return build_deps(settings, clock=clock, http=api.client())


Y2021 = "2021-01-01..2021-12-31"
Y2022 = "2022-01-01..2022-12-31"
Q1_2021 = "2021-01-01..2021-03-31"
JAN_2021 = "2021-01-01..2021-01-31"


def batch_api(
    counts: Mapping[str, int], *, records: Mapping[str, list[str]] | None = None
) -> FakeApi:
    """API sterowane **pełnym zakresem dat** partii: ile trafień i jakie identyfikatory zwrócić.

    Kluczem jest `"dataod..datado"`, bo sam rok nie odróżnia partii rocznej od jej pierwszego
    kwartału i pierwszego miesiąca — wszystkie trzy zaczynają się 1 stycznia. Test opisuje
    scenariusz tabelką, a nie ciągiem warunków.
    """
    api = FakeApi()

    def range_of(request: httpx.Request) -> str:
        return f"{request.url.params.get('dataod', '')}..{request.url.params.get('datado', '')}"

    def fallback(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/firmy" not in url:
            raise AssertionError(f"nieoczekiwane żądanie: {url}")
        key = range_of(request)
        count = counts.get(key, 0)
        if request.url.params.get("limit") == "1":
            return httpx.Response(200, json={"count": count, "firmy": []})
        ids = (records or {}).get(key, [])
        if not ids:
            return httpx.Response(204)
        firmy = [list_record(1, id=rec_id) for rec_id in ids]
        return httpx.Response(200, json={"count": count, "firmy": firmy, "links": {}})

    api.fallback = fallback
    return api


def yearly_plan(count: int) -> BatchPlan:
    """Plan, który dla zakresu 2021–2022 wypada partiami rocznymi (dekada jest za duża)."""
    plan = plan_batches(SOURCE, count, today=TODAY, threshold=THRESHOLD)
    assert [b.label for b in plan.batches] == ["2021", "2022"], plan.granularity
    return plan


def counted_requests(api: FakeApi) -> list[str]:
    return [r for r in api.requests if "limit=1" in r]


# ----------------------------------------------------------------------------- przebieg zwykły


def test_each_batch_is_counted_once_and_then_fetched(tmp_path: Path, clock: FakeClock) -> None:
    """Przed każdą partią jeden `count` — to on decyduje, czy partię pobrać, czy podzielić."""
    api = batch_api({Y2021: 2, Y2022: 1}, records={Y2021: ["a", "b"], Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)
    plan = yearly_plan(100_000)

    result = run_batched_fetch(plan, deps, threshold=THRESHOLD)

    assert len(counted_requests(api)) == 2
    assert [o.status for o in result.outcomes] == ["pobrana", "pobrana"]
    assert result.counted == 3
    assert result.records == 3
    assert len(result.run_ids) == 2
    deps.store.close()


def test_an_empty_batch_costs_only_its_count(tmp_path: Path, clock: FakeClock) -> None:
    """Partia bez trafień nie zasługuje na żądanie o listę — sam `count` wystarczy."""
    api = batch_api({Y2021: 0, Y2022: 1}, records={Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)

    statuses = {o.label: o.status for o in result.outcomes}
    assert statuses["2021"] == "pusta"
    assert statuses["2022"] == "pobrana"
    assert len(result.run_ids) == 1  # pusta partia nie zakłada runu
    deps.store.close()


def test_the_caller_is_notified_about_every_batch_as_it_finishes(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Postęp „partia 2 z 5” wymaga zgłoszenia każdej partii z osobna, w kolejności."""
    api = batch_api({Y2021: 1, Y2022: 1}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)
    seen: list[BatchOutcome] = []

    run_batched_fetch(
        yearly_plan(100_000),
        deps,
        threshold=THRESHOLD,
        on_batch=seen.append,
    )

    assert [o.label for o in seen] == ["2021", "2022"]
    assert all(o.run_id for o in seen)
    deps.store.close()


# --------------------------------------------------------------------- pomijanie i wznawianie


def test_a_batch_already_fetched_is_skipped_without_a_single_request(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Sedno wznawiania planu: powtórzone uruchomienie nie płaci drugi raz za tę samą partię."""
    api = batch_api({Y2021: 1, Y2022: 1}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)
    plan = yearly_plan(100_000)
    run_batched_fetch(plan, deps, threshold=THRESHOLD)
    requests_after_first = len(api.requests)

    second = run_batched_fetch(plan, deps, threshold=THRESHOLD)

    assert [o.status for o in second.outcomes] == ["pominięta (już pobrana)"] * 2
    assert len(api.requests) == requests_after_first  # ani jednego żądania więcej
    assert second.records == 2  # rekordy nadal raportowane z bazy
    deps.store.close()


def test_an_interrupted_batch_is_resumed_rather_than_restarted(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Partia przerwana ma checkpoint — plan ma ją wznowić, a nie zakładać nowego runu."""
    api = batch_api({Y2021: 1, Y2022: 1}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)
    plan = yearly_plan(100_000)
    first_batch = plan.batches[0]
    stale_run = "run-przerwany"
    deps.store.start_run(
        run_id=stale_run,
        criteria_json=first_batch.criteria.model_dump_json(),
        criteria_hash=first_batch.fingerprint(),
        profile_hash=deps.profile.profile_hash(),
        mode="lista",
        tool_version="0.1.0",
        cursor_mode="links",
    )
    deps.store.update_run_status(stale_run, "przerwany")

    result = run_batched_fetch(plan, deps, threshold=THRESHOLD)

    outcomes = {o.label: o for o in result.outcomes}
    assert outcomes["2021"].status == "wznowiona"
    assert outcomes["2021"].run_id == stale_run  # ten sam run, nie nowy
    assert stale_run in result.run_ids
    deps.store.close()


# ----------------------------------------------------------------------------- dalszy podział


def test_a_batch_still_over_the_threshold_is_subdivided_before_any_fetch(
    tmp_path: Path, clock: FakeClock
) -> None:
    """§C: partia ponad progiem nie startuje — najpierw schodzi o poziom drobniej."""
    api = batch_api({Y2021: 60_000, Q1_2021: 10, Y2022: 1}, records={Q1_2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)

    statuses = [o.status for o in result.outcomes]
    assert statuses[0] == "podzielona na 4"  # rok → cztery kwartały
    assert [o.label for o in result.outcomes[1:5]] == [
        "2021-Q1",
        "2021-Q2",
        "2021-Q3",
        "2021-Q4",
    ]
    # partia podzielona nie pobrała ani jednego rekordu przed podziałem
    assert result.outcomes[0].records == 0
    deps.store.close()


def test_subdivision_happens_before_the_remaining_batches(tmp_path: Path, clock: FakeClock) -> None:
    """Drobniejsze partie wchodzą na początek kolejki, żeby zachować porządek chronologiczny."""
    api = batch_api({Y2021: 60_000, Q1_2021: 10, Y2022: 1}, records={Q1_2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)

    labels = [o.label for o in result.outcomes]
    assert labels.index("2021-Q4") < labels.index("2022")
    deps.store.close()


def test_a_month_that_is_still_too_large_is_refused_with_advice(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Dno podziału: program mówi wprost, że datą już nie zawęzi, i podpowiada czym."""
    api = batch_api({JAN_2021: 60_000})
    deps = deps_for(tmp_path, clock, api)
    month = criteria(wojewodztwo="mazowieckie", data_od="2021-01-01", data_do="2021-01-31")
    plan = plan_batches(month, 60_000, today=TODAY, threshold=THRESHOLD)
    assert plan.granularity == "miesiac"

    with pytest.raises(ConfigError) as caught:
        run_batched_fetch(plan, deps, threshold=THRESHOLD)

    message = str(caught.value)
    assert "nie da się jej podzielić" in message
    assert "województwo, PKD, status" in message
    deps.store.close()


# ----------------------------------------------------------------------------- rozliczenie trafień


def test_hits_not_covered_by_any_batch_are_reported_not_hidden(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Różnica między `count` sprzed podziału a sumą partii musi być widoczna.

    Pierwotny `count` mówi 60 000, a partie obejmują 2. Ten docstring twierdził do
    2026-09-09, że reszta „to wpisy bez daty rozpoczęcia" — powtarzając za komunikatem
    przyczynę, której nikt nie zmierzył i której na bazie operatora nie ma (zero takich
    wpisów na 16 310). Prawdziwą przyczyną straty były dokładane granice dat, naprawione
    w ADR-0015. Test pilnuje **widoczności różnicy**, bo tylko to potrafi sprawdzić.
    """
    api = batch_api({Y2021: 1, Y2022: 1}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(60_000), deps, threshold=THRESHOLD)

    assert result.expected == 60_000
    assert result.counted == 2
    assert result.missing == 59_998
    deps.store.close()


def test_no_missing_hits_are_reported_when_the_batches_add_up(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Zgodny rachunek nie może produkować fałszywego ostrzeżenia."""
    api = batch_api({Y2021: 30_000, Y2022: 30_000}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(60_000), deps, threshold=THRESHOLD)

    assert result.counted == result.expected == 60_000
    assert result.missing == 0
    deps.store.close()


def test_more_records_than_expected_never_reports_negative_missing(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Rejestr żyje między zapytaniami — nadwyżka nie może dać ujemnej „różnicy”."""
    api = batch_api({Y2021: 40_000, Y2022: 40_000}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)

    result = run_batched_fetch(yearly_plan(60_000), deps, threshold=THRESHOLD)

    assert result.counted > result.expected
    assert result.missing == 0
    deps.store.close()


# ----------------------------------------------------------------------------- eksport z partii


def test_a_batched_run_exports_into_one_workbook(tmp_path: Path, clock: FakeClock) -> None:
    """ADR-0008: pobranie w partiach ma dać jeden skoroszyt, a nie plik na partię."""
    api = batch_api({Y2021: 2, Y2022: 1}, records={Y2021: ["a", "b"], Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)
    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)
    requests_before_export = len(api.requests)

    summary = run_export(list(result.run_ids), tmp_path / "partie.xlsx", deps)

    assert summary.records == 3
    assert summary.run_ids == result.run_ids
    assert load_workbook(tmp_path / "partie.xlsx")["Firmy"].max_row - 1 == 3
    assert len(api.requests) == requests_before_export  # eksport wyłącznie z bazy
    deps.store.close()


def test_a_company_appearing_in_two_batches_is_exported_once(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Firma może zmienić datę rozpoczęcia między partiami — skoroszyt nie może jej zdublować."""
    api = batch_api({Y2021: 2, Y2022: 2}, records={Y2021: ["a", "b"], Y2022: ["b", "c"]})
    deps = deps_for(tmp_path, clock, api)
    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)

    summary = run_export(list(result.run_ids), tmp_path / "bez_duplikatow.xlsx", deps)

    assert summary.records == 3  # a, b, c — nie 4
    assert load_workbook(tmp_path / "bez_duplikatow.xlsx")["Firmy"].max_row - 1 == 3
    deps.store.close()


def test_the_metadata_sheet_names_every_batch_run(tmp_path: Path, clock: FakeClock) -> None:
    """Skoroszyt musi dać się powiązać z runami, z których powstał — inaczej audyt jest ślepy."""
    api = batch_api({Y2021: 1, Y2022: 1}, records={Y2021: ["a"], Y2022: ["b"]})
    deps = deps_for(tmp_path, clock, api)
    result = run_batched_fetch(yearly_plan(100_000), deps, threshold=THRESHOLD)

    run_export(list(result.run_ids), tmp_path / "meta.xlsx", deps)

    sheet = load_workbook(tmp_path / "meta.xlsx")["Metadane"]
    text = "\n".join(str(row[1].value) for row in sheet.iter_rows(min_row=2))
    for run_id in result.run_ids:
        assert run_id in text
    deps.store.close()


def test_exporting_zero_runs_is_refused(tmp_path: Path, clock: FakeClock) -> None:
    """Pusta lista runów znaczy „nie ma czego eksportować” — lepszy błąd niż pusty plik."""
    api = batch_api({})
    deps = deps_for(tmp_path, clock, api)

    with pytest.raises(Exception, match="przynajmniej jednego pobrania"):
        run_export([], tmp_path / "nic.xlsx", deps)

    deps.store.close()


# ------------------------------------------------------- próg wędruje w planie, nie obok


Q2_2021 = "2021-04-01..2021-06-30"
Q3_2021 = "2021-07-01..2021-09-30"
Q4_2021 = "2021-10-01..2021-12-31"


def test_execute_splits_by_the_threshold_the_operator_was_shown(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Próg jest polem `FetchPlan`, więc wykonanie dzieli partie dokładnie tak, jak
    zapowiedziała tabela. Gdyby `execute` brało próg skądinąd (np. domyślne 50 000),
    partia o trzech trafieniach zostałaby pobrana w całości zamiast podzielona."""
    api = batch_api(
        {Y2021: 3, Y2022: 0, Q1_2021: 2, Q2_2021: 0, Q3_2021: 0, Q4_2021: 1},
        records={Q1_2021: ["a", "b"], Q4_2021: ["c"]},
    )
    deps = deps_for(tmp_path, clock, api)
    plan = FetchPlan(criteria=SOURCE, count=3, threshold=2, batches=yearly_plan(100_000))

    result = flow.execute("partie", plan, deps, RecordingView())

    # 2021 (podzielony) + jego 4 kwartały + 2022 = 6 zapytań o `count`
    assert len(counted_requests(api)) == 6
    assert result.records == 3
    deps.store.close()


def test_the_default_threshold_would_not_have_split_that_batch(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Kontrola do testu wyżej: ten sam plan przy zwykłym progu pobiera partie w całości."""
    api = batch_api({Y2021: 3, Y2022: 0}, records={Y2021: ["a", "b", "c"]})
    deps = deps_for(tmp_path, clock, api)
    plan = FetchPlan(criteria=SOURCE, count=3, threshold=THRESHOLD, batches=yearly_plan(100_000))

    result = flow.execute("partie", plan, deps, RecordingView())

    assert len(counted_requests(api)) == 2
    assert result.records == 3
    deps.store.close()


# ------------------------------------------------------- blokada po ubitym procesie a partie


def stale_lock_of_a_dead_process(deps: Deps, clock: FakeClock) -> None:
    """Wpis blokady po procesie, który jej nie zwolnił — dokładnie to zostawia `kill -9`."""
    deps.store._conn.execute(
        "INSERT INTO run_lock(environment, pid, started_utc, heartbeat_epoch) VALUES (?,?,?,?)",
        ("test", 999_999, "2026-09-06T10:00:00Z", clock.wall()),
    )
    deps.store._conn.commit()


def test_force_reaches_the_first_batch_that_has_never_run(tmp_path: Path, clock: FakeClock) -> None:
    """Poprawka z przeglądu: `force_lock` szło tylko do partii wznawianych.

    Świeży plan nie ma czego wznawiać, więc `pobierz --partie --force` po ubitym procesie
    odbijał się o blokadę tak samo, jakby flagi nie było — a flaga jest jedynym, co po
    awarii odblokowuje pracę przed upływem dziesięciu minut.
    """
    api = batch_api({Y2021: 2, Y2022: 1}, records={Y2021: ["a", "b"], Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)
    stale_lock_of_a_dead_process(deps, clock)  # heartbeat świeży: sama nie wygasnie
    plan = yearly_plan(100_000)

    result = run_batched_fetch(plan, deps, threshold=THRESHOLD, force_lock=True)

    assert result.records == 3
    deps.store.close()


def test_without_force_a_stale_lock_stops_a_batched_fetch(tmp_path: Path, clock: FakeClock) -> None:
    """Kontrola dla testu wyżej: bez flagi ta sama blokada ma zatrzymać pobieranie.

    Bez tej pary poprzedni test przechodziłby też wtedy, gdyby blokada w ogóle nie działała.
    """
    from ceidg_tool.errors import StoreLockedError

    api = batch_api({Y2021: 2, Y2022: 1}, records={Y2021: ["a", "b"], Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)
    stale_lock_of_a_dead_process(deps, clock)
    plan = yearly_plan(100_000)

    with pytest.raises(StoreLockedError):
        run_batched_fetch(plan, deps, threshold=THRESHOLD)
    deps.store.close()


def test_force_is_spent_once_and_not_re_used_for_every_batch(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Wymuszenie ma przejąć martwą blokadę raz, a nie odbierać ją przy każdej partii.

    Partia zwalnia blokadę po sobie, więc gdyby flaga działała przy każdej z nich, dwunasta
    partia odebrałaby ją procesowi, który w międzyczasie wziął ją uczciwie.
    """
    api = batch_api({Y2021: 2, Y2022: 1}, records={Y2021: ["a", "b"], Y2022: ["c"]})
    deps = deps_for(tmp_path, clock, api)
    forced: list[bool] = []
    real_acquire = type(deps.store).acquire_lock

    def spy(self: object, *, force: bool = False) -> bool:
        forced.append(force)
        return bool(real_acquire(self, force=force))  # type: ignore[arg-type]

    deps.store.acquire_lock = spy.__get__(deps.store)  # type: ignore[method-assign]
    plan = yearly_plan(100_000)

    run_batched_fetch(plan, deps, threshold=THRESHOLD, force_lock=True)

    assert forced == [True, False]  # pierwsza partia wymusza, druga już nie
    deps.store.close()


def test_an_open_ended_batch_is_refetched_while_closed_ones_are_skipped(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Kafel otwarty rośnie, kafel zamknięty nie — więc „już pobrana" jest prawdą tylko
    o drugim (ADR-0015, uwaga z przeglądu 2026-09-09).

    Przed otwarciem krawędzi ostatnia partia niosła `data_do = dzisiaj` i jej odcisk zmieniał
    się z dnia na dzień, więc pomijanie jej nie groziło. Otwarcie ustabilizowało odcisk: bez
    tego rozróżnienia powtórzony `pobierz --partie` meldowałby „pominięta" dla **wszystkich**
    partii i nie przyniósł ani jednego nowego rekordu — a ścieżka niepodzielona zawsze pobiera
    od nowa.
    """
    otwarty_ogon = "2020-01-01.."
    dekady = {"..1999-12-31": 1, "2000-01-01..2009-12-31": 1, "2010-01-01..2019-12-31": 1}
    api = batch_api(
        {**dekady, otwarty_ogon: 1},
        records={
            "..1999-12-31": ["a"],
            "2000-01-01..2009-12-31": ["b"],
            "2010-01-01..2019-12-31": ["c"],
            otwarty_ogon: ["d"],
        },
    )
    deps = deps_for(tmp_path, clock, api)
    plan = plan_batches(
        criteria(wojewodztwo="mazowieckie"), 100_000, today=TODAY, threshold=THRESHOLD
    )
    assert [b.otwarty_do for b in plan.batches] == [False, False, False, True]

    run_batched_fetch(plan, deps, threshold=THRESHOLD)
    po_pierwszym = len(api.requests)

    drugi = run_batched_fetch(plan, deps, threshold=THRESHOLD)
    deps.store.close()

    statusy = [o.status for o in drugi.outcomes]
    assert statusy[:3] == ["pominięta (już pobrana)"] * 3, statusy
    assert statusy[3] != "pominięta (już pobrana)", "otwarty ogon musi zostać pobrany ponownie"
    assert len(api.requests) > po_pierwszym, "ogon kosztuje żądania, i o to chodzi"
