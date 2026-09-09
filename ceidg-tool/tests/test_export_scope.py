"""Zakres eksportu: co skoroszyt mówi o sobie i który run w ogóle wybiera (audyt A4, A7).

Trzy defekty jednego kształtu — **opis pierwszej partii podany jako opis całości**:

* **A4** — `Metadane` brały status, czas, `count`, liczbę stron i ostrzeżenie z `parts[0]`,
  podczas gdy `liczba_pobranych_rekordow` liczyła wszystkie partie. Skoroszyt z dwunastu
  partii twierdził o sobie rzeczy nieprawdziwe, a partia przerwana nie zostawiała śladu.
* **A7** — `eksportuj` bez `--run-id` brało najnowszy run **dowolnego** stanu, choć pomoc
  flagi obiecuje ostatni zakończony. Run przerwany ma komplet kolumn i połowę wierszy, więc
  wychodził jako plik kompletny.
* Sąsiedztwo znalezione przy okazji: `list_resumable` odsiewało stany **po** `LIMIT 20`, więc
  przerwany run znikał z menu, gdy tylko powstało dwadzieścia nowszych zakończonych.

Wspólna diagnoza: filtr nałożony na wynik zapytania zamiast na zapytanie, a agregat wzięty
z pierwszego elementu zamiast ze wszystkich.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

import ceidg_tool.config as config
from ceidg_tool import cli
from ceidg_tool.cli import app
from ceidg_tool.config import Settings
from ceidg_tool.pipeline import build_deps, run_export
from ceidg_tool.store import Store
from ceidg_tool.ui import flow
from ceidg_tool.ui.texts import SummaryInput, no_finished_run, summary_notes
from tests.conftest import FakeClock, detail_record, list_record
from tests.support import RecordingView

EXIT_OK = 0


@pytest.fixture
def store(tmp_path: Path, clock: FakeClock) -> Iterator[Store]:
    s = Store(tmp_path / "store-test.sqlite", environment="test", clock=clock)
    yield s
    s.close()


def _pusty_run(store: Store, run_id: str, status: str) -> str:
    store.start_run(
        run_id=run_id,
        criteria_json='{"wojewodztwo":["podlaskie"]}',
        criteria_hash="abc",
        profile_hash="prof",
        mode="lista",
        tool_version="0.1.0",
        cursor_mode="links",
    )
    if status != "w_toku":
        store.update_run_status(run_id, status)
    return run_id


# --------------------------------------------------------------- okno LIMIT (list_resumable)


def test_przerwany_run_poza_oknem_limitu_nadal_jest_do_wznowienia(store: Store) -> None:
    """Rdzeń defektu okna. Dwadzieścia nowszych zakończonych runów wypychało przerwany
    poza `LIMIT 20`, a filtr stanu działał dopiero na tym, co zapytanie zwróciło — więc
    `wznow` meldowało brak czegokolwiek do wznowienia przy żywym checkpoincie w bazie."""
    stary = _pusty_run(store, "run-przerwany", "przerwany")
    for i in range(20):
        _pusty_run(store, f"run-nowszy-{i:02d}", "zakonczony")

    okno = [r.run_id for r in store.list_runs()]
    assert stary not in okno, "kontrola: run musi naprawdę wypaść poza domyślne okno"

    znalezione = [r.run_id for r in store.list_runs(statuses=("przerwany", "w_toku"))]
    assert stary in znalezione


def test_filtr_stanu_nie_przepuszcza_stanow_niepasujacych(store: Store) -> None:
    """Kontrola pozytywna do testu wyżej: gdyby `statuses=` było ignorowane, tamten test
    przechodziłby przy każdej implementacji zwracającej cokolwiek."""
    _pusty_run(store, "run-a", "zakonczony")
    _pusty_run(store, "run-b", "przerwany")

    assert [r.run_id for r in store.list_runs(statuses=("przerwany",))] == ["run-b"]
    assert [r.run_id for r in store.list_runs(statuses=("zakonczony",))] == ["run-a"]
    assert store.list_runs(statuses=()) == []


def test_filtr_rodzaju_nie_miesza_pobran_z_raportami(store: Store) -> None:
    """`kinds=` idzie tą samą drogą co `statuses=`; menu wznowienia pyta o oba naraz."""
    store.start_run(
        run_id="run-raport",
        criteria_json="{}",
        criteria_hash="h",
        profile_hash="p",
        mode="lista",
        tool_version="0.1.0",
        cursor_mode="links",
        kind="raport",
    )
    _pusty_run(store, "run-firmy", "w_toku")

    tylko_firmy = store.list_runs(statuses=("w_toku",), kinds=("firmy",))
    assert [r.run_id for r in tylko_firmy] == ["run-firmy"]


# --------------------------------------------------------------- A4: Metadane opisują całość


def _partia(
    store: Store,
    run_id: str,
    *,
    od: str,
    do: str,
    firmy: range,
    status: str,
    blad: str | None = None,
) -> str:
    """Partia z własnym zakresem dat, własnymi firmami i własnym stanem."""
    store.start_run(
        run_id=run_id,
        criteria_json=f'{{"wojewodztwo":["podlaskie"],"data_od":"{od}","data_do":"{do}"}}',
        criteria_hash=run_id,
        profile_hash="prof",
        mode="szczegoly",
        tool_version="0.1.0",
        cursor_mode="links",
    )
    store.save_page(
        run_id,
        page_index=0,
        records=[list_record(i) for i in firmy],
        next_cursor=None,
    )
    store.save_details(details=[detail_record(i) for i in firmy])
    store.set_run_count(run_id, len(firmy))
    store.update_run_status(run_id, status, blad)
    return run_id


def _metadane(path: Path) -> dict[str, object]:
    wb = load_workbook(path)
    return {r[0].value: r[1].value for r in wb["Metadane"].iter_rows(min_row=2)}


def test_metadane_eksportu_z_partii_opisuja_wszystkie_partie(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Rdzeń A4, sprawdzony na komórkach skoroszytu, nie na wartości zwracanej.

    Druga partia jest przerwana i niesie błąd. Przed poprawką arkusz meldował status,
    czas, `count` i liczbę stron pierwszej partii, a o błędzie drugiej milczał — plik
    twierdził o sobie, że jest kompletnym pobraniem zakończonym."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    pierwsza = _partia(
        deps.store,
        "run-p1",
        od="2020-01-01",
        do="2020-12-31",
        firmy=range(1, 4),
        status="zakonczony",
    )
    druga = _partia(
        deps.store,
        "run-p2",
        od="2021-01-01",
        do="2021-12-31",
        firmy=range(4, 6),
        status="przerwany",
        blad="połączenie zerwane",
    )

    run_export([pierwsza, druga], tmp_path / "partie.xlsx", deps)
    meta = _metadane(tmp_path / "partie.xlsx")
    deps.store.close()

    assert meta["status_runu"] == "mieszany: przerwany×1, zakonczony×1"
    assert meta["liczba_trafien_count"] == 5, "count musi sumować partie, nie brać pierwszej"
    assert meta["liczba_stron"] == 2, "każda partia zapisała jedną stronę"
    assert str(meta["run_id"]) == "run-p1, run-p2"
    # Zakres dat całości, nie pierwszej partii: podział idzie po dacie rozpoczęcia, więc
    # sumą jest kryterium partii pierwszej rozciągnięte na skrajne daty.
    assert "2020-01-01" in str(meta["kryteria"]) and "2021-12-31" in str(meta["kryteria"])
    assert "2020-12-31" not in str(meta["kryteria"])
    # Błąd partii, która padła — z jej identyfikatorem, bo inaczej nie wiadomo której.
    assert "run-p2" in str(meta["ostrzezenie"]) and "połączenie zerwane" in str(meta["ostrzezenie"])
    # Etykiety partii niosą stan, inaczej przerwana wygląda jak każda inna.
    assert "(przerwany)" in str(meta["partie_kryteria"])


def test_metadane_pojedynczego_runu_opisuja_ten_run(tmp_path: Path, clock: FakeClock) -> None:
    """Kontrola pozytywna. Bez niej „agreguj wszystko" mogłoby zepsuć zwykły eksport
    jednego runu — a to jest ścieżka, którą operator chodzi codziennie."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = _partia(
        deps.store,
        "run-solo",
        od="2020-01-01",
        do="2020-12-31",
        firmy=range(1, 4),
        status="zakonczony",
    )

    run_export(run_id, tmp_path / "solo.xlsx", deps)
    meta = _metadane(tmp_path / "solo.xlsx")
    deps.store.close()

    assert meta["status_runu"] == "zakonczony"
    assert meta["liczba_trafien_count"] == 3
    assert meta["run_id"] == "run-solo"
    assert "liczba_partii" not in meta, "jeden run to nie pobranie w partiach"
    assert "ostrzezenie" not in meta


def test_nieznany_count_partii_nie_udaje_zmierzonego(tmp_path: Path, clock: FakeClock) -> None:
    """`count` partii, która padła przed pierwszą odpowiedzią, jest **nieznany**.

    Sama suma po znanych wyglądałaby jak pomiar całości — czyli byłaby tym samym
    przekłamaniem co A4, tylko o poziom rzadszym."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    znana = _partia(
        deps.store,
        "run-znana",
        od="2020-01-01",
        do="2020-12-31",
        firmy=range(1, 4),
        status="zakonczony",
    )
    nieznana = _partia(
        deps.store,
        "run-nieznana",
        od="2021-01-01",
        do="2021-12-31",
        firmy=range(4, 6),
        status="przerwany",
    )
    deps.store.set_run_count(nieznana, None)

    run_export([znana, nieznana], tmp_path / "mieszane.xlsx", deps)
    meta = _metadane(tmp_path / "mieszane.xlsx")
    deps.store.close()

    assert "znany dla 1 z 2" in str(meta["liczba_trafien_count"])


# --------------------------------------------------------------- A7: który run i co o nim mówimy


def test_podsumowanie_mowi_ze_pobranie_nie_jest_zakonczone() -> None:
    """Ekran, nie tylko arkusz. Operator bez wiedzy o API otwiera przede wszystkim ekran,
    a `Metadane` bywa arkuszem, do którego nigdy nie dojdzie."""
    uwagi = summary_notes(
        SummaryInput(
            paths=(),
            records=3,
            by_status={},
            with_phone=0,
            with_email=0,
            sheets=(),
            kind="firmy",
            run_ids=("run-a",),
            log_path=Path("log"),
            statuses=("przerwany",),
        )
    )
    assert any("nie jest zakończone" in u and "wznow" in u for u in uwagi)


def test_podsumowanie_milczy_o_kompletnosci_gdy_wszystko_zakonczone() -> None:
    """Kontrola pozytywna, i zarazem decyzja projektowa: ostrzeżenie stałe uczy operatora
    je pomijać — ta sama zasada, co przy `vintage_skipped`."""
    uwagi = summary_notes(
        SummaryInput(
            paths=(),
            records=3,
            by_status={},
            with_phone=0,
            with_email=0,
            sheets=(),
            kind="firmy",
            run_ids=("run-a", "run-b"),
            log_path=Path("log"),
            statuses=("zakonczony", "zakonczony"),
        )
    )
    assert not any("zakończone" in u for u in uwagi)
    assert any("link_ceidg" in u for u in uwagi), "zdanie o źródle musi przetrwać przebudowę"


def test_zdanie_o_braku_zakonczonego_runu_prowadzi_do_wznow() -> None:
    """Komunikat musi powiedzieć, co zrobić: `wznow` albo jawne `--run-id`."""
    zdanie = no_finished_run(3)
    assert "3" in zdanie and "wznow" in zdanie and "--run-id" in zdanie


@pytest.fixture
def runner(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> CliRunner:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(config, "read_token_from_keyring", lambda *_: None)
    monkeypatch.setenv("COLUMNS", "200")
    monkeypatch.delenv("FORCE_COLOR", raising=False)
    monkeypatch.setenv("TERM", "dumb")
    monkeypatch.setattr(cli.console, "width", 200)
    monkeypatch.setenv("CEIDG_TOKEN", "token-testowy-nie-jwt")
    monkeypatch.setenv("CEIDG_DATA_DIR", str(tmp_path / "dane"))
    return CliRunner()


def test_eksportuj_bez_run_id_odmawia_gdy_zaden_run_nie_jest_zakonczony(
    runner: CliRunner, tmp_path: Path, clock: FakeClock
) -> None:
    """Rdzeń A7 na poziomie okablowania, nie samej decyzji.

    Przed poprawką ta sama sytuacja kończyła się **plikiem**: `list_runs()[0]` zwracał run
    przerwany, eksport szedł, a operator dostawał skoroszyt z połową wierszy i bez słowa
    o tym, że pobranie nie doszło do końca."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    store = Store(settings.store_path, environment="test", clock=clock)
    _pusty_run(store, "run-przerwany", "przerwany")
    store.close()

    wynik = runner.invoke(app, ["eksportuj", "--srodowisko", "test"])

    assert wynik.exit_code == EXIT_OK
    assert "nie jest zakończone" in wynik.output
    assert not list((tmp_path / "dane" / "wyniki").glob("*.xlsx"))


def test_eksportuj_bez_run_id_mowi_wprost_o_pustej_bazie(
    runner: CliRunner, tmp_path: Path, clock: FakeClock
) -> None:
    """Kontrola pozytywna: „brak pobrań" i „żadne nie jest zakończone" to dwie różne
    sytuacje i dwa różne zdania — zlanie ich kazałoby operatorowi zgadywać, którą ma."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    Store(settings.store_path, environment="test", clock=clock).close()

    wynik = runner.invoke(app, ["eksportuj", "--srodowisko", "test"])

    assert wynik.exit_code == EXIT_OK
    assert "Brak pobrań" in wynik.output
    assert "nie jest zakończone" not in wynik.output


# ------------------------------------------------- okablowanie ostrzeżenia (przegląd 2026-09-09)


def test_eksport_niedokonczonego_runu_dowozi_ostrzezenie_az_na_ekran(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Przegląd wykazał, że skasowanie `statuses=summary.statuses` w `flow` zostawia **cały
    pakiet 1182 testów zielony**: `ExportSummary.statuses` było policzone poprawnie i
    nieprzekazane. To jest dokładnie kształt A7 — wartość obliczona i niedowieziona — więc
    test idzie przez prawdziwe `export_and_report` aż do notatek bloku, a nie przez ręcznie
    zbudowane `SummaryInput`."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = _partia(
        deps.store,
        "run-urwany",
        od="2020-01-01",
        do="2020-12-31",
        firmy=range(1, 4),
        status="przerwany",
    )
    view = RecordingView()

    flow.export_and_report((run_id,), deps, view, out=tmp_path / "urwany.xlsx")
    deps.store.close()

    uwagi = [u for blok in view.blocks for u in blok.notes]
    assert any("nie jest zakończone" in u for u in uwagi), uwagi


def test_eksport_zakonczonego_runu_nie_dokleja_ostrzezenia(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Kontrola pozytywna do testu wyżej — inaczej „zawsze ostrzegaj" przeszłoby oba."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = _partia(
        deps.store,
        "run-caly",
        od="2020-01-01",
        do="2020-12-31",
        firmy=range(1, 4),
        status="zakonczony",
    )
    view = RecordingView()

    flow.export_and_report((run_id,), deps, view, out=tmp_path / "caly.xlsx")
    deps.store.close()

    uwagi = [u for blok in view.blocks for u in blok.notes]
    assert not any("nie jest zakończone" in u for u in uwagi), uwagi


def test_liczba_runow_w_komunikacie_nie_jest_rozmiarem_strony(store: Store) -> None:
    """`len(list_runs())` zwracało 20 przy dwudziestu pięciu runach, więc zdanie „żadne
    z 20 pobrań" mówiło o rozmiarze okna podanym jako liczba w bazie."""
    for i in range(25):
        _pusty_run(store, f"run-{i:02d}", "przerwany")

    assert len(store.list_runs()) == 20
    assert store.count_runs() == 25
