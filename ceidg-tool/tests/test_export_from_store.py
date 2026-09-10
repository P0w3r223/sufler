"""Eksport wyłącznie z bazy: podział na części, powtarzalność, zero żądań.

Istniejące testy eksportera karmią zapisywacz listą w pamięci. Tutaj źródłem jest
kursor SQLite, więc sprawdzamy to, co widać dopiero na produkcji: fabryka iteratorów
musi odtworzyć zapytanie dla planu i dla każdej części, bez gubienia i dublowania firm
(UZUPELNIENIE_01 §C i §E — „`ceidg eksportuj` odtwarza identyczny skoroszyt z bazy”).
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ceidg_tool.config import Settings
from ceidg_tool.errors import ConfigError
from ceidg_tool.exporter import write_workbook
from ceidg_tool.normalizer import (
    KOLUMNY_TYLKO_ZE_SZCZEGOLOW,
    NormalizedRecord,
    normalize,
)
from ceidg_tool.pipeline import UkryteKolumny, build_deps, count_hits, run_export
from ceidg_tool.records import RowContext
from ceidg_tool.store import Store
from tests.conftest import FakeClock, detail_record, list_record
from tests.support import criteria

CTX = RowContext(srodowisko="test", pobrano_utc="2026-09-05T10:00:00Z")
FIRMS = 7
PKD_PER_FIRM = 4


@pytest.fixture
def store(tmp_path: Path, clock: FakeClock) -> Iterator[Store]:
    s = Store(tmp_path / "store-test.sqlite", environment="test", clock=clock)
    yield s
    s.close()


def filled_run(store: Store, firms: int = FIRMS) -> str:
    """Run z pełnymi szczegółami — każda firma daje 1 wiersz Firmy, 4 PKD i 1 Spolki."""
    run_id = store.start_run(
        run_id="run-eksport",
        criteria_json='{"wojewodztwo":["podlaskie"]}',
        criteria_hash="abc",
        profile_hash="prof",
        mode="szczegoly",
        tool_version="0.1.0",
        cursor_mode="links",
    )
    store.save_page(
        run_id,
        page_index=0,
        records=[list_record(i) for i in range(1, firms + 1)],
        next_cursor=None,
    )
    store.save_details(details=[detail_record(i) for i in range(1, firms + 1)])
    store.set_stage(run_id, "gotowe")
    store.update_run_status(run_id, "zakonczony")
    return run_id


def source_from_store(store: Store, run_id: str):  # type: ignore[no-untyped-def]
    def source() -> Iterator[NormalizedRecord]:
        for raw in store.iter_run_records(run_id):
            yield normalize(raw, CTX)

    return source


def firmy_ids(path: Path) -> list[str]:
    ws = load_workbook(path)["Firmy"]
    column = [c.value for c in ws[1]].index("id")
    return [row[column] for row in ws.iter_rows(min_row=2, values_only=True)]


def test_multipart_export_from_store_keeps_every_firm_exactly_once(
    store: Store, tmp_path: Path
) -> None:
    run_id = filled_run(store)
    source = source_from_store(store, run_id)

    paths = write_workbook(
        tmp_path / "duzy.xlsx", source, metadata=[("srodowisko", "test")], row_limit=9
    )

    assert [p.name for p in paths] == [f"duzy_czesc{i:02d}.xlsx" for i in range(1, len(paths) + 1)]
    collected: list[str] = []
    for index, path in enumerate(paths, start=1):
        wb = load_workbook(path)
        assert wb.sheetnames == ["Firmy", "PKD", "Spolki", "Slownik", "Metadane"]
        assert wb["PKD"].max_row - 1 <= 9  # żaden arkusz części nie przekracza limitu
        meta = {r[0].value: r[1].value for r in wb["Metadane"].iter_rows(min_row=2)}
        assert meta["czesc"] == f"{index}/{len(paths)}"
        assert meta["firm_w_czesci"] == wb["Firmy"].max_row - 1
        collected.extend(firmy_ids(path))

    assert len(collected) == FIRMS
    assert len(set(collected)) == FIRMS  # kursor odtworzony dla każdej części, bez dubli
    assert collected == [r.id for r in store.iter_run_records(run_id)]  # kolejność zachowana


def test_single_part_export_has_no_part_metadata(store: Store, tmp_path: Path) -> None:
    run_id = filled_run(store)
    (dest,) = write_workbook(
        tmp_path / "caly.xlsx", source_from_store(store, run_id), metadata=[("srodowisko", "test")]
    )
    meta = {r[0].value: r[1].value for r in load_workbook(dest)["Metadane"].iter_rows(min_row=2)}
    assert "czesc" not in meta
    assert len(firmy_ids(dest)) == FIRMS


def test_export_repeated_from_the_same_database_gives_the_same_workbook(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Kryterium odbioru §E: ponowny eksport z bazy odtwarza ten sam skoroszyt."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = filled_run(deps.store)

    first = run_export(run_id, tmp_path / "a.xlsx", deps, cel_pobrania="odtworzenie")
    second = run_export(run_id, tmp_path / "b.xlsx", deps, cel_pobrania="odtworzenie")

    assert first.records == second.records == FIRMS
    assert first.by_status == second.by_status
    assert first.sheets == second.sheets == ("Firmy", "PKD", "Spolki", "Slownik", "Metadane")
    assert firmy_ids(tmp_path / "a.xlsx") == firmy_ids(tmp_path / "b.xlsx")

    left = load_workbook(tmp_path / "a.xlsx")
    right = load_workbook(tmp_path / "b.xlsx")
    for sheet in ("Firmy", "PKD", "Spolki", "Slownik"):
        assert list(left[sheet].values) == list(right[sheet].values), sheet
    volatile = {"eksport_utc"}
    meta_left = {r[0].value: r[1].value for r in left["Metadane"].iter_rows(min_row=2)}
    meta_right = {r[0].value: r[1].value for r in right["Metadane"].iter_rows(min_row=2)}
    assert {k: v for k, v in meta_left.items() if k not in volatile} == {
        k: v for k, v in meta_right.items() if k not in volatile
    }
    deps.store.close()


def test_offline_deps_have_no_client_so_export_cannot_reach_the_api(
    tmp_path: Path, clock: FakeClock
) -> None:
    """`eksportuj` buduje zależności bez klienta — każda próba żądania jest błędem konfiguracji."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = filled_run(deps.store)

    assert deps.client is None and deps.limiter is None
    summary = run_export(run_id, tmp_path / "out.xlsx", deps)
    assert summary.records == FIRMS

    from ceidg_tool.criteria import Criteria  # noqa: F401

    with pytest.raises(ConfigError, match="wymaga połączenia"):
        count_hits(criteria(wojewodztwo="podlaskie"), deps)
    deps.store.close()


def test_export_refuses_run_recorded_under_another_environment(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Skopiowana albo podmieniona baza: run oznaczony `prod` nie może wyjść jako `test`."""
    import sqlite3

    from ceidg_tool.errors import StoreError

    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = filled_run(deps.store)
    deps.store.close()

    conn = sqlite3.connect(settings.store_path)
    conn.execute("UPDATE run SET environment = 'prod' WHERE run_id = ?", (run_id,))
    conn.commit()
    conn.close()

    reopened = build_deps(settings, clock=clock, online=False)
    with pytest.raises(StoreError, match="środowiska prod"):
        run_export(run_id, tmp_path / "x.xlsx", reopened)
    reopened.store.close()


def test_export_of_unknown_run_id_fails_with_a_clear_message(
    tmp_path: Path, clock: FakeClock
) -> None:
    from ceidg_tool.errors import StoreError

    deps = build_deps(
        Settings(token="tok", environment="test", data_dir=tmp_path / "dane"),
        clock=clock,
        online=False,
    )
    with pytest.raises(StoreError, match="brak runu"):
        run_export("nie-ma-takiego", tmp_path / "x.xlsx", deps)
    deps.store.close()


# --- ukrywanie kolumn niedostępnych w źródle (ocena ergonomii bramki 2) ------------------


def report_run(store: Store, firms: int = 3, *, kind: str = "raport", run_id: str = "run-raport"):  # type: ignore[no-untyped-def]
    """Run ze ścieżki raportu. `kind` jest parametrem, bo o ukryciu decyduje **nie on**,
    tylko `zrodlo` przy rekordach — pierwsza wersja tego helpera zakładała odwrotnie."""
    run_id = store.start_run(
        run_id=run_id,
        criteria_json='{"wojewodztwo":["podlaskie"]}',
        criteria_hash="abc",
        profile_hash="prof",
        mode="lista",
        tool_version="0.1.0",
        cursor_mode="links",
        kind=kind,
    )
    store.save_page(
        run_id,
        page_index=0,
        records=[list_record(i) for i in range(1, firms + 1)],
        next_cursor=None,
        zrodlo="CEIDG_RAPORT",
    )
    store.set_stage(run_id, "gotowe")
    store.update_run_status(run_id, "zakonczony")
    return run_id


def hidden_columns_of(path: Path, sheet: str = "Firmy") -> set[str]:
    ws = load_workbook(path)[sheet]
    from openpyxl.utils import get_column_letter

    return {
        cell.value
        for i, cell in enumerate(ws[1], start=1)
        if ws.column_dimensions[get_column_letter(i)].hidden
    }


def test_report_export_hides_the_columns_the_csv_cannot_fill(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Ścieżka raportu zostawia kilkanaście kolumn strukturalnie pustych; w skoroszycie
    schodzą z oczu, zostają w schemacie, a `Metadane` mówią, że tam są."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = report_run(deps.store)

    summary = run_export(run_id, tmp_path / "raport.xlsx", deps)

    hidden = hidden_columns_of(summary.paths[0])
    assert "link_ceidg" in hidden and "terc" in hidden
    assert "nip" not in hidden and "nazwa" not in hidden and "telefon" not in hidden

    meta = {
        r[0].value: r[1].value
        for r in load_workbook(summary.paths[0])["Metadane"].iter_rows(min_row=2)
    }
    assert "link_ceidg" in meta["kolumny_ukryte"]
    assert "CEIDG_RAPORT" in meta["kolumny_ukryte"]
    deps.store.close()


def test_list_export_blames_the_list_and_not_the_report(tmp_path: Path, clock: FakeClock) -> None:
    """Run z API bez szczegółów też chowa kolumny — ale z innego powodu i innym zdaniem.

    Do 2026-09-10 `Metadane` miały jedno zdanie na dwa niepełne źródła, więc plik z trybu
    `lista` twierdził, że kolumny są „niedostępne w źródle CEIDG_RAPORT" — źródle, którego
    w tym pobraniu nie było. Ekran mówił prawdę, arkusz nie, a to arkusz zostaje przy pliku.
    """
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = deps.store.start_run(
        run_id="run-lista",
        criteria_json='{"wojewodztwo":["podlaskie"]}',
        criteria_hash="abc",
        profile_hash="prof",
        mode="lista",
        tool_version="0.1.0",
        cursor_mode="links",
    )
    deps.store.save_page(run_id, page_index=0, records=[list_record(1)], next_cursor=None)
    deps.store.set_stage(run_id, "gotowe")
    deps.store.update_run_status(run_id, "zakonczony")

    summary = run_export(run_id, tmp_path / "lista.xlsx", deps)

    assert "telefon" in hidden_columns_of(summary.paths[0])
    meta = {
        r[0].value: r[1].value
        for r in load_workbook(summary.paths[0])["Metadane"].iter_rows(min_row=2)
    }
    assert "telefon" in meta["kolumny_ukryte"]
    assert "CEIDG_RAPORT" not in meta["kolumny_ukryte"]
    assert "liście podstawowej" in meta["kolumny_ukryte"]
    assert meta["zrodlo"].startswith("CEIDG_API")
    deps.store.close()


def test_api_export_hides_nothing(tmp_path: Path, clock: FakeClock) -> None:
    """Ten sam kod, run typu `firmy` — nic nie znika i w `Metadane` nie ma o tym wiersza."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = filled_run(deps.store)

    summary = run_export(run_id, tmp_path / "api.xlsx", deps)

    assert hidden_columns_of(summary.paths[0]) == set()
    meta = {
        r[0].value: r[1].value
        for r in load_workbook(summary.paths[0])["Metadane"].iter_rows(min_row=2)
    }
    assert "kolumny_ukryte" not in meta
    deps.store.close()


def test_a_run_mislabelled_as_firmy_is_still_recognised_by_its_records(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Dokładnie kształt runu `0a06df59-…` z bazy produkcyjnej: etykieta `kind='firmy'`
    (zapis starszej wersji kodu), a wszystkie rekordy z `zrodlo='CEIDG_RAPORT'`.

    Predykat liczony z danych rozpoznaje go poprawnie, więc kolumny są ukryte, a podsumowanie
    mówi prawdę o `link_ceidg` — bez poprawiania czegokolwiek w bazie."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    run_id = report_run(deps.store, kind="firmy", run_id="run-zle-oznaczony")
    assert deps.store.get_run(run_id).kind == "firmy"  # etykieta nadal kłamie

    summary = run_export(run_id, tmp_path / "raport.xlsx", deps)

    assert summary.kind == "raport"
    assert "link_ceidg" in hidden_columns_of(summary.paths[0])
    deps.store.close()


def test_a_mixed_export_hides_nothing_and_keeps_the_run_label(
    tmp_path: Path, clock: FakeClock
) -> None:
    """Zestaw mieszany: te same kolumny bywają wypełnione przez ścieżkę API, więc ukrycie
    zgubiłoby dane. Jeden predykat rozstrzyga i o ukryciu, i o zdaniu w podsumowaniu."""
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    deps = build_deps(settings, clock=clock, online=False)
    z_api = filled_run(deps.store)
    z_raportu = report_run(deps.store, run_id="run-mieszany")

    summary = run_export([z_api, z_raportu], tmp_path / "mieszany.xlsx", deps)

    # `== "firmy"`, nie `!= "raport"`: nierówność przeszłaby też dla pustego napisu albo
    # dla etykiety, której nie zna `texts.summary_table`
    assert summary.kind == "firmy"
    assert hidden_columns_of(summary.paths[0]) == set()
    deps.store.close()


def test_the_export_command_speaks_while_it_writes(tmp_path: Path, clock: FakeClock) -> None:
    """`eksportuj` musi mieć kanał postępu — inaczej pełne województwo to dziesięć minut ciszy.

    `cli.eksportuj` budowało zależności bez `events`, więc dostawało `NullEvents` i znikało
    z niego nie tylko zadanie paska, ale i zdanie „Zapisuję skoroszyt…". Komentarz nad tym
    zdaniem w `pipeline` obiecuje, że pada ono „nawet tam, gdzie paska nie widać (log, tryb
    cichy)" — a `eksportuj` był jedyną drogą, na której nie padało nigdzie (audyt 2026-09-07).
    Zmierzone tempo to 453 firmy na sekundę, więc cisza rosła liniowo z rozmiarem pobrania.
    """
    import io as _io

    from rich.console import Console

    from ceidg_tool.console import ConsoleEvents

    bufor = _io.StringIO()
    settings = Settings(token="tok", environment="test", data_dir=tmp_path / "dane")
    events = ConsoleEvents(Console(file=bufor, width=120))
    deps = build_deps(settings, clock=clock, events=events, online=False)
    run_id = filled_run(deps.store)

    run_export(run_id, tmp_path / "out.xlsx", deps)
    events.close()

    assert "Zapisuję skoroszyt" in bufor.getvalue()
    deps.store.close()


def test_ukryte_kolumny_bez_kolumn_sa_odmawiane() -> None:
    """Pusty zbiór z powodem to ostrzeżenie o niczym — `None` znaczy „nic nie ukryto".

    Docstring `UkryteKolumny` obiecywał to od początku, ale `dataclass` sam z siebie niczego
    nie sprawdza: `UkryteKolumny("lista", frozenset())` przechodziło, a `build_metadata`
    dopisywało wtedy do arkusza wiersz `kolumny_ukryte` zakończony dwukropkiem i niczym.
    Dziś nie ma jak tego wywołać — oba źródła są niepustymi stałymi — więc test pilnuje
    trzeciego źródła, którego jeszcze nie ma.
    """
    with pytest.raises(ValueError, match="użyj None"):
        UkryteKolumny("lista", frozenset())

    dozwolone = UkryteKolumny("lista", KOLUMNY_TYLKO_ZE_SZCZEGOLOW)
    assert dozwolone.kolumny == KOLUMNY_TYLKO_ZE_SZCZEGOLOW
