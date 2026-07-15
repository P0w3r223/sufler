"""Materializacja załączników Teams: referencja → bajty → ``Attachment`` (base64/tekst).

Tu żyje I/O (pobranie z Graph przez wstrzyknięty port) + kodowanie base64 + rozpoznanie i
przetworzenie obrazów (sniff po ZAWARTOŚCI + downscaling) + ekstrakcja tekstu z dokumentów
(``.docx``/``.xlsx``/``.pptx`` i plików tekstowych — Claude API nie przyjmuje ich natywnie) +
walidacja limitów. Biblioteki (``python-docx``/``openpyxl``/``python-pptx``/``Pillow``)
importowane LENIWIE (extra ``teams-graph``). Błąd/limit/nieobsługiwany typ NIE kładzie pollera
— zamiast bajtów wstawiamy krótką notkę tekstową (agent poinformuje użytkownika). Treść
załącznika to DANE, nie polecenia — nie interpretujemy jej tutaj.
"""
from __future__ import annotations

import base64
import io
import logging
from dataclasses import dataclass
from typing import TYPE_CHECKING

from workmate.core.ports.llm import Attachment

if TYPE_CHECKING:
    from workmate.adapters.inbound.teams_graph.poller import GraphChannelClient
    from workmate.adapters.inbound.teams_graph.selection import (
        AttachmentRef,
        ChannelMessage,
    )

logger = logging.getLogger(__name__)

_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
_JPEG_MAGIC = b"\xff\xd8\xff"
# Formaty obrazów akceptowane przez blok ``image`` API → nasz media_type.
_SUPPORTED_IMAGE = {
    "PNG": "image/png",
    "JPEG": "image/jpeg",
    "GIF": "image/gif",
    "WEBP": "image/webp",
}
# Formaty źródłowe FOTOGRAFICZNE nieakceptowane natywnie przez API — konwertuj do JPEG, nie PNG:
# HEIF/HEIC (domyślny format iPhone) jako PNG dałby wielomegabajtowy plik, który po downscalingu
# i tak może przekroczyć max_bytes; JPEG q85 mieści fotografię w setkach kB. Pillow+pillow-heif
# ustawia ``img.format="HEIF"`` dla .heic i .heif.
_PHOTO_SOURCE_FORMATS = {"HEIF"}
_JPEG_QUALITY = 85
# Rozszerzenia traktowane jako czysty tekst (dekodowanie UTF-8, bez base64).
_TEXT_EXTS = {"txt", "md", "csv", "log", "json", "xml", "yaml", "yml"}
# Górne capy ekstrakcji — chronią przed absurdalnie dużym plikiem (bloki wracają co turę).
_MAX_TEXT_CHARS = 200_000
_MAX_SHEET_ROWS = 2000
# Sufit liczby pikseli obrazu do LOKALNEGO dekodowania (downscaling). Powyżej nie dekodujemy
# (bomba dekompresji / wielki skan mógłby zjeść setki MB RAM i położyć pollera) — oddajemy
# oryginał (Anthropic skaluje serwerowo) albo degradujemy do notki. 40 MP ≈ 160 MB RGBA.
_MAX_IMAGE_PIXELS = 40_000_000


@dataclass(frozen=True)
class AttachmentLimits:
    """Granice materializacji: rozmiary, liczba, ŁĄCZNY budżet base64 i próg downscalingu.

    Rozróżniamy dwie klasy: pliki wysyłane jako BASE64 (obraz/PDF) obowiązuje ``max_bytes`` i
    ŁĄCZNY ``max_total_bytes`` (sufit 32 MB żądania API); pliki EKSTRAHOWANE do tekstu
    (docx/xlsx/pptx/txt) dostają wyższy ``max_extract_bytes`` i NIE liczą się do budżetu API
    (wysyłamy z nich sam tekst, nie bajty). ``max_extract_bytes`` jest zarazem UNIWERSALNYM
    twardym capem surowego pobrania (pierwsza bramka dla każdego typu). Budżet liczymy w bajtach
    SUROWYCH — przeliczik na base64 (~1.33×) siedzi w suficie configu (24 MB raw ≈ 32 MB API).
    """

    max_bytes: int  # pojedynczy plik base64 (obraz/PDF), w bajtach surowych
    max_count: int  # liczba na wiadomość
    max_total_bytes: int  # łączny budżet base64 (bajty surowe; sufit 24 MB ≈ 32 MB API)
    max_extract_bytes: int = 50 * 1024 * 1024  # sufit tekstu ORAZ uniwersalny cap pobrania
    max_image_edge: int = 2048  # dłuższa krawędź obrazu (px) — powyżej skalujemy w dół


class AttachmentMaterializer:
    """Zamienia ``AttachmentRef`` na ``Attachment``: pobiera bajty i koduje/ekstrahuje.

    Klient Graph (port ``GraphChannelClient`` z metodami binarnymi) jest wstrzykiwany —
    pełną materializację testujemy atrapą portu bez sieci.
    """

    def __init__(
        self, client: GraphChannelClient, *, limits: AttachmentLimits
    ) -> None:
        self._client = client
        self._limits = limits

    async def materialize(
        self, team_id: str, channel_id: str, msg: ChannelMessage
    ) -> tuple[Attachment, ...]:
        """Zmaterializuj referencje wiadomości; ponad limit → notka, reszta pobrana.

        Egzekwuje trzy granice: liczbę (max_count), rozmiar pliku (max_bytes) i ŁĄCZNY
        budżet bajtów (max_total_bytes) — bez tego kilka plików mogłoby przekroczyć sufit
        32 MB żądania API (bloki są odtwarzane co turę).
        """
        refs = msg.attachment_refs
        out: list[Attachment] = []
        kept = refs[: self._limits.max_count]
        if len(refs) > len(kept):
            dropped = len(refs) - len(kept)
            out.append(_note(f"Pominięto {dropped} załącznik(ów) — limit na wiadomość."))
        total = 0
        for ref in kept:
            remaining = self._limits.max_total_bytes - total
            att, size = await self._materialize_one(
                team_id, channel_id, msg.thread_root_id, msg.id, ref, budget=remaining
            )
            out.append(att)
            total += size
        return tuple(out)

    async def _materialize_one(
        self,
        team_id: str,
        channel_id: str,
        root_id: str,
        message_id: str,
        ref: AttachmentRef,
        *,
        budget: int,
    ) -> tuple[Attachment, int]:
        """Pobierz i zbuduj załącznik; zwróć (Attachment, bajty_zaliczone_do_budżetu_API).

        Pobranie i budowanie mają osobne ``try``: błąd POBRANIA (np. 404 wklejonego obrazu AMS)
        to stan oczekiwany → ``warning`` + rzeczowa notka (bez straszącego tracebacku); błąd
        BUDOWANIA (uszkodzony plik) → ``exception`` + notka. Każdy błąd degraduje do notki, nigdy
        nie zapętla pollera. Limit base64 liczymy na WYNIKU (po downscalingu obrazu); pliki
        ekstrahowane do tekstu podlegają wyższemu ``max_extract_bytes`` i zaliczają 0 do budżetu.
        """
        try:
            if ref.kind == "hosted":
                data = await self._client.get_hosted_content(
                    team_id, channel_id, root_id, message_id, ref.hosted_id
                )
            elif ref.kind == "url":
                data = await self._client.download_public_url(ref.url)
            else:
                data = await self._client.download_shared_url(ref.url)
        except Exception as exc:
            logger.warning("Nie pobrano załącznika %s (status %s).", ref.name, _http_status(exc))
            if ref.kind == "hosted":
                return _note(
                    f"Załącznika „{ref.name}” (obraz wklejony w treści wiadomości) nie udało się "
                    "pobrać. Można go wysłać ponownie jako osobny plik."
                ), 0
            return _note(f"Załącznika „{ref.name}” nie udało się pobrać."), 0

        try:
            if len(data) > self._limits.max_extract_bytes:
                return _note(
                    f"Załącznika „{ref.name}” nie udało się odczytać: plik jest za duży."
                ), 0
            result = _build(ref, data, max_image_edge=self._limits.max_image_edge)
        except Exception:
            logger.exception("Nie udało się przetworzyć załącznika %s", ref.name)
            return _note(f"Załącznika „{ref.name}” nie udało się przetworzyć."), 0

        if result is None:
            return _note(
                f"Załącznika „{ref.name}” nie udało się odczytać: nieobsługiwany typ pliku."
            ), 0
        built, sent = result
        if sent > self._limits.max_bytes:
            return _note(
                f"Załącznika „{ref.name}” nie udało się odczytać: przekracza limit rozmiaru."
            ), 0
        if sent > budget:
            return _note(
                f"Załącznika „{ref.name}” nie udało się odczytać: przekroczony łączny limit "
                "załączników wiadomości."
            ), 0
        return built, sent


def _build(
    ref: AttachmentRef, data: bytes, *, max_image_edge: int
) -> tuple[Attachment, int] | None:
    """Zbuduj ``(Attachment, bajty_base64)`` z bajtów; ``None`` gdy typ nieobsługiwany.

    Drugi element to liczba bajtów, które FAKTYCZNIE pójdą do API jako base64 (obraz PO
    downscalingu, PDF w całości) — 0 dla plików ekstrahowanych do tekstu (nie wysyłamy z nich
    bajtów). Kolejność: najpierw OBRAZ po ZAWARTOŚCI (nie po rozszerzeniu) — łapie png/jpg/gif/
    webp wklejone inline ORAZ załączone jako plik, niezależnie od nazwy. Dopiero potem plik po
    rozszerzeniu (dokumenty/tekst). Obrazy inline (hosted) mogą być WYŁĄCZNIE obrazem.
    """
    image = _process_image(data, max_image_edge)
    if image is not None:
        media_type, out = image
        return Attachment("image", media_type, ref.name, data_base64=_b64(out)), len(out)
    if ref.kind == "hosted":
        return None  # inline, ale nie rozpoznany jako obraz
    ext = ref.name.rsplit(".", 1)[-1].lower() if "." in ref.name else ""
    if ext == "pdf":
        pdf = Attachment("document", "application/pdf", ref.name, data_base64=_b64(data))
        return pdf, len(data)
    if ext == "docx":
        return Attachment("text", "text/plain", ref.name, text=_extract_docx(data)), 0
    if ext == "xlsx":
        return Attachment("text", "text/plain", ref.name, text=_extract_xlsx(data)), 0
    if ext == "pptx":
        return Attachment("text", "text/plain", ref.name, text=_extract_pptx(data)), 0
    if ext in _TEXT_EXTS:
        return Attachment("text", "text/plain", ref.name, text=_extract_text(data)), 0
    return None


_heif_state: bool | None = None  # None=nie próbowano; True=zarejestrowano; False=brak wtyczki


def _ensure_heif_registered() -> None:
    """Zarejestruj opener HEIF w Pillow RAZ na proces (idempotentnie); brak wtyczki = cichy no-op.

    ``pillow-heif`` żyje w extra ``teams-graph``; jego brak NIE może wywalić importu modułu ani
    pollera — bez niego HEIC się nie otworzy i zejdzie notką (jak inne nieobsługiwane typy).
    """
    global _heif_state
    if _heif_state is not None:
        return
    try:
        import pillow_heif

        pillow_heif.register_heif_opener()
        _heif_state = True
    except Exception:  # brak extra albo błąd rejestracji — degradujemy, nie wywracamy
        _heif_state = False


def _process_image(data: bytes, max_edge: int) -> tuple[str, bytes] | None:
    """Rozpoznaj i przetwórz obraz: ``(media_type, bajty)`` albo ``None`` (to nie obraz).

    Pillow (jeśli dostępny) otwiera bajty — to sniff po ZAWARTOŚCI, a nie po rozszerzeniu.
    Gdy dłuższa krawędź > ``max_edge`` → skalujemy w dół (koszt tokenów; bloki wracają co turę).
    Format nieobsługiwany przez API (np. BMP/TIFF), ale otwieralny → konwersja do PNG. Gdy
    Pillow brak/nie otworzy → fallback na sniff magicznych bajtów (bez downscalingu), a gdy i to
    nie rozpozna → ``None``. Nigdy nie rzuca (egress nie może paść).
    """
    try:
        from PIL import Image
    except ImportError:
        return _sniff_image(data)
    _ensure_heif_registered()  # idempotentnie włącza dekoder HEIC/HEIF, jeśli extra obecny

    try:
        with Image.open(io.BytesIO(data)) as img:
            # Rozmiar czytamy z NAGŁÓWKA (bez dekodowania). Zbyt duży obraz nie jest dekodowany
            # lokalnie — oddajemy oryginał (Anthropic skaluje serwerowo) albo później zejdzie
            # notką na limicie rozmiaru; chroni pollera przed skokiem pamięci / bombą dekompresji.
            if img.width * img.height > _MAX_IMAGE_PIXELS:
                return _sniff_image(data)
            img.load()  # wymuś dekodowanie — atrapy/uszkodzone tu rzucą (→ fallback)
            fmt = img.format or ""
            width, height = img.size
            media_type = _SUPPORTED_IMAGE.get(fmt)
            longest = max(width, height)
            if media_type is not None and longest <= max_edge:
                return media_type, data  # bez zmian (zachowaj oryginał/animację)
            if media_type is not None:
                out_format, out_media = fmt, media_type
            elif fmt in _PHOTO_SOURCE_FORMATS:
                # Fotografia (HEIF/HEIC) → JPEG: PNG dałby wielomegabajtowy plik (ryzyko max_bytes).
                out_format, out_media = "JPEG", "image/jpeg"
            else:
                out_format, out_media = "PNG", "image/png"
            scaled: Image.Image = img
            if longest > max_edge:
                ratio = max_edge / longest
                scaled = img.resize(
                    (max(1, int(width * ratio)), max(1, int(height * ratio))),
                    Image.Resampling.LANCZOS,
                )
            # JPEG nie zapisze trybu z alfą/paletą; PNG nie zapisze CMYK — ujednolić do RGB.
            if (out_format == "JPEG" and scaled.mode not in ("RGB", "L")) or (
                out_format == "PNG" and scaled.mode == "CMYK"
            ):
                scaled = scaled.convert("RGB")
            buffer = io.BytesIO()
            if out_format == "JPEG":
                scaled.save(buffer, format="JPEG", quality=_JPEG_QUALITY, optimize=True)
            else:
                scaled.save(buffer, format=out_format)
            return out_media, buffer.getvalue()
    except Exception:
        return _sniff_image(data)


def _sniff_image(data: bytes) -> tuple[str, bytes] | None:
    """Fallback bez Pillow: rozpoznaj typ obrazu po magicznych bajtach (bez downscalingu)."""
    if data.startswith(_PNG_MAGIC):
        return "image/png", data
    if data.startswith(_JPEG_MAGIC):
        return "image/jpeg", data
    if data[:4] == b"GIF8":
        return "image/gif", data
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", data
    return None


def _extract_docx(data: bytes) -> str:
    """Wyciągnij tekst z .docx: akapity + komórki tabel (``python-docx``, import leniwy)."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    """Wyciągnij tekst z .xlsx: per arkusz nagłówek + wiersze (``openpyxl``, import leniwy)."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"# Arkusz: {sheet.title}")
            rows = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if not cells:
                    continue
                parts.append(" | ".join(cells))
                rows += 1
                if rows >= _MAX_SHEET_ROWS:
                    parts.append("… (obcięto wiersze)")
                    break
        return "\n".join(parts).strip()
    finally:
        workbook.close()


def _extract_pptx(data: bytes) -> str:
    """Wyciągnij tekst z .pptx: per slajd tekst z kształtów (``python-pptx``, import leniwy)."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"# Slajd {index}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts).strip()


def _extract_text(data: bytes) -> str:
    """Zdekoduj plik tekstowy (UTF-8, nieznane bajty zastąpione) z górnym capem długości."""
    text = data.decode("utf-8", errors="replace")
    if len(text) > _MAX_TEXT_CHARS:
        text = text[:_MAX_TEXT_CHARS] + "\n… (obcięto)"
    return text.strip()


def _b64(data: bytes) -> str:
    """Base64 bez znaków nowej linii (wymóg bloków ``image``/``document`` API)."""
    return base64.standard_b64encode(data).decode("ascii")


def _http_status(exc: Exception) -> int | None:
    """Kod HTTP z wyjątku httpx (duck typing, bez importu httpx) — do logu, ``None`` gdy brak."""
    return getattr(getattr(exc, "response", None), "status_code", None)


def _note(message: str) -> Attachment:
    """Rzeczowa notka o statusie załącznika zamiast bajtów (limit/błąd/nieobsługiwany typ).

    Sformułowana jak FAKT o załączniku, nie jak „instrukcja systemowa": agent czyta treść
    wiadomości jak dane (granica z prompt-injection), więc notka udająca instrukcję systemową
    była odrzucana i piętnowana. Neutralny opis statusu jest po prostu relacjonowany użytkownikowi.
    """
    return Attachment("text", "text/plain", "status załącznika", text=message)
