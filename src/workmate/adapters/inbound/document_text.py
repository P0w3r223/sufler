"""Ekstrakcja tekstu z dokumentów (docx/xlsx/pptx/pdf/html/tekst) — jedno źródło dla adapterów.

Powstało z materializera załączników teams-graph (ADR 0016), gdzie te same funkcje bytes→str
żyły prywatnie. Importer korpusu (W0, ``seed_corpus``) potrzebuje dokładnie tego samego —
zamiana Worda/Excela/PDF na czysty tekst notatki — więc wyciągamy ekstraktory tu, by nie
dublować logiki (DRY na TYM SAMYM pojęciu). Załączniki i seed czytają z jednego miejsca.

Wszystkie funkcje są CZYSTE (``bytes``/``Path`` → ``str``), a importy bibliotek są LENIWE:
python-docx/openpyxl/python-pptx/pypdf żyją w extra (``teams-graph`` / ``seed``), więc rdzeń i
serwer MCP zostają lekkie. Treść dokumentu to DANE — kopiujemy ją wiernie, nie interpretujemy.
"""

from __future__ import annotations

import io
from html.parser import HTMLParser
from pathlib import Path

# Rozszerzenia traktowane jako czysty tekst (dekodowanie UTF-8, bez ekstraktora binarnego).
TEXT_EXTS = frozenset({"txt", "md", "csv", "log", "json", "xml", "yaml", "yml"})
# Rozszerzenia dokumentów z dedykowanym ekstraktorem tekstu. HTML jest TUTAJ, a nie w
# ``TEXT_EXTS`` (ADR 0064): zdekodowany jako zwykły tekst oddałby modelowi surowy znacznik —
# treść utopioną w atrybutach i stylach zamiast tego, co człowiek na tej stronie widzi.
BINARY_EXTS = frozenset({"docx", "xlsx", "pptx", "pdf", "html", "htm"})
# Wszystkie rozszerzenia, które importer/materializer umie zamienić na tekst.
SUPPORTED_EXTS = TEXT_EXTS | BINARY_EXTS

# Górne capy ekstrakcji — chronią przed absurdalnie dużym plikiem (materializer wysyła tekst co
# turę; importer trzyma go w jednej notatce). Wspólne dla obu konsumentów.
_MAX_TEXT_CHARS = 200_000
_MAX_SHEET_ROWS = 2000


class DocumentExtractionError(Exception):
    """Ekstrakcja pliku nie powiodła się (uszkodzony/zaszyfrowany dokument, brak biblioteki)."""


def _cap(text: str) -> str:
    """Przytnij tekst do ``_MAX_TEXT_CHARS`` — bezpiecznik przed olbrzymią pojedynczą notatką."""
    if len(text) > _MAX_TEXT_CHARS:
        return text[:_MAX_TEXT_CHARS] + "\n… (obcięto)"
    return text


def extract_docx(data: bytes) -> str:
    """Wyciągnij tekst z .docx: akapity + komórki tabel (``python-docx``, import leniwy)."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _cap("\n".join(parts).strip())


def extract_xlsx(data: bytes) -> str:
    """Wyciągnij tekst z .xlsx: per arkusz nagłówek + wiersze (``openpyxl``, import leniwy)."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            parts.append(f"# Arkusz: {sheet.title}")
            rows = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if not cells:
                    continue
                parts.append(" | ".join(cells))
                rows += 1
                if rows >= _MAX_SHEET_ROWS:
                    parts.append("… (obcięto wiersze)")
                    break
        return "\n".join(parts).strip()
    finally:
        workbook.close()


def extract_pptx(data: bytes) -> str:
    """Wyciągnij tekst z .pptx: per slajd tekst z kształtów (``python-pptx``, import leniwy)."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"# Slajd {index}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return _cap("\n".join(parts).strip())


def extract_pdf(data: bytes) -> str:
    """Wyciągnij tekst z .pdf: per strona nagłówek + tekst warstwy tekstowej (``pypdf``, leniwy).

    Bez OCR — skan bez warstwy tekstowej zwróci pusto (importer potraktuje to jak pusty dokument,
    nie jak błąd). Twarde szyfrowanie/uszkodzenie podnosi ``DocumentExtractionError``.
    """
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(f"# Strona {index}")
                parts.append(text)
        joined = "\n".join(parts).strip()
    except PyPdfError as exc:
        raise DocumentExtractionError(f"nieczytelny PDF: {exc}") from exc
    return _cap(joined)


def extract_text(data: bytes) -> str:
    """Zdekoduj plik tekstowy (UTF-8, nieznane bajty zastąpione) z górnym capem długości."""
    return _cap(data.decode("utf-8", errors="replace").strip())


# Znaczniki, których ZAWARTOŚĆ nie jest treścią strony — kod i grafika wektorowa. Pomijamy je
# w całości, inaczej model dostałby JavaScript i reguły CSS jako „tekst dokumentu".
_HTML_SKIP_TAGS = frozenset({"script", "style", "noscript", "template", "svg"})
# Znaczniki blokowe — kończą linię, żeby akapity i wiersze tabel nie skleiły się w jeden ciąg.
_HTML_BLOCK_TAGS = frozenset(
    {
        "p",
        "div",
        "br",
        "hr",
        "li",
        "tr",
        "td",
        "th",
        "section",
        "article",
        "header",
        "footer",
        "nav",
        "aside",
        "main",
        "table",
        "ul",
        "ol",
        "dl",
        "dt",
        "dd",
        "blockquote",
        "pre",
        "figure",
        "figcaption",
        "form",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
    }
)


class _HtmlTextExtractor(HTMLParser):
    """Zbiera widoczny tekst strony; treść to DANE — przepisujemy ją, nie interpretujemy.

    Sedno anty-maskowania (ADR 0064): strona, której ~95% powierzchni zajmuje obraz, ma realną
    treść w resztce. Parser tekstowy jest na to z natury odporny — grafiki nie widzi wcale —
    ale gubiłby to, co obraz NIESIE. Dlatego ``alt`` wchodzi do tekstu, a obrazy bez opisu są
    LICZONE i raportowane jedną linią na końcu: model dowiaduje się, że strona jest graficzna,
    nie dostając setki pustych znaczników.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0
        self._images_without_alt = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _HTML_SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "img":
            alt = (dict(attrs).get("alt") or "").strip()
            if alt:
                self._parts.append(f"\n[obraz: {alt}]\n")
            else:
                self._images_without_alt += 1
            return
        if tag in _HTML_BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _HTML_SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag in _HTML_BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        collapsed = " ".join(data.split())
        if collapsed:
            self._parts.append(collapsed + " ")

    def result(self) -> str:
        """Złóż fragmenty w linie + dopisz podsumowanie obrazów bez opisu.

        Puste linie znikają: znacznik blokowy zamyka i otwiera linię, więc ``</p><p>`` dałoby
        pustkę przy każdym akapicie, a tabela podwoiłaby swoją długość. Tak samo składają tekst
        pozostałe ekstraktory (docx/pptx łączą niepuste akapity pojedynczym ``\\n``) — jedna
        konwencja dla wszystkich formatów.
        """
        text = "".join(self._parts)
        lines = [line.strip() for line in text.split("\n")]
        joined = "\n".join(line for line in lines if line)
        if self._images_without_alt:
            joined += f"\n\n[{self._images_without_alt} obraz(ów) bez opisu tekstowego]"
        return joined.strip()


def extract_html(data: bytes) -> str:
    """Wyciągnij widoczny tekst ze strony HTML (``html.parser`` ze stdlib — bez zależności).

    Nie renderujemy i nie liczymy powierzchni — o „maskowaniu" rozstrzyga tu sam fakt, że
    ekstraktor czyta strukturę, a nie wygląd: tekst zajmujący 5% ekranu waży tyle samo, co
    baner na całą stronę. Uzupełnia to budżet materializacji, który pilnuje, żeby bajty obrazu
    nie wyparły tego tekstu z kontekstu.
    """
    parser = _HtmlTextExtractor()
    parser.feed(data.decode("utf-8", errors="replace"))
    parser.close()
    return _cap(parser.result())


_BINARY_EXTRACTORS = {
    "docx": extract_docx,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "pdf": extract_pdf,
    "html": extract_html,
    "htm": extract_html,
}


def extract_text_from_bytes(data: bytes, ext: str) -> str:
    """Zamień bajty na tekst wg rozszerzenia (bez kropki, lowercase). ``KeyError``-free dyspozytor.

    Nieobsługiwane rozszerzenie → ``DocumentExtractionError`` (wołający decyduje: pominąć/zgłosić).
    Błąd biblioteki ekstrahującej opakowujemy tym samym wyjątkiem, by konsument miał jeden typ.
    """
    ext = ext.lower()
    if ext in TEXT_EXTS:
        return extract_text(data)
    extractor = _BINARY_EXTRACTORS.get(ext)
    if extractor is None:
        raise DocumentExtractionError(f"nieobsługiwane rozszerzenie: .{ext}")
    try:
        return extractor(data)
    except DocumentExtractionError:
        raise  # zachowuje precyzyjny komunikat (np. „nieczytelny PDF: …")
    except Exception as exc:  # brak biblioteki (extra) lub uszkodzony plik nie kładzie partii
        raise DocumentExtractionError(f"błąd ekstrakcji .{ext}: {exc}") from exc


def extract_text_from_path(path: Path) -> str:
    """Wczytaj plik i zwróć tekst wg rozszerzenia; ``DocumentExtractionError`` gdy się nie da."""
    ext = path.suffix.lstrip(".").lower()
    return extract_text_from_bytes(path.read_bytes(), ext)
